import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import analyze_transferred_dicom_by_center as core
import analyze_transferred_dicom_by_center_linux as linux_center_web
from build_center_image_workbook_linux import write_center_workbook_linux


OPENPYXL_AVAILABLE = importlib.util.find_spec("openpyxl") is not None


@unittest.skipUnless(OPENPYXL_AVAILABLE, "openpyxl未安装")
class LinuxCenterWorkbookTests(unittest.TestCase):
    def test_reads_center_workbook_without_node(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "centers.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "分中心影像"
            worksheet.append(["中心名称", "姓名", "住院号/门诊号/放射号", "手术日期"])
            worksheet.append(["宜昌市中心人民医院", "测试患者", "H001", "2025/1/1"])
            workbook.save(path)
            workbook.close()

            sheets = linux_center_web.read_center_workbook_linux(path)

        self.assertEqual(sheets[0]["name"], "分中心影像")
        self.assertEqual(sheets[0]["values"][1][2], "H001")

    def test_linux_writer_preserves_layout_white_fill_and_black_borders(self):
        from openpyxl import load_workbook

        payload = {
            "rule_version": core.RULE_VERSION,
            "generated_at": "2026-09-01T13:30:00",
            "headers": core.WEB_HEADERS,
            "centers": [
                {
                    "sheet_name": "宜昌",
                    "source_center_name": "宜昌",
                    "reference_center_name": "宜昌市中心人民医院",
                    "rows": [
                        {
                            "住院号": "H001",
                            "患者": "测试患者",
                            "StudyInstanceUID": "1.2.3",
                            "SeriesUID": "1.2.3.4",
                            "SOPUID": "NA",
                            "AcqusitionDate": 20250101,
                            "时期": "术中",
                            "影像类型": "['ORIGINAL']",
                            "序列描述": "Test",
                            "SliceThickness": "NA",
                            "NumberOfFrames": "NA",
                            "帧数": 1,
                            "modality": "XA",
                            "第一拍摄角度": "NA",
                            "第二拍摄角度": "NA",
                        }
                    ],
                    "counters": {
                        "center_match_basis": "中心目录简称唯一包含匹配",
                        "files_seen": 1,
                        "dicom": 1,
                        "patients": 1,
                        "studies": 1,
                        "series": 1,
                        "output_rows": 1,
                        "matched_rows": 1,
                        "exceptions": 0,
                    },
                }
            ],
            "exceptions": [],
        }
        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "影像-web.xlsx"
            write_center_workbook_linux(output, payload)
            workbook = load_workbook(output, data_only=False)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    ["宜昌", "处理异常", "运行摘要", "字段说明"],
                )
                worksheet = workbook["宜昌"]
                self.assertEqual(worksheet["A2"].value, "H001")
                self.assertEqual(worksheet["G2"].value, "术中")
                self.assertFalse(worksheet.sheet_view.showGridLines)
                self.assertEqual(worksheet["A2"].fill.fgColor.rgb[-6:], "FFFFFF")
                self.assertEqual(worksheet["A2"].border.left.style, "thin")
                self.assertEqual(worksheet["A2"].border.left.color.rgb[-6:], "000000")
                self.assertEqual(worksheet["A1"].border.bottom.style, "medium")
                self.assertEqual(worksheet.freeze_panes, "A2")
                self.assertEqual(len(worksheet.tables), 1)
                self.assertEqual(workbook["运行摘要"]["J2"].value, 1)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
