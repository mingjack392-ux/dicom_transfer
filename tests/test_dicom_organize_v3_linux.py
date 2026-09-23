import importlib.util
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import dicom_organize_v3_linux as linux_v3


class LinuxReportOnlyTests(unittest.TestCase):
    def test_review_rows_are_recovered_from_grouped_mapping(self):
        rows = [
            {
                "master_patient_key": "G000001",
                "canonical_name": "王玉梅",
                "patient_id": "A001",
                "birth_date": "19740317",
                "sex": "F",
                "institutions": "xiongan_xuanwu_hospital",
                "match_status": "grouped_needs_review",
                "match_basis": "same_name;sex+institution;conflict=birth_date",
            },
            {
                "master_patient_key": "G000001",
                "canonical_name": "王玉梅",
                "patient_id": "B002",
                "birth_date": "19740226",
                "sex": "F",
                "institutions": "xiongan_xuanwu_hospital",
                "match_status": "grouped_needs_review",
                "match_basis": "same_name;sex+institution;conflict=birth_date",
            },
        ]

        review = linux_v3._review_rows_from_mapping(rows)

        self.assertEqual(len(review), 1)
        self.assertEqual(review[0]["review_status"], "grouped_needs_review")
        self.assertEqual(review[0]["left_patient_id"], "A001")
        self.assertEqual(review[0]["right_patient_id"], "B002")

    def test_report_only_uses_existing_json_state_without_scanning(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            destination = root / "output"
            state = destination / ".dicom_v3_state"
            state.mkdir(parents=True)
            (state / "study_inventory.json").write_text(
                json.dumps(
                    {
                        "rows": [
                            {
                                "patient_id": "A001",
                                "patient_name": "甲",
                                "study_uid": "1.2.3",
                                "file_count": 10,
                                "series_count": 1,
                                "confidence": 0.99,
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            (state / "identity_mapping.json").write_text(
                json.dumps(
                    {
                        "rows": [
                            {
                                "master_patient_key": "G000001",
                                "canonical_name": "甲",
                                "patient_id": "A001",
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            with (
                patch.object(linux_v3, "require_openpyxl"),
                patch.object(linux_v3, "write_study_workbook_linux") as writer,
            ):
                status = linux_v3.report_only(
                    [str(root / "source"), str(destination), "--report-only"]
                )

            self.assertEqual(status, 0)
            writer.assert_called_once()
            self.assertEqual(writer.call_args.args[1][0]["study_uid"], "1.2.3")
            reports = writer.call_args.args[2]
            mapping_report = next(
                report for report in reports if report["sheet_name"] == "患者身份映射"
            )
            self.assertEqual(mapping_report["rows"][0]["patient_id"], "A001")

    def test_report_only_groups_multiple_patient_ids_by_master_identity(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            destination = root / "output"
            state = destination / ".dicom_v3_state"
            state.mkdir(parents=True)
            (state / "study_inventory.json").write_text(
                json.dumps(
                    {
                        "rows": [
                            {
                                "patient_id": "A001",
                                "patient_name": "Alpha Patient",
                                "exam_date": "2026-07-12",
                                "study_uid": "study-alpha-late",
                            },
                            {
                                "patient_id": "B001",
                                "patient_name": "Beta Patient",
                                "exam_date": "2026-01-18",
                                "study_uid": "study-beta",
                            },
                            {
                                "patient_id": "C001",
                                "patient_name": "Alpha Patient",
                                "exam_date": "2025-04-17",
                                "study_uid": "study-alpha-early",
                            },
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            (state / "identity_mapping.json").write_text(
                json.dumps(
                    {
                        "rows": [
                            {
                                "master_patient_key": "G000001",
                                "canonical_name": "Alpha Patient",
                                "patient_id": "A001",
                            },
                            {
                                "master_patient_key": "G000002",
                                "canonical_name": "Beta Patient",
                                "patient_id": "B001",
                            },
                            {
                                "master_patient_key": "G000001",
                                "canonical_name": "Alpha Patient",
                                "patient_id": "C001",
                            },
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            with (
                patch.object(linux_v3, "require_openpyxl"),
                patch.object(linux_v3, "write_study_workbook_linux") as writer,
            ):
                status = linux_v3.report_only(
                    [str(root / "source"), str(destination), "--report-only"]
                )

            self.assertEqual(status, 0)
            self.assertEqual(
                [row["study_uid"] for row in writer.call_args.args[1]],
                ["study-alpha-early", "study-alpha-late", "study-beta"],
            )


@unittest.skipUnless(
    importlib.util.find_spec("openpyxl"),
    "本测试环境未安装openpyxl；Linux requirements安装后执行",
)
class LinuxWorkbookTests(unittest.TestCase):
    def test_linux_writer_creates_expected_sheets_and_types(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "report.xlsx"
            linux_v3.write_study_workbook_linux(
                output,
                [
                    {
                        "patient_id": "04553303",
                        "patient_name": "王玉梅",
                        "exam_date": "2025-04-17",
                        "modalities": "CT",
                        "has_3d": "无",
                        "display_type": "CT",
                        "study_uid": "1.2.3",
                        "series_count": 3,
                        "file_count": 315,
                        "confidence": 0.99,
                    }
                ],
                linux_v3._reports([], [], [], []),
                {"dicom": 315, "studies_current_run": 1, "studies": 1},
                rule_version="test",
            )

            workbook = load_workbook(output, data_only=False)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [
                        "影像检查明细",
                        "患者身份映射",
                        "患者身份待确认",
                        "来源目录身份审计",
                        "转存异常",
                        "字段说明",
                    ],
                )
                sheet = workbook["影像检查明细"]
                self.assertEqual(sheet["A5"].value, "04553303")
                self.assertEqual(sheet["G5"].value, "1.2.3")
                self.assertEqual(sheet["I5"].value, 315)
                self.assertEqual(sheet["J5"].value, 0.99)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
