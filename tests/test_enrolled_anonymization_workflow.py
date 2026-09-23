import argparse
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path


sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from enrolled_anonymization_common import anonymous_code, classify_period
from build_enrolled_study_selection import (
    assign_anonymous_codes,
    build_parser,
    create_anonymous_map_workbook,
    create_workbook,
    load_anonymous_map,
    prepare_selection,
    run as build_selection,
    validate_anonymous_map,
)


OPENPYXL_AVAILABLE = importlib.util.find_spec("openpyxl") is not None
PYDICOM_AVAILABLE = importlib.util.find_spec("pydicom") is not None


class PeriodRuleTests(unittest.TestCase):
    def test_surgical_periods_are_aggregated_per_study(self):
        self.assertEqual(classify_period(["术前"]).category, "术前")
        self.assertEqual(classify_period(["术中", "术中"]).category, "术中")
        self.assertEqual(classify_period(["术前", "术中"]).category, "术前与术中")

    def test_follow_up_boundaries_and_text(self):
        for value in (3, 6, 8.999, "术后6个月", "术后7.5个月"):
            with self.subTest(value=value):
                self.assertEqual(classify_period([value]).category, "6M")
        for value in (9, 12, 15, "术后9个月", "术后12个月"):
            with self.subTest(value=value):
                self.assertEqual(classify_period([value]).category, "12M")

    def test_after_nine_months_is_a_twelve_month_candidate(self):
        for value in (15.01, 18.4, "术后19个月"):
            with self.subTest(value=value):
                result = classify_period([value])
                self.assertEqual(result.category, "12M")
                self.assertTrue(result.valid)
                self.assertTrue(result.is_12m_candidate)

    def test_outside_and_mixed_periods_are_not_approved(self):
        for values in ([2.99], [-1], [""], ["术中", 6], [6, 12]):
            with self.subTest(values=values):
                result = classify_period(values)
                self.assertEqual(result.category, "时期待确认")
                self.assertFalse(result.valid)

    def test_anonymous_number_is_zero_padded(self):
        self.assertEqual(anonymous_code(1), "001")
        self.assertEqual(anonymous_code(17), "017")
        self.assertEqual(anonymous_code(206), "206")


class SelectionRuleTests(unittest.TestCase):
    def test_missing_id_patient_can_enter_only_with_explicit_manual_study_binding(self):
        sequence = [
            {
                "住院号": "",
                "患者": "ZFYI",
                "StudyInstanceUID": "1.2.840.10008.15.1",
                "SeriesUID": "1.2.840.10008.15.1.1",
                "SOPUID": "1.2.840.10008.15.1.1.1",
                "AcqusitionDate": "20250901",
                "时期": 6,
                "__row__": 2,
            }
        ]
        enrolled = [{"住院号": "", "患者": "ZFYI", "__row__": 2}]
        registry = [
            {
                "#": 17,
                "中心编号": 15,
                "中心名称": "测试中心",
                "受试者筛选号": "15-H001",
                "受试者编号": 1,
                "姓名": "ZFYI",
                "住院号/门诊号/放射号": "",
                "手术日期": "2025-03-01",
                "__row__": 18,
            }
        ]
        manual = [
            {
                "受试者筛选号": "15-H001",
                "患者": "ZFYI",
                "StudyInstanceUID": "1.2.840.10008.15.1",
                "是否确认同一人": "是",
                "确认依据": "数据来源方确认",
                "确认人": "测试确认人",
                "确认时间": "2026-09-14",
                "__row__": 2,
            }
        ]

        main, details, mapping, exceptions = prepare_selection(
            sequence,
            enrolled,
            registry,
            "15",
            "测试中心",
            manual_match_rows=manual,
        )

        self.assertEqual(exceptions, [])
        self.assertEqual(len(main), 1)
        self.assertEqual(main[0]["住院号"], "")
        self.assertEqual(main[0]["受试者筛选号"], "15-H001")
        self.assertEqual(main[0]["是否进入匿名化"], "是")
        self.assertEqual(main[0]["匹配状态"], "人工确认进入（缺少患者号）")
        self.assertIn("数据来源方确认", main[0]["说明"])
        self.assertEqual(mapping[0]["匹配依据"], "人工确认（缺少患者号；筛选号+StudyInstanceUID）")
        self.assertEqual(details[0]["人工确认表来源行"], 2)
        self.assertEqual(details[0]["人工确认依据"], "数据来源方确认")

    def test_missing_id_manual_row_not_marked_yes_does_not_enter(self):
        sequence = [
            {"住院号": "", "患者": "ZFYI", "StudyInstanceUID": "1.2.3", "时期": 6}
        ]
        enrolled = [{"住院号": "", "患者": "ZFYI", "__row__": 2}]
        registry = [
            {
                "#": 17,
                "中心编号": 15,
                "中心名称": "测试中心",
                "受试者筛选号": "15-H001",
                "受试者编号": 1,
                "姓名": "ZFYI",
                "住院号/门诊号/放射号": "",
            }
        ]
        manual = [
            {
                "受试者筛选号": "15-H001",
                "患者": "ZFYI",
                "StudyInstanceUID": "1.2.3",
                "是否确认同一人": "否",
                "确认依据": "尚未确认",
                "__row__": 2,
            }
        ]

        main, _details, mapping, exceptions = prepare_selection(
            sequence,
            enrolled,
            registry,
            "15",
            "测试中心",
            manual_match_rows=manual,
        )

        self.assertEqual(main, [])
        self.assertEqual(mapping, [])
        self.assertTrue(any(row["异常类型"] == "入组患者缺少住院号" for row in exceptions))

    def test_manual_same_study_cannot_be_assigned_to_two_screening_numbers(self):
        sequence = [{"住院号": "", "患者": "", "StudyInstanceUID": "1.2.3", "时期": 6}]
        enrolled = [
            {"住院号": "", "患者": "AAAA", "__row__": 2},
            {"住院号": "", "患者": "BBBB", "__row__": 3},
        ]
        registry = [
            {
                "#": 17,
                "中心编号": 15,
                "中心名称": "测试中心",
                "受试者筛选号": "15-H001",
                "受试者编号": 1,
                "姓名": "AAAA",
                "住院号/门诊号/放射号": "",
            },
            {
                "#": 18,
                "中心编号": 15,
                "中心名称": "测试中心",
                "受试者筛选号": "15-H002",
                "受试者编号": 2,
                "姓名": "BBBB",
                "住院号/门诊号/放射号": "",
            },
        ]
        manual = [
            {
                "受试者筛选号": "15-H001",
                "患者": "AAAA",
                "StudyInstanceUID": "1.2.3",
                "是否确认同一人": "是",
                "确认依据": "人工核实",
                "__row__": 2,
            },
            {
                "受试者筛选号": "15-H002",
                "患者": "BBBB",
                "StudyInstanceUID": "1.2.3",
                "是否确认同一人": "是",
                "确认依据": "人工核实",
                "__row__": 3,
            },
        ]

        main, _details, mapping, exceptions = prepare_selection(
            sequence,
            enrolled,
            registry,
            "15",
            "测试中心",
            manual_match_rows=manual,
        )

        self.assertEqual(main, [])
        self.assertEqual(mapping, [])
        self.assertEqual(
            sum(row["异常类型"] == "人工确认归属冲突" for row in exceptions),
            2,
        )

    def test_selection_matches_patient_and_deduplicates_study(self):
        sequence = [
            {
                "住院号": "1059564",
                "患者": "朱连绪",
                "StudyInstanceUID": "1.2.3",
                "SeriesUID": "1.2.3.1",
                "SOPUID": "1.2.3.1.1",
                "AcqusitionDate": "20250401",
                "时期": "术前",
                "__row__": 2,
            },
            {
                "住院号": "1059564",
                "患者": "朱连绪",
                "StudyInstanceUID": "1.2.3",
                "SeriesUID": "1.2.3.2",
                "SOPUID": "1.2.3.2.1",
                "AcqusitionDate": "20250401",
                "时期": "术中",
                "__row__": 3,
            },
            {
                "住院号": "999",
                "患者": "未入组",
                "StudyInstanceUID": "9.9.9",
                "SeriesUID": "9.9.9.1",
                "SOPUID": "9.9.9.1.1",
                "AcqusitionDate": "20250401",
                "时期": 6,
                "__row__": 4,
            },
        ]
        enrolled = [{"住院号": "1059564", "患者": "朱连绪", "__row__": 2}]
        registry = [
            {
                "#": 17,
                "中心编号": 1,
                "中心名称": "首都医科大学宣武医院",
                "受试者筛选号": "01-H017",
                "受试者编号": 17,
                "姓名": "朱连绪",
                "住院号/门诊号/放射号": 1059564,
                "手术日期": "2025-04-01",
                "__row__": 18,
            }
        ]

        main, details, mapping, exceptions = prepare_selection(
            sequence, enrolled, registry, "1", "首都医科大学宣武医院"
        )

        self.assertEqual(len(main), 1)
        self.assertEqual(main[0]["匿名编号"], "001")
        self.assertEqual(main[0]["标准访视期"], "术前与术中")
        self.assertEqual(main[0]["目标二级目录"], "001-术前与术中")
        self.assertEqual(main[0]["原明细行数"], 2)
        self.assertEqual(len(details), 2)
        self.assertEqual(len(mapping), 1)
        self.assertEqual(exceptions, [])

    def test_name_conflict_is_not_auto_matched(self):
        sequence = [{"住院号": "1", "患者": "甲", "StudyInstanceUID": "1.2.3"}]
        enrolled = [{"住院号": "1", "患者": "乙", "__row__": 2}]
        registry = []
        main, _details, _mapping, exceptions = prepare_selection(
            sequence, enrolled, registry, "1", "测试中心"
        )
        self.assertEqual(main, [])
        self.assertTrue(any(row["异常类型"] == "姓名不一致" for row in exceptions))

    def test_twelve_month_selects_closest_study_per_patient_with_late_fallback(self):
        sequence = [
            {"住院号": "1001", "患者": "患者甲", "StudyInstanceUID": "1.2.12", "时期": "术后12个月"},
            {"住院号": "1001", "患者": "患者甲", "StudyInstanceUID": "1.2.19", "时期": "术后19个月"},
            {"住院号": "1002", "患者": "患者乙", "StudyInstanceUID": "1.2.18", "时期": 18.4},
        ]
        enrolled = [
            {"住院号": "1001", "患者": "患者甲", "__row__": 2},
            {"住院号": "1002", "患者": "患者乙", "__row__": 3},
        ]
        registry = [
            {
                "#": 17,
                "中心编号": 1,
                "中心名称": "测试中心",
                "受试者筛选号": "01-H017",
                "姓名": "患者甲",
                "住院号/门诊号/放射号": "1001",
            },
            {
                "#": 18,
                "中心编号": 1,
                "中心名称": "测试中心",
                "受试者筛选号": "01-H018",
                "姓名": "患者乙",
                "住院号/门诊号/放射号": "1002",
            },
        ]

        main, _details, _mapping, exceptions = prepare_selection(
            sequence, enrolled, registry, "1", "测试中心"
        )
        by_uid = {row["StudyInstanceUID"]: row for row in main}

        self.assertEqual(by_uid["1.2.12"]["是否进入匿名化"], "是")
        self.assertEqual(by_uid["1.2.12"]["匹配状态"], "已选为最接近12M")
        self.assertEqual(by_uid["1.2.19"]["是否进入匿名化"], "否")
        self.assertEqual(by_uid["1.2.19"]["匹配状态"], "12M候选未选")
        self.assertEqual(by_uid["1.2.18"]["是否进入匿名化"], "是")
        self.assertEqual(by_uid["1.2.18"]["标准访视期"], "12M")
        self.assertIn("无更接近数据", by_uid["1.2.18"]["说明"])
        self.assertEqual(exceptions, [])

    def test_equal_distance_twelve_month_candidates_require_manual_choice(self):
        sequence = [
            {"住院号": "1001", "患者": "患者甲", "StudyInstanceUID": "1.2.10", "时期": 10},
            {"住院号": "1001", "患者": "患者甲", "StudyInstanceUID": "1.2.14", "时期": 14},
        ]
        enrolled = [{"住院号": "1001", "患者": "患者甲", "__row__": 2}]
        registry = [
            {
                "#": 17,
                "中心编号": 1,
                "中心名称": "测试中心",
                "受试者筛选号": "01-H017",
                "姓名": "患者甲",
                "住院号/门诊号/放射号": "1001",
            }
        ]

        main, _details, _mapping, exceptions = prepare_selection(
            sequence, enrolled, registry, "1", "测试中心"
        )

        self.assertEqual({row["是否进入匿名化"] for row in main}, {"否"})
        self.assertEqual({row["匹配状态"] for row in main}, {"12M并列待确认"})
        self.assertTrue(any(row["异常类型"] == "12M候选并列" for row in exceptions))

    def test_cross_center_codes_are_continuous_and_rerun_reuses_codes(self):
        shared_map = []
        patients_a = {
            "1002": {"住院号": "1002", "患者": "乙"},
            "1001": {"住院号": "1001", "患者": "甲"},
        }
        matches_a = {
            "1002": {
                "#": 127,
                "中心编号": 1,
                "中心名称": "宣武",
                "受试者筛选号": "01-H002",
                "受试者编号": 2,
            },
            "1001": {
                "#": 126,
                "中心编号": 1,
                "中心名称": "宣武",
                "受试者筛选号": "01-H001",
                "受试者编号": 1,
            },
        }
        assign_anonymous_codes(matches_a, patients_a, "1", "宣武", shared_map)
        self.assertEqual(matches_a["1001"]["匿名编号"], "001")
        self.assertEqual(matches_a["1002"]["匿名编号"], "002")

        patients_b = {"2001": {"住院号": "2001", "患者": "丙"}}
        matches_b = {
            "2001": {
                "#": 301,
                "中心编号": 2,
                "中心名称": "下一中心",
                "受试者筛选号": "02-H001",
                "受试者编号": 1,
            }
        }
        assign_anonymous_codes(matches_b, patients_b, "2", "下一中心", shared_map)
        self.assertEqual(matches_b["2001"]["匿名编号"], "003")

        rerun = {"1001": dict(matches_a["1001"]), "1002": dict(matches_a["1002"])}
        assign_anonymous_codes(rerun, patients_a, "1", "宣武", shared_map)
        self.assertEqual(rerun["1001"]["匿名编号"], "001")
        self.assertEqual(rerun["1002"]["匿名编号"], "002")
        self.assertEqual(len(shared_map), 3)

    def test_mapping_conflicts_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "同时对应"):
            validate_anonymous_map(
                [
                    {"匿名编号": "001", "中心编号": "1", "受试者筛选号": "01-H001"},
                    {"匿名编号": "001", "中心编号": "2", "受试者筛选号": "02-H001"},
                ]
            )


@unittest.skipUnless(OPENPYXL_AVAILABLE, "需要openpyxl")
class WorkbookLayoutTests(unittest.TestCase):
    def test_cli_run_reads_manual_workbook_and_writes_auditable_output(self):
        from openpyxl import Workbook, load_workbook

        def write_workbook(path, title, headers, rows):
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = title
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            workbook.save(path)
            workbook.close()

        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as temporary_directory:
            root = Path(temporary_directory)
            sequence = root / "sequence.xlsx"
            enrolled = root / "enrolled.xlsx"
            registry = root / "registry.xlsx"
            manual = root / "manual.xlsx"
            output = root / "selection.xlsx"
            anonymous_map = root / "anonymous-map.xlsx"
            write_workbook(
                sequence,
                "筛选序列",
                ["住院号", "患者", "StudyInstanceUID", "SeriesUID", "SOPUID", "AcqusitionDate", "时期"],
                [["", "ZFYI", "1.2.3", "1.2.3.1", "1.2.3.1.1", "20250901", 6]],
            )
            write_workbook(
                enrolled,
                "符合入组患者",
                ["住院号", "患者"],
                [["", "ZFYI"]],
            )
            write_workbook(
                registry,
                "分中心影像",
                ["#", "中心编号", "中心名称", "受试者筛选号", "受试者编号", "姓名", "住院号/门诊号/放射号", "手术日期"],
                [[17, 15, "测试中心", "15-H001", 1, "ZFYI", "", "2025-03-01"]],
            )
            write_workbook(
                manual,
                "缺号人工确认",
                ["受试者筛选号", "患者", "StudyInstanceUID", "是否确认同一人", "确认依据", "确认人", "确认时间"],
                [["15-H001", "ZFYI", "1.2.3", "是", "数据来源方确认", "确认人", "2026-09-14"]],
            )
            arguments = build_parser().parse_args(
                [
                    "--center", "15",
                    "--sequence-workbook", str(sequence),
                    "--enrolled-workbook", str(enrolled),
                    "--registry-workbook", str(registry),
                    "--manual-match-workbook", str(manual),
                    "--anonymous-map", str(anonymous_map),
                    "--output", str(output),
                ]
            )

            summary = build_selection(arguments)

            self.assertEqual(summary["patients"], 1)
            self.assertEqual(summary["studies"], 1)
            self.assertEqual(summary["manual_patients"], 1)
            self.assertEqual(summary["manual_studies"], 1)
            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                main_rows = list(workbook["入组Study清单"].iter_rows(values_only=True))
                headers = list(main_rows[0])
                values = dict(zip(headers, main_rows[1]))
                self.assertEqual(values["受试者筛选号"], "15-H001")
                self.assertEqual(values["住院号"], None)
                self.assertEqual(values["匹配状态"], "人工确认进入（缺少患者号）")
                self.assertEqual(workbook["匹配异常"].max_row, 1)
            finally:
                workbook.close()

    def test_control_workbook_is_white_with_visible_borders(self):
        main_row = {
            "中心编号": "1",
            "中心名称": "测试中心",
            "住院号": "H001",
            "患者": "测试患者",
            "受试者筛选号": "01-H001",
            "全局编号": "1",
            "匿名编号": "001",
            "StudyInstanceUID": "1.2.3",
            "标准访视期": "术前",
            "目标二级目录": "001-术前",
            "原明细行数": 1,
            "是否进入匿名化": "是",
            "匹配状态": "已匹配",
        }
        workbook = create_workbook([main_row], [], [], [])
        sheet = workbook["入组Study清单"]
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["A1"].fill.fgColor.rgb[-6:], "FFFFFF")
        self.assertEqual(sheet["A2"].fill.fgColor.rgb[-6:], "FFFFFF")
        self.assertEqual(sheet["A1"].border.left.style, "thin")
        self.assertEqual(sheet["A2"].border.bottom.style, "thin")
        self.assertEqual(len(sheet.data_validations.dataValidation), 1)
        workbook.close()

    def test_shared_mapping_workbook_round_trip(self):
        rows = [
            {
                "匿名编号": "001",
                "中心编号": "1",
                "中心名称": "宣武",
                "受试者筛选号": "01-H001",
                "原全局编号": "126",
                "住院号": "1001",
                "患者": "甲",
                "编号状态": "已分配",
            }
        ]
        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as temporary_directory:
            path = Path(temporary_directory) / "患者匿名编号映射.xlsx"
            workbook = create_anonymous_map_workbook(rows)
            workbook.save(path)
            workbook.close()
            loaded = load_anonymous_map(path)
        self.assertEqual(loaded, rows)


class LinuxEntrypointTests(unittest.TestCase):
    def test_selection_linux_entrypoint_reuses_shared_main(self):
        import build_enrolled_study_selection as shared
        import build_enrolled_study_selection_linux as linux_entrypoint

        self.assertIs(linux_entrypoint.main, shared.main)

    @unittest.skipUnless(PYDICOM_AVAILABLE, "需要pydicom")
    def test_anonymization_linux_entrypoint_reuses_shared_main(self):
        import anonymize_enrolled_studies as shared
        import anonymize_enrolled_studies_linux as linux_entrypoint

        self.assertIs(linux_entrypoint.main, shared.main)


@unittest.skipUnless(OPENPYXL_AVAILABLE and PYDICOM_AVAILABLE, "需要openpyxl和pydicom")
class EndToEndTests(unittest.TestCase):
    @staticmethod
    def _make_dicom(path: Path, study_uid: str, series_uid: str, sop_uid: str) -> bytes:
        import pydicom
        from pydicom.dataset import Dataset, FileDataset, FileMetaDataset
        from pydicom.sequence import Sequence
        from pydicom.uid import ExplicitVRLittleEndian, SecondaryCaptureImageStorage

        file_meta = FileMetaDataset()
        file_meta.MediaStorageSOPClassUID = SecondaryCaptureImageStorage
        file_meta.MediaStorageSOPInstanceUID = sop_uid
        file_meta.TransferSyntaxUID = ExplicitVRLittleEndian
        dataset = FileDataset(str(path), {}, file_meta=file_meta, preamble=b"\0" * 128)
        dataset.is_little_endian = True
        dataset.is_implicit_VR = False
        dataset.SOPClassUID = SecondaryCaptureImageStorage
        dataset.SOPInstanceUID = sop_uid
        dataset.SpecificCharacterSet = "ISO_IR 192"
        dataset.StudyInstanceUID = study_uid
        dataset.SeriesInstanceUID = series_uid
        dataset.PatientName = "测试患者"
        dataset.PatientID = "PID-KEEP-001"
        dataset.PatientBirthDate = "19800102"
        dataset.InstitutionName = "四川省人民医院"
        dataset.StudyDate = "20250901"
        dataset.SeriesDate = "20250902"
        dataset.AcquisitionDate = "20250903"
        dataset.StudyTime = "101112.123"
        dataset.StationName = "四川省人民医院DSA"
        nested = Dataset()
        nested.PatientName = "嵌套姓名"
        nested.PatientBirthDate = "19700101"
        nested.InstitutionName = "四川省人民医院"
        dataset.ReferencedStudySequence = Sequence([nested])
        dataset.Rows = 1
        dataset.Columns = 4
        dataset.SamplesPerPixel = 1
        dataset.PhotometricInterpretation = "MONOCHROME2"
        dataset.BitsAllocated = 8
        dataset.BitsStored = 8
        dataset.HighBit = 7
        dataset.PixelRepresentation = 0
        dataset.PixelData = b"\x01\x02\x03\x04"
        path.parent.mkdir(parents=True, exist_ok=True)
        pydicom.dcmwrite(path, dataset, enforce_file_format=True)
        return dataset.PixelData

    @staticmethod
    def _make_list(path: Path, study_uid: str) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "入组Study清单"
        sheet.append(
            [
                "中心编号",
                "中心名称",
                "匿名编号",
                "StudyInstanceUID",
                "标准访视期",
                "目标二级目录",
                "是否进入匿名化",
            ]
        )
        sheet.append([23, "四川省人民医院", "023", study_uid, "6M", "023-6M", "是"])
        workbook.save(path)
        workbook.close()

    def test_anonymizer_rejects_multiple_selected_twelve_month_studies(self):
        from openpyxl import Workbook
        from anonymize_enrolled_studies import load_plans

        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as temporary_directory:
            path = Path(temporary_directory) / "study-list.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "入组Study清单"
            sheet.append(
                [
                    "中心编号",
                    "中心名称",
                    "匿名编号",
                    "StudyInstanceUID",
                    "标准访视期",
                    "目标二级目录",
                    "是否进入匿名化",
                ]
            )
            sheet.append([23, "四川省人民医院", "023", "1.2.3", "12M", "023-12M", "是"])
            sheet.append([23, "四川省人民医院", "023", "1.2.4", "12M", "023-12M", "是"])
            workbook.save(path)
            workbook.close()

            plans, rejected = load_plans(str(path), "入组Study清单")

        self.assertEqual(plans, [])
        self.assertTrue(any("选择了多个12M" in item["message"] for item in rejected))

    def test_preview_then_execute_preserves_required_fields(self):
        import pydicom
        from anonymize_enrolled_studies import run

        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as temporary_directory:
            root = Path(temporary_directory)
            source = root / "source" / "23_H005"
            output = root / "delivery"
            control = root / "control"
            study_uid = "1.2.826.0.1.3680043.8.498.1001"
            series_uid = "1.2.826.0.1.3680043.8.498.2001"
            sop_uid = "1.2.826.0.1.3680043.8.498.3001"
            source_file = source / study_uid / series_uid / f"{sop_uid}.dcm"
            pixel_data = self._make_dicom(source_file, study_uid, series_uid, sop_uid)
            study_list = root / "study-list.xlsx"
            self._make_list(study_list, study_uid)
            arguments = argparse.Namespace(
                source_center=str(root / "source"),
                study_list=str(study_list),
                output_root=str(output),
                control_dir=str(control),
                sheet="入组Study清单",
                center_alias=["四川省人民医院"],
                execute=False,
            )

            summary, _audits, audit_path = run(arguments)
            self.assertEqual(summary.get("计划写入"), 1)
            self.assertIsNone(audit_path)
            self.assertFalse(output.exists())
            self.assertFalse(control.exists())

            arguments.execute = True
            summary, _audits, audit_path = run(arguments)
            self.assertEqual(summary.get("已写入"), 1)
            self.assertTrue(audit_path and audit_path.is_file())
            target = output / "023" / "023-6M" / study_uid / series_uid / f"{sop_uid}.dcm"
            result = pydicom.dcmread(target)
            self.assertNotIn("PatientName", result)
            self.assertNotIn("PatientBirthDate", result)
            self.assertNotIn("InstitutionName", result)
            self.assertNotIn("StationName", result)
            self.assertEqual(result.PatientID, "PID-KEEP-001")
            self.assertEqual(result.StudyDate, "20250901")
            self.assertEqual(result.SeriesDate, "20250902")
            self.assertEqual(result.AcquisitionDate, "20250903")
            self.assertEqual(result.StudyTime, "101112.123")
            self.assertEqual(result.StudyInstanceUID, study_uid)
            self.assertEqual(result.SeriesInstanceUID, series_uid)
            self.assertEqual(result.SOPInstanceUID, sop_uid)
            self.assertEqual(result.PixelData, pixel_data)
            self.assertNotIn("PatientName", result.ReferencedStudySequence[0])
            self.assertNotIn("PatientBirthDate", result.ReferencedStudySequence[0])
            self.assertNotIn("InstitutionName", result.ReferencedStudySequence[0])

    def test_legacy_flat_output_layout_is_rejected(self):
        from anonymize_enrolled_studies import run

        with tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent) as temporary_directory:
            root = Path(temporary_directory)
            source = root / "source"
            output = root / "delivery"
            control = root / "control"
            source.mkdir()
            (output / "023-6M").mkdir(parents=True)
            study_list = root / "study-list.xlsx"
            self._make_list(study_list, "1.2.826.0.1.3680043.8.498.1001")
            arguments = argparse.Namespace(
                source_center=str(source),
                study_list=str(study_list),
                output_root=str(output),
                control_dir=str(control),
                sheet="入组Study清单",
                center_alias=["四川省人民医院"],
                execute=False,
                workers=1,
            )

            with self.assertRaisesRegex(ValueError, "v1.3旧层级目录"):
                run(arguments)


if __name__ == "__main__":
    unittest.main()
