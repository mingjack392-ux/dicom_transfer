from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import shutil
import subprocess
import sys
import unittest
import uuid
from copy import copy
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch

import pydicom
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pydicom.dataset import FileDataset, FileMetaDataset
from pydicom.uid import ExplicitVRLittleEndian, SecondaryCaptureImageStorage, generate_uid


SRC_DIR = Path(__file__).resolve().parents[1] / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

import fill_dicom_equipment as equipment


HEADERS = [
    "匿名编号", "目标二级目录", "StudyInstanceUID", "SeriesUID", "SOPUID",
    "AcquisitionDate", "时期", "modality", "是否进入匿名化", "备注",
] + ["保留列%d" % column for column in range(11, 26)]
SHEET = "筛选序列明细"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def dicom_fixture(path: Path, study: str, series: str, sop: str,
                  manufacturer="Vendor A", model="Model A", **extra) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    meta = FileMetaDataset()
    meta.TransferSyntaxUID = ExplicitVRLittleEndian
    meta.MediaStorageSOPClassUID = SecondaryCaptureImageStorage
    meta.MediaStorageSOPInstanceUID = sop
    ds = FileDataset(str(path), {}, file_meta=meta, preamble=b"\0" * 128)
    ds.SOPClassUID = SecondaryCaptureImageStorage
    ds.SOPInstanceUID = sop
    ds.StudyInstanceUID = study
    ds.SeriesInstanceUID = series
    ds.Modality = "XA"
    ds.PatientID = "SYNTHETIC-001"
    ds.SpecificCharacterSet = "ISO_IR 192"
    if manufacturer is not None:
        ds.Manufacturer = manufacturer
    if model is not None:
        ds.ManufacturerModelName = model
    for keyword, value in extra.items():
        setattr(ds, keyword, value)
    ds.Rows = 1
    ds.Columns = 2
    ds.SamplesPerPixel = 1
    ds.PhotometricInterpretation = "MONOCHROME2"
    ds.BitsAllocated = 8
    ds.BitsStored = 8
    ds.HighBit = 7
    ds.PixelRepresentation = 0
    ds.PixelData = b"\x01\x02"
    ds.save_as(path, enforce_file_format=True)


def workbook_fixture(path: Path, rows: list[list], *, headers=None,
                     edge_styles=False) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(headers or HEADERS)
    for row in rows:
        ws.append(row + ["原值-%d" % n for n in range(len(row), 25)])
    for row in ws.iter_rows(min_row=1, max_row=len(rows) + 1, max_col=25):
        for cell in row:
            cell.font = Font(name="Arial", size=11, bold=cell.row == 1,
                             color="123456")
            cell.fill = PatternFill("solid", fgColor="DDEEFF")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="thin", color="445566"))
    ws["A2"].number_format = "000"
    ws.row_dimensions[1].height = 32
    ws.column_dimensions["H"].width = 17
    ws.column_dimensions["I"].width = 19
    ws.auto_filter.ref = "A1:Y%d" % (len(rows) + 1)
    ws.freeze_panes = "I2"
    if edge_styles:
        ws["XFC2"].fill = PatternFill("solid", fgColor="FFEEDD")
        ws["XFD2"].fill = PatternFill("solid", fgColor="FFEEDD")
    other = wb.create_sheet("说明")
    other["A1"] = "Synthetic workbook, no private patient data"
    other.merge_cells("A1:C1")
    other["A1"].font = Font(italic=True, color="AA1100")
    wb.save(path)
    wb.close()


class FillDicomEquipmentTests(unittest.TestCase):
    def setUp(self) -> None:
        output_root = Path(__file__).resolve().parents[1] / "outputs"
        output_root.mkdir(exist_ok=True)
        self.temp_parent = output_root.resolve()
        self.root = output_root / ("equipment_test_" + uuid.uuid4().hex)
        self.root.mkdir()
        self.source = self.root / "筛选影像"
        self.source.mkdir()
        self.xlsx = self.root / "输入工作簿.xlsx"
        self.output = self.root / "已补充设备.xlsx"
        self.control = self.root / "内部控制文件"
        self.study = generate_uid()
        self.series = generate_uid()
        self.sop = generate_uid()
        self.code = "001"
        self.visit = "001-术前"

    def tearDown(self) -> None:
        if self.root.resolve().parent != self.temp_parent or not self.root.name.startswith("equipment_test_"):
            raise RuntimeError("Refusing to clean up outside the test output directory")
        shutil.rmtree(self.root, ignore_errors=True)

    def path_for(self, series=None, sop=None) -> Path:
        return (self.source / self.code / self.visit / self.study /
                (series or self.series) / ((sop or self.sop) + ".dcm"))

    def add_dicom(self, *, series=None, sop=None, manufacturer="Vendor A",
                  model="Model A", study_header=None, series_header=None,
                  sop_header=None, **extra) -> Path:
        series = series or self.series
        sop = sop or self.sop
        path = self.path_for(series, sop)
        dicom_fixture(path, study_header or self.study, series_header or series,
                      sop_header or sop, manufacturer, model, **extra)
        return path

    def row(self, *, series=None, sop="NA", code=None, visit=None,
            included="是") -> list:
        return [self.code if code is None else code,
                self.visit if visit is None else visit,
                self.study, series or self.series, sop, "20240101", "术前",
                "3D-DSA", included, "保留的原始内容"]

    def run_fill(self, *, execute=False, **kwargs) -> dict:
        options = dict(sheet_name=SHEET, control_dir=self.control, workers=2,
                       progress_every=0, execute=execute)
        options.update(kwargs)
        with redirect_stdout(io.StringIO()):
            return equipment.run(self.source, self.xlsx, self.output, **options)

    def audit(self, result: dict) -> list[dict]:
        with Path(result["audit_path"]).open(encoding="utf-8-sig", newline="") as stream:
            return list(csv.DictReader(stream))

    def assert_row(self, result, status, manufacturer="", model="", index=0):
        record = self.audit(result)[index]
        self.assertEqual(status, record["status"])
        self.assertEqual(manufacturer, record["Manufacturer"])
        self.assertEqual(model, record["ManufacturerModelName"])
        self.assertEqual(str(index + 2), record["excel_row"])

    def test_preview_reads_header_only_and_preserves_inputs(self):
        self.add_dicom()
        self.add_dicom(sop=generate_uid())
        workbook_fixture(self.xlsx, [self.row()])
        before = {p: sha256(p) for p in [self.xlsx] + list(self.source.rglob("*.dcm"))}
        real_read = pydicom.dcmread
        with patch.object(equipment, "dcmread", wraps=real_read) as reader:
            result = self.run_fill()
        self.assertEqual(0, result["exit_code"])
        self.assertEqual({"ok": 1}, {k: v for k, v in result["counts"].items() if v})
        self.assert_row(result, "ok", "Vendor A", "Model A")
        self.assertFalse(self.output.exists())
        self.assertGreaterEqual(reader.call_count, 2)
        for call in reader.call_args_list:
            self.assertTrue(call.kwargs.get("stop_before_pixels"))
        self.assertEqual(before, {p: sha256(p) for p in before})
        self.assertTrue(Path(result["summary_path"]).is_file())
        with Path(result["summary_path"]).open(encoding="utf-8") as stream:
            self.assertIsInstance(json.load(stream), dict)

    def test_exact_sop_does_not_use_other_sop_equipment(self):
        self.add_dicom(manufacturer="Chosen", model="Chosen model")
        self.add_dicom(sop=generate_uid(), manufacturer="Unselected", model="Different")
        workbook_fixture(self.xlsx, [self.row(sop=self.sop)])
        self.assert_row(self.run_fill(), "ok", "Chosen", "Chosen model")

    def test_missing_exact_sop_does_not_fall_back_to_series(self):
        self.add_dicom(sop=generate_uid())
        workbook_fixture(self.xlsx, [self.row(sop=self.sop)])
        result = self.run_fill()
        self.assertEqual(1, result["exit_code"])
        self.assert_row(result, "source_missing")

    def test_distinct_series_in_one_study_keep_distinct_equipment(self):
        other_series = generate_uid()
        self.add_dicom(manufacturer="Vendor A", model="Model A")
        self.add_dicom(series=other_series, sop=generate_uid(), manufacturer="Vendor B", model="Model B")
        workbook_fixture(self.xlsx, [self.row(), self.row(series=other_series)])
        result = self.run_fill(execute=True)
        self.assert_row(result, "ok", "Vendor A", "Model A")
        self.assert_row(result, "ok", "Vendor B", "Model B", index=1)
        wb = load_workbook(self.output)
        self.assertEqual(("Vendor A", "Model A"), (wb[SHEET]["I2"].value, wb[SHEET]["J2"].value))
        self.assertEqual(("Vendor B", "Model B"), (wb[SHEET]["I3"].value, wb[SHEET]["J3"].value))
        wb.close()

    def test_series_partial_missing_fills_unique_nonempty_values(self):
        self.add_dicom(manufacturer="Vendor A", model=None)
        self.add_dicom(sop=generate_uid(), manufacturer="", model="Model A")
        workbook_fixture(self.xlsx, [self.row()])
        result = self.run_fill()
        self.assertEqual(1, result["exit_code"])
        self.assert_row(result, "partial_missing", "Vendor A", "Model A")

    def test_missing_field_is_blank_without_secondary_capture_fallback(self):
        self.add_dicom(manufacturer=None, model="Model A",
                       SecondaryCaptureDeviceManufacturer="Capture manufacturer",
                       SecondaryCaptureDeviceManufacturerModelName="Capture model")
        workbook_fixture(self.xlsx, [self.row()])
        result = self.run_fill()
        self.assertEqual(1, result["exit_code"])
        self.assert_row(result, "missing_tag", "", "Model A")

    def test_all_equipment_tags_missing_stays_blank(self):
        self.add_dicom(manufacturer=None, model=None)
        workbook_fixture(self.xlsx, [self.row()])
        self.assert_row(self.run_fill(), "missing_tag")

    def test_conflicting_manufacturers_or_models_blank_both_fields(self):
        for mismatch in ("manufacturer", "model"):
            with self.subTest(mismatch=mismatch):
                series = generate_uid()
                self.add_dicom(series=series, sop=generate_uid())
                options = {mismatch: "Conflicting"}
                self.add_dicom(series=series, sop=generate_uid(), **options)
                workbook_fixture(self.xlsx, [self.row(series=series)])
                result = self.run_fill()
                self.assertEqual(1, result["exit_code"])
                self.assert_row(result, "multiple_values")

    def test_uid_mismatch_in_any_series_file_blocks_row(self):
        self.add_dicom()
        self.add_dicom(sop=generate_uid(), study_header=generate_uid())
        workbook_fixture(self.xlsx, [self.row()])
        self.assert_row(self.run_fill(), "uid_mismatch")

    def test_header_sop_must_match_filename(self):
        self.add_dicom(sop_header=generate_uid())
        workbook_fixture(self.xlsx, [self.row(sop=self.sop)])
        self.assert_row(self.run_fill(), "uid_mismatch")

    def test_unreadable_series_file_blocks_otherwise_valid_values(self):
        self.add_dicom()
        self.path_for(sop=generate_uid()).write_bytes(b"not a DICOM file")
        workbook_fixture(self.xlsx, [self.row()])
        self.assert_row(self.run_fill(), "unreadable_dicom")

    def test_missing_series_and_ignored_rows(self):
        workbook_fixture(self.xlsx, [self.row(), self.row(series=generate_uid(), included="否")])
        result = self.run_fill(execute=True)
        self.assert_row(result, "source_missing")
        self.assert_row(result, "ignored", index=1)
        self.assertEqual(1, result["exit_code"])
        wb = load_workbook(self.output)
        self.assertIsNone(wb[SHEET]["I2"].value)
        self.assertIsNone(wb[SHEET]["J3"].value)
        wb.close()

    def test_all_whole_series_markers_and_header_aliases(self):
        self.add_dicom()
        rows = [self.row(sop=marker) for marker in (None, "", "NA", "/", "N/A", "NONE", "NULL")]
        headers = list(HEADERS)
        headers[3] = "SeriesInstanceUID"
        headers[4] = "SOPInstanceUID"
        headers[7] = "Modality"
        workbook_fixture(self.xlsx, rows, headers=headers)
        result = self.run_fill()
        self.assertEqual(0, result["exit_code"])
        for index in range(len(rows)):
            self.assert_row(result, "ok", "Vendor A", "Model A", index=index)

    def test_path_traversal_and_absolute_components_are_invalid_rules(self):
        self.add_dicom()
        rows = [self.row(code="../outside"), self.row(visit="/absolute"),
                self.row(code=".."), self.row(visit="C:\\outside")]
        workbook_fixture(self.xlsx, rows)
        result = self.run_fill()
        self.assertEqual(1, result["exit_code"])
        for index in range(len(rows)):
            self.assert_row(result, "invalid_rule", index=index)

    def test_numeric_and_scientific_uid_cells_are_not_guessed(self):
        self.add_dicom()
        numeric, scientific = self.row(), self.row()
        numeric[2] = 1.2840113
        scientific[3] = "1.2840113E+14"
        workbook_fixture(self.xlsx, [numeric, scientific])
        result = self.run_fill()
        self.assertEqual(1, result["exit_code"])
        for index in range(2):
            self.assert_row(result, "invalid_rule", index=index)

    def test_modifying_file_during_read_blocks_row(self):
        changed = self.add_dicom()
        workbook_fixture(self.xlsx, [self.row()])
        real_read = equipment.dcmread

        def read_then_change(path, **kwargs):
            ds = real_read(path, **kwargs)
            changed.write_bytes(changed.read_bytes() + b"changed-during-read")
            return ds

        with patch.object(equipment, "dcmread", side_effect=read_then_change):
            result = self.run_fill()
        self.assertEqual(1, result["exit_code"])
        self.assert_row(result, "source_changed")

    def test_symlink_outside_source_is_blocked(self):
        external = self.root / "outside.dcm"
        dicom_fixture(external, self.study, self.series, self.sop)
        link = self.path_for()
        link.parent.mkdir(parents=True)
        try:
            link.symlink_to(external)
        except (OSError, NotImplementedError) as error:
            self.skipTest("Symlinks not available in this environment: %s" % error)
        workbook_fixture(self.xlsx, [self.row()])
        result = self.run_fill()
        self.assertEqual(1, result["exit_code"])
        self.assert_row(result, "invalid_rule")

    def test_resolved_file_outside_root_is_rejected_before_header_read(self):
        selected = self.add_dicom()
        workbook_fixture(self.xlsx, [self.row()])
        real_resolve = Path.resolve

        def resolve_changed_path(path, *args, **kwargs):
            if path == selected:
                return self.root / "outside.dcm"
            return real_resolve(path, *args, **kwargs)

        with patch.object(Path, "resolve", new=resolve_changed_path):
            with patch.object(equipment, "dcmread") as reader:
                result = self.run_fill()
        self.assert_row(result, "invalid_rule")
        reader.assert_not_called()

    def test_output_and_control_paths_cannot_modify_source_tree(self):
        self.add_dicom()
        workbook_fixture(self.xlsx, [self.row()])
        for location in ("output", "control"):
            with self.subTest(location=location):
                kwargs = {"control_dir": self.source / "new-control"} if location == "control" else {}
                original_output = self.output
                if location == "output":
                    self.output = self.source / "output.xlsx"
                try:
                    with self.assertRaises((ValueError, OSError)):
                        self.run_fill(execute=True, **kwargs)
                finally:
                    self.output = original_output
        self.assertFalse((self.source / "output.xlsx").exists())
        self.assertFalse((self.source / "new-control").exists())

    def test_existing_output_and_source_workbook_cannot_be_overwritten(self):
        self.add_dicom()
        workbook_fixture(self.xlsx, [self.row()])
        self.output.write_bytes(b"pre-existing output")
        before = sha256(self.xlsx)
        with self.assertRaises((ValueError, OSError)):
            self.run_fill(execute=True)
        self.assertEqual(b"pre-existing output", self.output.read_bytes())
        self.output = self.xlsx
        with self.assertRaises((ValueError, OSError)):
            self.run_fill(execute=True)
        self.assertEqual(before, sha256(self.xlsx))

    def test_formula_workbook_is_refused_before_writing(self):
        self.add_dicom()
        workbook_fixture(self.xlsx, [self.row()])
        wb = load_workbook(self.xlsx)
        wb[SHEET]["Y2"] = "=SUM(1,2)"
        wb.save(self.xlsx)
        wb.close()
        before = sha256(self.xlsx)
        with self.assertRaises(ValueError):
            self.run_fill(execute=True)
        self.assertFalse(self.output.exists())
        self.assertEqual(before, sha256(self.xlsx))

    def test_nonempty_right_edge_cell_is_not_dropped(self):
        self.add_dicom()
        workbook_fixture(self.xlsx, [self.row()])
        wb = load_workbook(self.xlsx)
        wb[SHEET]["XFD2"] = "must not be discarded"
        wb.save(self.xlsx)
        wb.close()
        with self.assertRaises(ValueError):
            self.run_fill(execute=True)
        self.assertFalse(self.output.exists())

    def test_formula_like_dicom_value_remains_text(self):
        manufacturer = "=1+1"
        self.add_dicom(manufacturer=manufacturer)
        workbook_fixture(self.xlsx, [self.row()])
        result = self.run_fill(execute=True)
        self.assertEqual(0, result["exit_code"])
        wb = load_workbook(self.output, data_only=False)
        self.assertEqual(manufacturer, wb[SHEET]["I2"].value)
        self.assertEqual("s", wb[SHEET]["I2"].data_type)
        wb.close()
        self.assertEqual("'" + manufacturer, self.audit(result)[0]["Manufacturer"])

    def test_writer_inserts_after_modality_and_preserves_original_cells_and_styles(self):
        self.add_dicom(manufacturer="厂商 & <设备>", model='型号 "A"')
        workbook_fixture(self.xlsx, [self.row()], edge_styles=True)
        original_hash = sha256(self.xlsx)
        before = load_workbook(self.xlsx)
        result = self.run_fill(execute=True)
        self.assertEqual(0, result["exit_code"])
        self.assertEqual(original_hash, sha256(self.xlsx))
        after = load_workbook(self.output)
        old_ws, new_ws = before[SHEET], after[SHEET]
        self.assertEqual("Manufacturer", new_ws["I1"].value)
        self.assertEqual("ManufacturerModelName", new_ws["J1"].value)
        self.assertEqual("厂商 & <设备>", new_ws["I2"].value)
        self.assertEqual('型号 "A"', new_ws["J2"].value)
        for row in old_ws.iter_rows(min_row=1, max_row=2, max_col=25):
            for cell in row:
                column = cell.column + (2 if cell.column > 8 else 0)
                shifted = new_ws.cell(cell.row, column)
                self.assertEqual(cell.value, shifted.value, cell.coordinate)
                self.assertEqual(copy(cell.font), copy(shifted.font), cell.coordinate)
                self.assertEqual(copy(cell.fill), copy(shifted.fill), cell.coordinate)
                self.assertEqual(copy(cell.border), copy(shifted.border), cell.coordinate)
                self.assertEqual(copy(cell.alignment), copy(shifted.alignment), cell.coordinate)
                self.assertEqual(cell.number_format, shifted.number_format, cell.coordinate)
        self.assertEqual(32, new_ws.row_dimensions[1].height)
        self.assertEqual("A1:AA2", new_ws.auto_filter.ref)
        self.assertEqual("K2", new_ws.freeze_panes)
        self.assertEqual(19, new_ws.column_dimensions["K"].width)
        self.assertEqual(before["说明"]["A1"].value, after["说明"]["A1"].value)
        self.assertEqual(str(before["说明"].merged_cells), str(after["说明"].merged_cells))
        self.assertLessEqual(new_ws.max_column, 16384)
        before.close()
        after.close()

    def test_cli_returns_two_for_missing_workbook(self):
        command = [sys.executable, str(SRC_DIR / "fill_dicom_equipment.py"),
                   str(self.source), str(self.xlsx), str(self.output),
                   "--sheet", SHEET, "--control-dir", str(self.control)]
        completed = subprocess.run(command, capture_output=True, text=True,
                                   encoding="utf-8", errors="replace", env=dict(os.environ))
        self.assertEqual(2, completed.returncode, completed.stdout + completed.stderr)
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
