#!/usr/bin/env python3
"""Linux entrypoint for the center image-web report.

Business rules are imported from ``analyze_transferred_dicom_by_center.py``.
Only Excel reading/writing is replaced with the public openpyxl implementation,
so a normal Linux server does not require Node.js or the desktop artifact runtime.
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Optional, Sequence

import analyze_transferred_dicom_by_center as core
from build_center_image_workbook_linux import (
    require_openpyxl,
    write_center_workbook_linux,
)


def read_center_workbook_linux(workbook_path: Path) -> list[dict[str, Any]]:
    require_openpyxl()
    if workbook_path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Linux版分中心表目前仅支持.xlsx或.xlsm文件")
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        return [
            {
                "name": worksheet.title,
                "values": [list(row) for row in worksheet.iter_rows(values_only=True)],
            }
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def write_workbook_linux(
    output_path: Path,
    center_results: Sequence[core.CenterResult],
    center_table_exceptions: Sequence[dict[str, Any]],
    transferred_root: Path,
    center_workbook: Path,
    preview_dir: Optional[Path],
) -> None:
    if preview_dir is not None:
        print(
            "提示: Linux openpyxl版不生成PNG预览，--preview-dir已忽略；"
            "请用LibreOffice或Excel打开结果核对。",
            file=sys.stderr,
        )
    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "rule_version": core.RULE_VERSION,
        "transferred_root": str(transferred_root),
        "center_workbook": str(center_workbook),
        "headers": core.WEB_HEADERS,
        "centers": [
            {
                "sheet_name": result.sheet_name,
                "source_center_name": result.source_center_name,
                "reference_center_name": result.reference_center_name,
                "source_path": result.source_path,
                "rows": result.rows,
                "counters": result.counters,
            }
            for result in center_results
        ],
        "exceptions": [
            *center_table_exceptions,
            *(exception for result in center_results for exception in result.exceptions),
        ],
    }
    write_center_workbook_linux(output_path, payload)


def main(argv: Optional[Sequence[str]] = None) -> int:
    core.read_center_workbook = read_center_workbook_linux
    core.write_workbook = write_workbook_linux
    return core.main(argv)


if __name__ == "__main__":
    raise SystemExit(main())
