#!/usr/bin/env python3
"""按Excel筛选表从已转存DICOM中提取术前3D影像和术后2D DSA。

规则：

* ``时期=术前``：使用 ``SeriesUID``，复制命中Series的全部DICOM实例；
* ``时期=术中``：使用 ``SOPUID``，只复制命中的SOP实例；
* 其他时期不参与筛选；
* Excel中存在、但源目录中不存在的患者按业务约定直接忽略。

脚本默认只预检并生成CSV审计。只有显式传入 ``--execute`` 才复制DICOM。
源目录、目标目录必须彼此独立，脚本不会修改源DICOM。
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import threading
import unicodedata
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence

from openpyxl import load_workbook

import dicom_transfer_by_screening as transfer_core


RULE_VERSION = "2026.09.08-selection-filter-v1.0"
PREOP_PERIOD = "术前"
INTRAOP_PERIOD = "术中"
PREOP_CATEGORY = "术前3D影像"
INTRAOP_CATEGORY = "术后2DDSA"
UID_PATTERN = re.compile(r"[0-9]+(?:\.[0-9]+)*")
FOLDER_SEPARATOR_PATTERN = re.compile(r"[\s_]+")
_COPY_LOCKS = tuple(threading.Lock() for _ in range(257))


@dataclass(frozen=True)
class SelectionRule:
    row_number: int
    hospital_number: str
    patient_name: str
    period: str
    selector_field: str
    selector_uid: str
    category: str

    @property
    def patient_key(self) -> str:
        return patient_match_key(self.hospital_number, self.patient_name)


@dataclass(frozen=True)
class CopyTask:
    rule: SelectionRule
    patient_folder_name: str
    record: transfer_core.DicomRecord
    target: Path
    rule_match_count: int


@dataclass
class RunSummary:
    sheet_name: str
    selection_rows: int = 0
    selection_patients: int = 0
    source_patient_directories: int = 0
    matched_patients: int = 0
    ignored_missing_source_patients: int = 0
    ignored_source_directories: int = 0
    matched_rules: int = 0
    uid_not_found: int = 0
    invalid_rules: int = 0
    planned_files: int = 0
    copied: int = 0
    duplicate_same: int = 0
    conflicts: int = 0
    errors: int = 0
    detail_audit_path: str = ""
    patient_summary_path: str = ""

    @property
    def needs_attention(self) -> bool:
        return bool(
            self.uid_not_found
            or self.invalid_rules
            or self.ignored_source_directories
            or self.conflicts
            or self.errors
        )


DETAIL_HEADERS = [
    "Excel行号",
    "住院号",
    "患者",
    "原患者文件夹",
    "时期",
    "筛选字段",
    "筛选UID",
    "源文件",
    "目标文件",
    "StudyInstanceUID",
    "SeriesInstanceUID",
    "SOPInstanceUID",
    "文件大小",
    "SHA256",
    "状态",
    "说明",
]


PATIENT_HEADERS = [
    "住院号",
    "患者",
    "源患者目录",
    "处理状态",
    "术前规则数",
    "术中规则数",
    "命中规则数",
    "未命中规则数",
    "计划或处理文件数",
    "成功复制",
    "内容重复",
    "UID冲突",
    "错误",
    "说明",
]


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_text(value: str) -> str:
    return unicodedata.normalize("NFKC", value or "").strip()


def compact_match_text(value: str) -> str:
    return FOLDER_SEPARATOR_PATTERN.sub("", normalized_text(value)).casefold()


def patient_match_key(hospital_number: str, patient_name: str) -> str:
    return compact_match_text(hospital_number) + compact_match_text(patient_name)


def folder_match_key(folder_name: str) -> str:
    return compact_match_text(folder_name)


def uid_problem(uid: str) -> str:
    if not uid:
        return "UID为空"
    if not UID_PATTERN.fullmatch(uid):
        return "UID不是数字点号格式"
    if uid.startswith(".") or uid.endswith("."):
        return "UID以点开头或结尾"
    if len(uid) > 240:
        return "UID超过安全路径长度240"
    return ""


def uid_warning(uid: str) -> str:
    return f"UID长度{len(uid)}超过DICOM标准上限64" if len(uid) > 64 else ""


def choose_sheet(workbook: Any, requested_name: Optional[str]) -> Any:
    if requested_name:
        if requested_name not in workbook.sheetnames:
            raise ValueError(
                f"筛选表不存在工作表: {requested_name}；可用工作表: "
                + ", ".join(workbook.sheetnames)
            )
        return workbook[requested_name]
    for worksheet in workbook.worksheets:
        if worksheet.max_row >= 2 and worksheet.max_column >= 1:
            return worksheet
    raise ValueError("筛选表中没有可读取的数据工作表")


def _resolve_header(headers: dict[str, int], *aliases: str) -> int:
    for alias in aliases:
        if alias in headers:
            return headers[alias]
    raise ValueError("筛选表缺少字段: " + "/".join(aliases))


def read_selection_rules(
    workbook_path: Path, sheet_name: Optional[str] = None
) -> tuple[str, list[SelectionRule]]:
    if not workbook_path.is_file():
        raise ValueError(f"筛选表不存在: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = choose_sheet(workbook, sheet_name)
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise ValueError("筛选表为空") from exc
        headers: dict[str, int] = {}
        for index, value in enumerate(header_row):
            header = normalized_text(cell_text(value))
            if header and header not in headers:
                headers[header] = index

        hospital_index = _resolve_header(headers, "住院号")
        patient_index = _resolve_header(headers, "患者", "患者姓名")
        period_index = _resolve_header(headers, "时期")
        series_index = _resolve_header(headers, "SeriesUID", "SeriesInstanceUID")
        sop_index = _resolve_header(headers, "SOPUID", "SOPInstanceUID")

        rules: list[SelectionRule] = []
        for row_number, row in enumerate(rows, start=2):
            period = normalized_text(cell_text(row[period_index]))
            if period not in {PREOP_PERIOD, INTRAOP_PERIOD}:
                continue
            hospital_number = normalized_text(cell_text(row[hospital_index]))
            patient_name = normalized_text(cell_text(row[patient_index]))
            if period == PREOP_PERIOD:
                selector_field = "SeriesInstanceUID"
                selector_uid = normalized_text(cell_text(row[series_index]))
                category = PREOP_CATEGORY
            else:
                selector_field = "SOPInstanceUID"
                selector_uid = normalized_text(cell_text(row[sop_index]))
                category = INTRAOP_CATEGORY
            if selector_uid.upper() in {"NA", "N/A", "NONE", "NULL"}:
                selector_uid = ""
            rules.append(
                SelectionRule(
                    row_number=row_number,
                    hospital_number=hospital_number,
                    patient_name=patient_name,
                    period=period,
                    selector_field=selector_field,
                    selector_uid=selector_uid,
                    category=category,
                )
            )
        return worksheet.title, rules
    finally:
        workbook.close()


def group_rules_by_patient(
    rules: Sequence[SelectionRule],
) -> dict[str, list[SelectionRule]]:
    grouped: dict[str, list[SelectionRule]] = defaultdict(list)
    for rule in rules:
        grouped[rule.patient_key].append(rule)
    return dict(grouped)


def source_patient_directories(source_root: Path) -> list[Path]:
    return sorted(
        (
            path
            for path in source_root.iterdir()
            if path.is_dir() and not path.name.startswith((".", "_"))
        ),
        key=lambda path: normalized_text(path.name).casefold(),
    )


def build_folder_index(folders: Iterable[Path]) -> dict[str, list[Path]]:
    index: dict[str, list[Path]] = defaultdict(list)
    for folder in folders:
        index[folder_match_key(folder.name)].append(folder)
    return dict(index)


def selection_issue(rule: SelectionRule) -> str:
    missing: list[str] = []
    if not rule.hospital_number:
        missing.append("住院号为空")
    if not rule.patient_name:
        missing.append("患者为空")
    uid_issue = uid_problem(rule.selector_uid)
    if uid_issue:
        missing.append(uid_issue)
    return "；".join(missing)


def safe_target_for(
    destination_root: Path,
    patient_folder_name: str,
    category: str,
    record: transfer_core.DicomRecord,
) -> tuple[Optional[Path], str]:
    missing = [
        label
        for label, value in (
            ("StudyInstanceUID", record.study_uid),
            ("SeriesInstanceUID", record.series_uid),
            ("SOPInstanceUID", record.sop_uid),
        )
        if not value
    ]
    problems = [
        f"{label}: {problem}"
        for label, value in (
            ("StudyInstanceUID", record.study_uid),
            ("SeriesInstanceUID", record.series_uid),
            ("SOPInstanceUID", record.sop_uid),
        )
        if value and (problem := uid_problem(value))
    ]
    if missing or problems:
        messages: list[str] = []
        if missing:
            messages.append("缺少必要标签: " + ", ".join(missing))
        messages.extend(problems)
        return None, "；".join(messages)
    target = (
        destination_root
        / patient_folder_name
        / category
        / record.study_uid
        / record.series_uid
        / f"{record.sop_uid}.dcm"
    )
    warnings = [
        warning
        for warning in (
            uid_warning(record.study_uid),
            uid_warning(record.series_uid),
            uid_warning(record.sop_uid),
        )
        if warning
    ]
    return target, "；".join(warnings)


def base_detail_row(
    rule: SelectionRule, patient_folder_name: str = ""
) -> dict[str, Any]:
    return {
        "Excel行号": rule.row_number,
        "住院号": rule.hospital_number,
        "患者": rule.patient_name,
        "原患者文件夹": patient_folder_name,
        "时期": rule.period,
        "筛选字段": rule.selector_field,
        "筛选UID": rule.selector_uid,
        "源文件": "",
        "目标文件": "",
        "StudyInstanceUID": "",
        "SeriesInstanceUID": "",
        "SOPInstanceUID": "",
        "文件大小": "",
        "SHA256": "",
        "状态": "",
        "说明": "",
    }


def task_detail_row(
    task: CopyTask,
    *,
    status: str,
    destination_path: Path,
    file_size: int | str = "",
    sha256: str = "",
    message: str = "",
) -> dict[str, Any]:
    row = base_detail_row(task.rule, task.patient_folder_name)
    row.update(
        {
            "源文件": str(task.record.path),
            "目标文件": str(destination_path),
            "StudyInstanceUID": task.record.study_uid,
            "SeriesInstanceUID": task.record.series_uid,
            "SOPInstanceUID": task.record.sop_uid,
            "文件大小": file_size,
            "SHA256": sha256,
            "状态": status,
            "说明": message,
        }
    )
    return row


def copy_task(task: CopyTask, durable: bool) -> dict[str, Any]:
    try:
        lock = _COPY_LOCKS[hash(str(task.target).casefold()) % len(_COPY_LOCKS)]
        with lock:
            status, actual_target, size, digest = transfer_core.copy_atomic_with_hash(
                task.record.path, task.target, durable=durable
            )
        message_parts = list(task.record.warnings)
        if task.rule_match_count > 1 and task.rule.selector_field == "SOPInstanceUID":
            message_parts.append(
                f"同一SOPInstanceUID在源患者目录命中{task.rule_match_count}个文件，已保留处理结果"
            )
        return task_detail_row(
            task,
            status=status,
            destination_path=actual_target,
            file_size=size,
            sha256=digest,
            message="；".join(message_parts),
        )
    except Exception as exc:
        return task_detail_row(
            task,
            status="error",
            destination_path=task.target,
            message=str(exc),
        )


def unique_csv_path(directory: Path, prefix: str, timestamp: str) -> Path:
    candidate = directory / f"{prefix}_{timestamp}.csv"
    suffix = 1
    while candidate.exists():
        candidate = directory / f"{prefix}_{timestamp}_{suffix}.csv"
        suffix += 1
    return candidate


def write_csv(path: Path, headers: Sequence[str], rows: Sequence[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def validate_audit_path(source_root: Path, audit_dir: Path) -> None:
    source = source_root.resolve()
    audit = audit_dir.resolve()
    if audit == source or transfer_core.is_relative_to(audit, source):
        raise ValueError("审计目录不能位于源目录内部")


def _count_patient_statuses(rows: Sequence[dict[str, Any]]) -> Counter[str]:
    return Counter(str(row["状态"]) for row in rows)


def run_filter(
    source_root: Path,
    workbook_path: Path,
    destination_root: Path,
    *,
    sheet_name: Optional[str] = None,
    execute: bool = False,
    workers: int = 4,
    copy_workers: int = 2,
    audit_dir: Optional[Path] = None,
    progress_every: int = 500,
    durable: bool = False,
) -> RunSummary:
    source_root = Path(os.path.abspath(source_root))
    workbook_path = Path(os.path.abspath(workbook_path))
    destination_root = Path(os.path.abspath(destination_root))
    transfer_core.validate_paths(source_root, destination_root)
    audit_dir = Path(os.path.abspath(audit_dir or destination_root / "_selection_audit"))
    validate_audit_path(source_root, audit_dir)

    actual_sheet, rules = read_selection_rules(workbook_path, sheet_name)
    grouped_rules = group_rules_by_patient(rules)
    folders = source_patient_directories(source_root)
    folder_index = build_folder_index(folders)

    summary = RunSummary(
        sheet_name=actual_sheet,
        selection_rows=len(rules),
        selection_patients=len(grouped_rules),
        source_patient_directories=len(folders),
    )
    detail_rows: list[dict[str, Any]] = []
    patient_rows: list[dict[str, Any]] = []
    matched_folder_paths: set[Path] = set()

    for patient_key, patient_rules in grouped_rules.items():
        first_rule = patient_rules[0]
        preop_rules = sum(rule.period == PREOP_PERIOD for rule in patient_rules)
        intraop_rules = sum(rule.period == INTRAOP_PERIOD for rule in patient_rules)
        candidates = folder_index.get(patient_key, [])
        if not candidates:
            summary.ignored_missing_source_patients += 1
            for rule in patient_rules:
                row = base_detail_row(rule)
                row.update(
                    {
                        "状态": "ignored_missing_source_patient",
                        "说明": "筛选表有记录，但转存目录无对应患者目录；按业务约定忽略",
                    }
                )
                detail_rows.append(row)
            patient_rows.append(
                {
                    "住院号": first_rule.hospital_number,
                    "患者": first_rule.patient_name,
                    "源患者目录": "",
                    "处理状态": "ignored_missing_source_patient",
                    "术前规则数": preop_rules,
                    "术中规则数": intraop_rules,
                    "命中规则数": 0,
                    "未命中规则数": 0,
                    "计划或处理文件数": 0,
                    "成功复制": 0,
                    "内容重复": 0,
                    "UID冲突": 0,
                    "错误": 0,
                    "说明": "源目录中不存在，已忽略",
                }
            )
            continue
        if len(candidates) > 1:
            summary.errors += 1
            for rule in patient_rules:
                row = base_detail_row(rule)
                row.update(
                    {
                        "状态": "ambiguous_patient_folder",
                        "说明": "多个一级目录与住院号+患者匹配，未处理: "
                        + " | ".join(str(path) for path in candidates),
                    }
                )
                detail_rows.append(row)
            patient_rows.append(
                {
                    "住院号": first_rule.hospital_number,
                    "患者": first_rule.patient_name,
                    "源患者目录": " | ".join(str(path) for path in candidates),
                    "处理状态": "ambiguous_patient_folder",
                    "术前规则数": preop_rules,
                    "术中规则数": intraop_rules,
                    "命中规则数": 0,
                    "未命中规则数": 0,
                    "计划或处理文件数": 0,
                    "成功复制": 0,
                    "内容重复": 0,
                    "UID冲突": 0,
                    "错误": 1,
                    "说明": "患者目录匹配不唯一，需人工核对",
                }
            )
            continue

        patient_folder = candidates[0]
        matched_folder_paths.add(patient_folder.resolve())
        summary.matched_patients += 1
        try:
            records = transfer_core.scan_directory(
                patient_folder,
                destination_root,
                workers=max(1, workers),
                progress_every=max(0, progress_every),
            )
        except Exception as exc:
            summary.errors += 1
            for rule in patient_rules:
                row = base_detail_row(rule, patient_folder.name)
                row.update({"状态": "scan_error", "说明": str(exc)})
                detail_rows.append(row)
            patient_rows.append(
                {
                    "住院号": first_rule.hospital_number,
                    "患者": first_rule.patient_name,
                    "源患者目录": str(patient_folder),
                    "处理状态": "scan_error",
                    "术前规则数": preop_rules,
                    "术中规则数": intraop_rules,
                    "命中规则数": 0,
                    "未命中规则数": 0,
                    "计划或处理文件数": 0,
                    "成功复制": 0,
                    "内容重复": 0,
                    "UID冲突": 0,
                    "错误": 1,
                    "说明": str(exc),
                }
            )
            continue

        valid_records = [
            record for record in records if record.is_dicom and not record.unreadable
        ]
        series_index: dict[str, list[transfer_core.DicomRecord]] = defaultdict(list)
        sop_index: dict[str, list[transfer_core.DicomRecord]] = defaultdict(list)
        for record in valid_records:
            if record.series_uid:
                series_index[record.series_uid].append(record)
            if record.sop_uid:
                sop_index[record.sop_uid].append(record)

        patient_detail_start = len(detail_rows)
        tasks: list[CopyTask] = []
        matched_rule_count = 0
        missing_rule_count = 0
        invalid_rule_count = 0
        for rule in patient_rules:
            issue = selection_issue(rule)
            if issue:
                invalid_rule_count += 1
                summary.invalid_rules += 1
                row = base_detail_row(rule, patient_folder.name)
                row.update({"状态": "invalid_selection_rule", "说明": issue})
                detail_rows.append(row)
                continue
            index = series_index if rule.selector_field == "SeriesInstanceUID" else sop_index
            matches = index.get(rule.selector_uid, [])
            if not matches:
                missing_rule_count += 1
                summary.uid_not_found += 1
                row = base_detail_row(rule, patient_folder.name)
                row.update(
                    {
                        "状态": "uid_not_found",
                        "说明": f"在患者目录中未找到{rule.selector_field}",
                    }
                )
                detail_rows.append(row)
                continue
            matched_rule_count += 1
            summary.matched_rules += 1
            for record in matches:
                target, target_problem = safe_target_for(
                    destination_root, patient_folder.name, rule.category, record
                )
                if target is None:
                    summary.errors += 1
                    row = base_detail_row(rule, patient_folder.name)
                    row.update(
                        {
                            "源文件": str(record.path),
                            "StudyInstanceUID": record.study_uid,
                            "SeriesInstanceUID": record.series_uid,
                            "SOPInstanceUID": record.sop_uid,
                            "状态": "invalid_source_uids",
                            "说明": target_problem,
                        }
                    )
                    detail_rows.append(row)
                    continue
                tasks.append(
                    CopyTask(
                        rule=rule,
                        patient_folder_name=patient_folder.name,
                        record=record,
                        target=target,
                        rule_match_count=len(matches),
                    )
                )

        summary.planned_files += len(tasks)
        if execute and tasks:
            with ThreadPoolExecutor(max_workers=max(1, copy_workers)) as executor:
                processed_rows = list(
                    executor.map(lambda task: copy_task(task, durable), tasks)
                )
            detail_rows.extend(processed_rows)
        else:
            for task in tasks:
                messages = list(task.record.warnings)
                if (
                    task.rule_match_count > 1
                    and task.rule.selector_field == "SOPInstanceUID"
                ):
                    messages.append(
                        f"同一SOPInstanceUID在源患者目录命中{task.rule_match_count}个文件"
                    )
                detail_rows.append(
                    task_detail_row(
                        task,
                        status="planned",
                        destination_path=task.target,
                        file_size=task.record.path.stat().st_size,
                        message="；".join(messages),
                    )
                )

        patient_details = detail_rows[patient_detail_start:]
        status_counts = _count_patient_statuses(patient_details)
        summary.copied += status_counts["copied"]
        summary.duplicate_same += status_counts["duplicate_same"]
        summary.conflicts += status_counts["conflict"]
        summary.errors += status_counts["error"]
        patient_errors = (
            status_counts["error"]
            + status_counts["invalid_source_uids"]
            + invalid_rule_count
        )
        patient_attention = missing_rule_count + patient_errors + status_counts["conflict"]
        patient_rows.append(
            {
                "住院号": first_rule.hospital_number,
                "患者": first_rule.patient_name,
                "源患者目录": str(patient_folder),
                "处理状态": "needs_review" if patient_attention else ("completed" if execute else "preview_passed"),
                "术前规则数": preop_rules,
                "术中规则数": intraop_rules,
                "命中规则数": matched_rule_count,
                "未命中规则数": missing_rule_count,
                "计划或处理文件数": len(tasks),
                "成功复制": status_counts["copied"],
                "内容重复": status_counts["duplicate_same"],
                "UID冲突": status_counts["conflict"],
                "错误": patient_errors,
                "说明": f"扫描DICOM={len(valid_records)}，不可读DICOM={sum(record.unreadable for record in records)}",
            }
        )

    extra_folders = [
        folder for folder in folders if folder.resolve() not in matched_folder_paths
    ]
    matched_candidate_paths = {
        path.resolve()
        for patient_rules in grouped_rules.values()
        for path in folder_index.get(patient_rules[0].patient_key, [])
    }
    extra_folders = [
        folder for folder in extra_folders if folder.resolve() not in matched_candidate_paths
    ]
    summary.ignored_source_directories = len(extra_folders)
    for folder in extra_folders:
        patient_rows.append(
            {
                "住院号": "",
                "患者": "",
                "源患者目录": str(folder),
                "处理状态": "ignored_not_in_selection",
                "术前规则数": 0,
                "术中规则数": 0,
                "命中规则数": 0,
                "未命中规则数": 0,
                "计划或处理文件数": 0,
                "成功复制": 0,
                "内容重复": 0,
                "UID冲突": 0,
                "错误": 0,
                "说明": "转存目录存在，但筛选表无对应术前/术中规则；已忽略",
            }
        )

    timestamp = datetime.now().strftime("%Y%m%dT%H%M%S")
    detail_path = unique_csv_path(audit_dir, "筛选明细", timestamp)
    patient_path = unique_csv_path(audit_dir, "患者筛选汇总", timestamp)
    write_csv(detail_path, DETAIL_HEADERS, detail_rows)
    write_csv(patient_path, PATIENT_HEADERS, patient_rows)
    summary.detail_audit_path = str(detail_path)
    summary.patient_summary_path = str(patient_path)
    return summary


def print_summary(summary: RunSummary, execute: bool) -> None:
    print(f"规则版本: {RULE_VERSION}")
    print("运行模式: " + ("正式复制" if execute else "预检（不复制DICOM）"))
    print(f"筛选工作表: {summary.sheet_name}")
    print(f"筛选规则行: {summary.selection_rows}")
    print(f"筛选表患者: {summary.selection_patients}")
    print(f"源一级患者目录: {summary.source_patient_directories}")
    print(f"匹配并处理患者: {summary.matched_patients}")
    print(f"表内但源目录不存在（已忽略）: {summary.ignored_missing_source_patients}")
    print(f"源目录有但筛选表无规则（已忽略）: {summary.ignored_source_directories}")
    print(f"UID命中规则: {summary.matched_rules}")
    print(f"UID未命中规则: {summary.uid_not_found}")
    print(f"无效筛选规则: {summary.invalid_rules}")
    print(f"计划或处理DICOM: {summary.planned_files}")
    if execute:
        print(f"成功复制: {summary.copied}")
        print(f"内容重复: {summary.duplicate_same}")
        print(f"UID冲突保留: {summary.conflicts}")
    print(f"错误: {summary.errors}")
    print(f"筛选明细: {summary.detail_audit_path}")
    print(f"患者汇总: {summary.patient_summary_path}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="在Linux/Windows上按Excel的术前SeriesUID和术中SOPUID筛选已转存DICOM"
    )
    parser.add_argument("source_root", help="已转存中心目录，其一级目录为原患者文件夹")
    parser.add_argument("selection_xlsx", help="筛选序列.xlsx路径")
    parser.add_argument("destination_root", help="独立筛选结果根目录，不能位于源目录内部")
    parser.add_argument("--sheet", help="工作表名称；默认读取第一个非空工作表")
    parser.add_argument("--execute", action="store_true", help="正式复制；默认仅预检")
    parser.add_argument("--workers", type=int, default=4, help="每位患者的DICOM头读取线程数，默认4")
    parser.add_argument("--copy-workers", type=int, default=2, help="文件复制线程数，默认2")
    parser.add_argument("--audit-dir", help="审计CSV目录；默认目标目录/_selection_audit")
    parser.add_argument("--progress-every", type=int, default=500, help="每扫描多少文件显示进度，0表示关闭")
    parser.add_argument("--durable", action="store_true", help="逐文件fsync，更安全但速度较慢")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    args = build_parser().parse_args(argv)
    try:
        summary = run_filter(
            Path(args.source_root),
            Path(args.selection_xlsx),
            Path(args.destination_root),
            sheet_name=args.sheet,
            execute=args.execute,
            workers=args.workers,
            copy_workers=args.copy_workers,
            audit_dir=Path(args.audit_dir) if args.audit_dir else None,
            progress_every=args.progress_every,
            durable=args.durable,
        )
    except Exception as exc:
        print(f"筛选失败: {exc}", file=sys.stderr)
        return 2
    print_summary(summary, args.execute)
    return 1 if summary.needs_attention else 0


if __name__ == "__main__":
    raise SystemExit(main())
