#!/usr/bin/env python3
"""使用公开的 openpyxl 生成 DICOM V3 综合报表。

该模块专供普通 Linux 服务器使用，不依赖 Node.js、MJS 或 Codex 运行时。
"""

from __future__ import annotations

import csv
import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence


STUDY_HEADERS = [
    "PatientID",
    "PatientName",
    "检查日期",
    "基础影像类型",
    "是否3D_DSA断层",
    "影像类型汇总",
    "StudyInstanceUID",
    "序列数量",
    "文件数量",
    "判断置信度",
    "判断依据",
    "来源目录",
    "转存目录",
]

STUDY_KEYS = [
    "patient_id",
    "patient_name",
    "exam_date",
    "modalities",
    "has_3d",
    "display_type",
    "study_uid",
    "series_count",
    "file_count",
    "confidence",
    "evidence",
    "source_root",
    "destination_dir",
]

STUDY_WIDTHS = [16, 18, 13, 17, 19, 29, 42, 11, 11, 13, 48, 36, 42]

FIELD_DEFINITIONS = [
    ("一行的含义", "一次DICOM Study", "以StudyInstanceUID分组"),
    (
        "跨批次追加",
        "保留以前批次并追加本批Study",
        "状态保存在.dicom_v3_state/study_inventory.json；相同StudyInstanceUID不重复记行",
    ),
    (
        "明细排序",
        "同一自然患者的Study连续排列",
        "按身份组聚集；组内按检查日期、PatientID、StudyInstanceUID排序",
    ),
    (
        "基础影像类型",
        "同一Study内所有Series的模态去重合并",
        "Modality：XA、CT、MR；多个用顿号连接",
    ),
    (
        "是否3D_DSA断层",
        "该Study是否存在至少一个3D_DSA断层序列",
        "硬性条件：Modality=XA且SliceThickness有值；并且SeriesDescription命中3D/重建关键词",
    ),
    ("影像类型汇总", "3D_DSA断层标注在所属XA类型上", "示例：XA（含3D_DSA断层）、CT、OT"),
    ("检查日期", "用于以后与手术时间匹配", "优先StudyDate，缺失时使用AcquisitionDate"),
    (
        "判断置信度",
        "当前3D或基础模态判断可靠程度",
        "低于80%以黄色提示，需结合待确认记录复核",
    ),
    ("序列数量", "该Study内Series数量", "按SeriesInstanceUID去重"),
    ("文件数量", "该Study内影像实例数量", "按SOPInstanceUID去重"),
    (
        "患者身份映射",
        "跨批次 PatientID 与自然人身份组关系",
        "运行状态保存在.dicom_v3_state，Excel页用于查看与审计",
    ),
    (
        "患者身份待确认",
        "证据不足、人口学冲突或自动解决记录",
        "红色冲突、黄色待确认、绿色自动解决",
    ),
    (
        "来源目录身份审计",
        "同一来源目录出现多个PatientID时的分流结果",
        "按全局PatientID身份路由",
    ),
    (
        "转存异常",
        "UID缺失、内容冲突、身份缺失及复制错误",
        "不覆盖不同内容，异常文件保留",
    ),
    (
        "转存清单",
        "使用--manifest时生成的文件级清单",
        "大批量时会增加Excel生成时间和文件体积",
    ),
    ("下一阶段", "术中及术后6/12个月判断", "取得住院号和手术时间名单后再计算"),
]


def require_openpyxl() -> dict[str, Any]:
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import CellIsRule, FormulaRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise RuntimeError(
            "Linux版生成Excel需要openpyxl；请执行: "
            "python3 -m pip install --user openpyxl"
        ) from exc
    return {
        "Workbook": Workbook,
        "CellIsRule": CellIsRule,
        "FormulaRule": FormulaRule,
        "Alignment": Alignment,
        "Border": Border,
        "Font": Font,
        "PatternFill": PatternFill,
        "Side": Side,
        "get_column_letter": get_column_letter,
        "Table": Table,
        "TableStyleInfo": TableStyleInfo,
    }


def _safe_table_name(value: str) -> str:
    result = "".join(character for character in value if character.isalnum() or character == "_")
    if not result or result[0].isdigit():
        result = "T_" + result
    return result[:240]


def _parse_date(value: object) -> object:
    text = str(value or "").strip()
    if not text:
        return ""
    for pattern in ("%Y-%m-%d", "%Y%m%d"):
        try:
            return datetime.strptime(text[:10] if "-" in pattern else text[:8], pattern).date()
        except ValueError:
            continue
    return text


def _number(value: object, default: float = 0) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return default


def _configure_sheet(ws: Any) -> None:
    ws.sheet_view.showGridLines = False


def _header_style(ws: Any, row: int, start: int, end: int, api: dict[str, Any]) -> None:
    fill = api["PatternFill"]("solid", fgColor="4472C4")
    font = api["Font"](bold=True, color="FFFFFF")
    side = api["Side"](style="thin", color="B4C6E7")
    border = api["Border"](left=side, right=side, top=side, bottom=side)
    for column in range(start, end + 1):
        cell = ws.cell(row=row, column=column)
        cell.fill = fill
        cell.font = font
        cell.alignment = api["Alignment"](
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = border
    ws.row_dimensions[row].height = 32


def _add_table(
    ws: Any,
    ref: str,
    name: str,
    api: dict[str, Any],
) -> None:
    table = api["Table"](displayName=_safe_table_name(name), ref=ref)
    table.tableStyleInfo = api["TableStyleInfo"](
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def _set_widths(ws: Any, widths: Sequence[float], api: dict[str, Any]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[api["get_column_letter"](index)].width = width


def _style_body(ws: Any, first_row: int, last_row: int, last_column: int, api: dict[str, Any]) -> None:
    if last_row < first_row:
        return
    side = api["Side"](style="thin", color="D9E2F3")
    border = api["Border"](bottom=side)
    for row in ws.iter_rows(
        min_row=first_row,
        max_row=last_row,
        min_col=1,
        max_col=last_column,
    ):
        for cell in row:
            cell.alignment = api["Alignment"](vertical="top", wrap_text=True)
            cell.border = border


def _add_study_sheet(
    workbook: Any,
    studies: Sequence[dict[str, Any]],
    generated_at: str,
    rule_version: str,
    counters: dict[str, int],
    api: dict[str, Any],
) -> None:
    ws = workbook.active
    ws.title = "影像检查明细"
    _configure_sheet(ws)
    ws.merge_cells("A1:M1")
    ws["A1"] = "DICOM影像检查明细（每个Study一行）"
    ws["A1"].fill = api["PatternFill"]("solid", fgColor="1F4E78")
    ws["A1"].font = api["Font"](bold=True, color="FFFFFF", size=15)
    ws["A1"].alignment = api["Alignment"](horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:M2")
    ws["A2"] = (
        f"生成时间：{generated_at}    规则版本：{rule_version}    "
        f"本批DICOM：{counters.get('dicom', 0)}    "
        f"本批Study：{counters.get('studies_current_run', counters.get('studies', 0))}    "
        f"累计Study：{counters.get('studies', len(studies))}"
    )
    ws["A2"].fill = api["PatternFill"]("solid", fgColor="D9EAF7")
    ws["A2"].alignment = api["Alignment"](vertical="center")
    ws.row_dimensions[2].height = 22

    for column, header in enumerate(STUDY_HEADERS, start=1):
        ws.cell(row=4, column=column, value=header)
    _header_style(ws, 4, 1, len(STUDY_HEADERS), api)
    for column in range(4, 7):
        ws.cell(row=4, column=column).fill = api["PatternFill"](
            "solid", fgColor="FFD966"
        )
        ws.cell(row=4, column=column).font = api["Font"](bold=True, color="000000")

    for row_index, study in enumerate(studies, start=5):
        values: list[object] = []
        for key in STUDY_KEYS:
            value: object = study.get(key, "")
            if key == "exam_date":
                value = _parse_date(value)
            elif key in {"series_count", "file_count"}:
                value = int(_number(value))
            elif key == "confidence":
                value = _number(value)
            values.append(value)
        for column, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=column, value=value)
        ws.cell(row=row_index, column=1).number_format = "@"
        ws.cell(row=row_index, column=3).number_format = "yyyy-mm-dd"
        ws.cell(row=row_index, column=7).number_format = "@"
        ws.cell(row=row_index, column=8).number_format = "#,##0"
        ws.cell(row=row_index, column=9).number_format = "#,##0"
        ws.cell(row=row_index, column=10).number_format = "0%"

    if studies:
        last_row = 4 + len(studies)
        _style_body(ws, 5, last_row, len(STUDY_HEADERS), api)
        _add_table(ws, f"A4:M{last_row}", "StudySummaryTable", api)
        green = api["PatternFill"]("solid", fgColor="E2F0D9")
        yellow = api["PatternFill"]("solid", fgColor="FFF2CC")
        red = api["PatternFill"]("solid", fgColor="FCE4D6")
        ws.conditional_formatting.add(
            f"E5:E{last_row}",
            api["FormulaRule"](formula=['ISNUMBER(SEARCH("有",E5))'], fill=green),
        )
        ws.conditional_formatting.add(
            f"F5:F{last_row}",
            api["FormulaRule"](
                formula=['ISNUMBER(SEARCH("其他/未知",F5))'], fill=red
            ),
        )
        ws.conditional_formatting.add(
            f"J5:J{last_row}",
            api["CellIsRule"](operator="lessThan", formula=["0.8"], fill=yellow),
        )
        ws.auto_filter.ref = f"A4:M{last_row}"
    else:
        ws.merge_cells("A5:M6")
        ws["A5"] = "没有可汇总的Study，请查看转存异常页"
        ws["A5"].fill = api["PatternFill"]("solid", fgColor="FFF2CC")
        ws["A5"].alignment = api["Alignment"](horizontal="center", vertical="center")
    ws.freeze_panes = "A5"
    _set_widths(ws, STUDY_WIDTHS, api)


def _add_report_sheet(
    workbook: Any,
    sheet_name: str,
    columns: Sequence[str],
    rows: Sequence[dict[str, Any]],
    widths: Sequence[float],
    table_name: str,
    api: dict[str, Any],
) -> None:
    ws = workbook.create_sheet(sheet_name)
    _configure_sheet(ws)
    for column, header in enumerate(columns, start=1):
        ws.cell(row=1, column=column, value=header)
    _header_style(ws, 1, 1, len(columns), api)
    for row_index, row in enumerate(rows, start=2):
        for column, key in enumerate(columns, start=1):
            cell = ws.cell(row=row_index, column=column, value=row.get(key, ""))
            if "patient_id" in key.casefold() or key in {
                "StudyInstanceUID",
                "SeriesInstanceUID",
                "SOPInstanceUID",
            }:
                cell.number_format = "@"
    last_row = 1 + len(rows)
    _style_body(ws, 2, last_row, len(columns), api)
    if rows:
        _add_table(
            ws,
            f"A1:{api['get_column_letter'](len(columns))}{last_row}",
            table_name,
            api,
        )
        ws.auto_filter.ref = f"A1:{api['get_column_letter'](len(columns))}{last_row}"
    if sheet_name == "患者身份待确认" and rows:
        red = api["PatternFill"]("solid", fgColor="F4CCCC")
        yellow = api["PatternFill"]("solid", fgColor="FFF2CC")
        green = api["PatternFill"]("solid", fgColor="E2F0D9")
        ws.conditional_formatting.add(
            f"A2:A{last_row}",
            api["FormulaRule"](formula=['ISNUMBER(SEARCH("conflict",A2))'], fill=red),
        )
        ws.conditional_formatting.add(
            f"A2:A{last_row}",
            api["FormulaRule"](
                formula=['ISNUMBER(SEARCH("needs_review",A2))'], fill=yellow
            ),
        )
        ws.conditional_formatting.add(
            f"A2:A{last_row}",
            api["FormulaRule"](
                formula=['ISNUMBER(SEARCH("auto_resolved",A2))'], fill=green
            ),
        )
    ws.freeze_panes = "A2"
    _set_widths(ws, widths or [20] * len(columns), api)


def _add_manifest_sheet(
    workbook: Any,
    manifest_csv_path: Path,
    api: dict[str, Any],
) -> None:
    ws = workbook.create_sheet("转存清单")
    _configure_sheet(ws)
    with manifest_csv_path.open("r", newline="", encoding="utf-8-sig") as stream:
        reader = csv.reader(stream)
        try:
            columns = next(reader)
        except StopIteration:
            workbook.remove(ws)
            return
        ws.append(columns)
        row_count = 0
        for values in reader:
            ws.append(values)
            row_count += 1
    _header_style(ws, 1, 1, len(columns), api)
    # 百万级清单避免逐单元格设样式造成额外内存和数分钟固定开销。
    if 0 < row_count <= 50000:
        _style_body(ws, 2, row_count + 1, len(columns), api)
    if row_count:
        last_column = api["get_column_letter"](len(columns))
        _add_table(
            ws,
            f"A1:{last_column}{row_count + 1}",
            "TransferManifestTable",
            api,
        )
        ws.auto_filter.ref = f"A1:{last_column}{row_count + 1}"
    ws.freeze_panes = "A2"
    _set_widths(ws, [42, 42, 16, 38, 38, 38, 13, 67, 18, 32], api)


def _add_codebook(workbook: Any, api: dict[str, Any]) -> None:
    ws = workbook.create_sheet("字段说明")
    _configure_sheet(ws)
    ws.merge_cells("A1:C1")
    ws["A1"] = "字段与判断规则说明"
    ws["A1"].fill = api["PatternFill"]("solid", fgColor="1F4E78")
    ws["A1"].font = api["Font"](bold=True, color="FFFFFF", size=14)
    ws["A1"].alignment = api["Alignment"](horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    for column, value in enumerate(("字段", "说明", "来源/规则"), start=1):
        ws.cell(row=3, column=column, value=value)
    _header_style(ws, 3, 1, 3, api)
    for row_index, values in enumerate(FIELD_DEFINITIONS, start=4):
        for column, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=column, value=value)
    _style_body(ws, 4, 3 + len(FIELD_DEFINITIONS), 3, api)
    ws.freeze_panes = "A4"
    _set_widths(ws, [22, 38, 62], api)


def write_study_workbook_linux(
    output_path: Path,
    study_rows: Sequence[dict[str, Any]],
    reports: Sequence[dict[str, Any]],
    counters: dict[str, int],
    *,
    rule_version: str,
    generated_at: Optional[str] = None,
    manifest_csv_path: Optional[Path] = None,
) -> None:
    """原子生成 Linux 版 V3 综合 Excel。"""
    api = require_openpyxl()
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = api["Workbook"]()
    _add_study_sheet(
        workbook,
        study_rows,
        generated_at or datetime.now().isoformat(timespec="seconds"),
        rule_version,
        counters,
        api,
    )
    for report in reports:
        _add_report_sheet(
            workbook,
            str(report["sheet_name"]),
            list(report.get("columns", [])),
            list(report.get("rows", [])),
            list(report.get("widths", [])),
            str(report.get("table_name", "ReportTable")),
            api,
        )
    if manifest_csv_path and manifest_csv_path.is_file():
        _add_manifest_sheet(workbook, manifest_csv_path, api)
    _add_codebook(workbook, api)

    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=output_path.stem + "_",
        suffix=".xlsx",
        dir=str(output_path.parent),
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, output_path)
    finally:
        workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()
