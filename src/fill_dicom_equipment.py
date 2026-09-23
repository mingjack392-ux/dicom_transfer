#!/usr/bin/env python3
"""Read selected anonymized DICOM headers and fill equipment columns in a new XLSX.

Python 3.10+; works on Linux and Windows. Preview is the default. No DICOM is
written, copied, or decoded. --execute writes a new, never-overwritten workbook.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import sys
import threading
import uuid
import warnings
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

try:
    from openpyxl import load_workbook
    from pydicom import dcmread
except ImportError as exc:
    print(f"Missing dependency: {exc}. Run: python3 -m pip install -r requirements-dicom-equipment.txt",
          file=sys.stderr)
    raise SystemExit(2) from exc

from dicom_equipment_xlsx import inspect_workbook, write_equipment_workbook

VERSION = "1.0.0"
NO_SOP = {"", "NA", "N/A", "NONE", "NULL", "/"}
YES = {"是", "Y", "YES", "TRUE", "1"}
NO = {"", "否", "N", "NO", "FALSE", "0"}
HEADER_TAGS = ["StudyInstanceUID", "SeriesInstanceUID", "SOPInstanceUID",
               "Manufacturer", "ManufacturerModelName"]
ROW_FIELDS = ["excel_row", "anonymous_code", "target_folder", "StudyInstanceUID",
              "SeriesUID", "SOPUID", "modality", "selection_level", "status",
              "Manufacturer", "ManufacturerModelName", "files_found", "files_valid",
              "manufacturer_missing", "model_missing", "manufacturer_values",
              "model_values", "message"]
FILE_FIELDS = ["excel_row", "source_file", "status", "Manufacturer",
               "ManufacturerModelName", "actual_study_uid", "actual_series_uid",
               "actual_sop_uid", "message"]


def text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def digest(path: Path) -> str:
    result = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            result.update(chunk)
    return result.hexdigest()


def within(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def safe_component(value: str, uid: bool = False) -> bool:
    if not value or value in {".", ".."} or len(value.encode("utf-8")) > 240:
        return False
    if any(ord(c) < 32 or c in '/\\:*?"<>|' for c in value) or value.endswith((" ", ".")):
        return False
    if uid and (not re.fullmatch(r"[0-9.]+", value) or value.startswith(".")):
        return False
    return True


def _resolve_column(info: dict, *names: str) -> int:
    matches = []
    for header, cols in info.get("header_columns", {}).items():
        if header.casefold() in {name.casefold() for name in names}:
            matches.extend(cols)
    if not info.get("header_columns"):
        matches = [col for header, col in info["headers"].items()
                   if header.casefold() in {name.casefold() for name in names}]
    if len(matches) != 1:
        raise ValueError("字段缺失或重复: " + "/".join(names))
    return matches[0]


def read_rules(path: Path, sheet_name: Optional[str] = None) -> tuple[dict, list[dict]]:
    """Read only bounded data columns; no patient name/PatientID values are retained."""
    info = inspect_workbook(path, sheet_name)
    columns = {
        "anonymous_code": _resolve_column(info, "匿名编号"),
        "target_folder": _resolve_column(info, "目标二级目录"),
        "StudyInstanceUID": _resolve_column(info, "StudyInstanceUID"),
        "SeriesUID": _resolve_column(info, "SeriesUID", "SeriesInstanceUID"),
        "SOPUID": _resolve_column(info, "SOPUID", "SOPInstanceUID"),
        "include": _resolve_column(info, "是否进入匿名化"),
        "modality": _resolve_column(info, "modality"),
    }
    workbook = load_workbook(path, read_only=True, data_only=False)
    rules = []
    try:
        sheet = workbook[info["sheet_name"]]
        for row in sheet.iter_rows(min_row=2, max_row=info["max_data_row"],
                                   max_col=info["max_data_column"]):
            if not any(cell.value not in (None, "") for cell in row):
                continue
            rule = {key: text(row[col - 1].value) for key, col in columns.items()}
            code_cell = row[columns["anonymous_code"] - 1]
            # Respect an explicit 000-style display; do not guess zero-padding.
            if isinstance(code_cell.value, (int, float)) and not isinstance(code_cell.value, bool):
                if re.fullmatch(r"0+", code_cell.number_format or ""):
                    number = float(code_cell.value)
                    if number.is_integer() and number >= 0:
                        rule["anonymous_code"] = str(int(number)).zfill(len(code_cell.number_format))
            rule["excel_row"] = row[0].row
            for field in ("StudyInstanceUID", "SeriesUID", "SOPUID"):
                value = row[columns[field] - 1].value
                if value is not None and not isinstance(value, str):
                    rule["_invalid_reason"] = f"{field}在Excel中不是文本，无法保证UID精度"
            if rule["SOPUID"].upper() in NO_SOP:
                rule["SOPUID"] = ""
            rule["selection_level"] = "SOP" if rule["SOPUID"] else "Series"
            rules.append(rule)
    finally:
        workbook.close()
    if not rules:
        raise ValueError("工作表没有数据行")
    return info, rules


def snapshot(path: Path) -> tuple:
    stat = path.stat()
    return stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns, stat.st_ctime_ns


def read_header(path: Path, root: Path, rule: dict, expected_sop: str) -> dict:
    result = {"excel_row": rule["excel_row"], "source_file": str(path),
              "status": "ok", "Manufacturer": "", "ManufacturerModelName": "",
              "actual_study_uid": "", "actual_series_uid": "", "actual_sop_uid": "",
              "message": ""}
    try:
        resolved = path.resolve(strict=True)
        if not within(resolved, root) or not resolved.is_file():
            result.update(status="invalid_rule", message="文件解析路径越出影像根目录或不是普通文件")
            return result
        before = snapshot(path)
        # Legacy files without a preamble are accepted only if all three UIDs match.
        # A local catch block records parser warnings; data are never rewritten.
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            ds = dcmread(str(path), stop_before_pixels=True, specific_tags=HEADER_TAGS, force=True)
        result.update(
            Manufacturer=text(getattr(ds, "Manufacturer", "")),
            ManufacturerModelName=text(getattr(ds, "ManufacturerModelName", "")),
            actual_study_uid=text(getattr(ds, "StudyInstanceUID", "")),
            actual_series_uid=text(getattr(ds, "SeriesInstanceUID", "")),
            actual_sop_uid=text(getattr(ds, "SOPInstanceUID", "")),
        )
        result["_snapshot"] = before
        result["_resolved"] = str(resolved)
        if path.resolve(strict=True) != resolved or snapshot(path) != before:
            result.update(status="source_changed", message="读取过程中源文件发生变化，请稳定源目录后重跑")
        elif not any(result[k] for k in ("actual_study_uid", "actual_series_uid", "actual_sop_uid")):
            result.update(status="unreadable_dicom", message="未读到任何路由UID，不能认定为有效影像DICOM")
        elif (result["actual_study_uid"], result["actual_series_uid"], result["actual_sop_uid"]) != (
                rule["StudyInstanceUID"], rule["SeriesUID"], expected_sop):
            result.update(status="uid_mismatch", message="DICOM头Study/Series/SOP与Excel及文件路径不一致或缺失")
        elif any(getattr(ds, key, None) is not None and ds.data_element(key).VM > 1
                 for key in ("Manufacturer", "ManufacturerModelName")):
            result.update(status="multiple_values", message="单个DICOM设备字段包含多个值，拒绝作为单值回填")
        if caught:
            result["message"] = "; ".join(filter(None, [result["message"]] +
                                                list(dict.fromkeys(str(item.message) for item in caught))))
    except FileNotFoundError:
        result.update(status="source_missing", message="文件不存在或读取过程中消失")
    except Exception as exc:
        result.update(status="unreadable_dicom", message=f"{type(exc).__name__}: {exc}")
    return result


def inspect_rule(rule: dict, root: Path, progress: Any) -> tuple[dict, list[dict]]:
    result = {key: rule.get(key, "") for key in ROW_FIELDS}
    result.update(status="ok", Manufacturer="", ManufacturerModelName="",
                  files_found=0, files_valid=0, manufacturer_missing=0, model_missing=0,
                  manufacturer_values="[]", model_values="[]", message="")
    include = rule["include"].upper()
    if include in NO:
        result.update(status="ignored", message="该行未选择进入匿名化")
        return result, []
    if include not in YES:
        result.update(status="invalid_rule", message="是否进入匿名化的值无法识别")
        return result, []
    if rule.get("_invalid_reason"):
        result.update(status="invalid_rule", message=rule["_invalid_reason"])
        return result, []
    for key in ("anonymous_code", "target_folder", "StudyInstanceUID", "SeriesUID", "SOPUID"):
        if key == "SOPUID" and not rule[key]:
            continue
        if not safe_component(rule[key], uid=key.endswith("UID")):
            result.update(status="invalid_rule", message=f"{key}为空或不能安全作为路径")
            return result, []
    series = root / rule["anonymous_code"] / rule["target_folder"] / rule["StudyInstanceUID"] / rule["SeriesUID"]
    records = []
    try:
        if not within(series.resolve(), root):
            result.update(status="invalid_rule", message="Series解析路径越出影像根目录")
            return result, []
        if not series.is_dir():
            result.update(status="source_missing", message="未找到Series目录")
            return result, []
        before_dir = snapshot(series)
        if rule["SOPUID"]:
            files = [series / (rule["SOPUID"] + ".dcm")]
            if not files[0].exists():
                result.update(status="source_missing", message="未找到指定SOP文件；未回退为整Series")
                return result, []
        else:
            files = sorted((p for p in series.iterdir() if p.suffix.lower() == ".dcm"), key=lambda p: p.name)
            if not files:
                result.update(status="source_missing", message="Series目录下没有.dcm文件")
                return result, []
        result["files_found"] = len(files)
        for path in files:
            if not safe_component(path.stem, uid=True):
                record = {"excel_row": rule["excel_row"], "source_file": str(path),
                          "status": "invalid_rule", "message": "非标准UID文件名，无法验证SOP"}
            else:
                record = read_header(path, root, rule, path.stem)
            records.append(record)
            progress()
        if not rule["SOPUID"]:
            after_files = sorted(p.name for p in series.iterdir() if p.suffix.lower() == ".dcm")
            if snapshot(series) != before_dir or after_files != [p.name for p in files]:
                result.update(status="source_changed", message="扫描过程中Series目录内容发生变化")
        for record in records:
            if record["status"] == "ok":
                path = Path(record["source_file"])
                if snapshot(path) != record["_snapshot"] or str(path.resolve()) != record["_resolved"]:
                    record.update(status="source_changed", message="完成本行验证前源文件发生变化")
    except OSError as exc:
        result.update(status="unreadable_dicom", message=f"目录或文件访问失败: {exc}")
    valid = [r for r in records if r["status"] == "ok"]
    result["files_valid"] = len(valid)
    manufacturers = sorted({r["Manufacturer"] for r in valid if r["Manufacturer"]})
    models = sorted({r["ManufacturerModelName"] for r in valid if r["ManufacturerModelName"]})
    result["manufacturer_values"] = json.dumps(manufacturers, ensure_ascii=False)
    result["model_values"] = json.dumps(models, ensure_ascii=False)
    result["manufacturer_missing"] = sum(not r["Manufacturer"] for r in valid)
    result["model_missing"] = sum(not r["ManufacturerModelName"] for r in valid)
    failures = [r for r in records if r["status"] != "ok"]
    if result["status"] != "ok":
        return result, records
    if failures:
        result.update(status=failures[0]["status"], message="至少一个文件验证失败，整行留空；详见文件审计")
    elif len(manufacturers) > 1 or len(models) > 1:
        result.update(status="multiple_values", message="同一规则出现多个厂家或型号，两列均留空")
    else:
        result["Manufacturer"] = manufacturers[0] if manufacturers else ""
        result["ManufacturerModelName"] = models[0] if models else ""
        if not manufacturers or not models:
            result.update(status="missing_tag", message="至少一个字段全部缺失；另一字段如唯一则保留")
        elif result["manufacturer_missing"] or result["model_missing"]:
            result.update(status="partial_missing", message="部分文件字段缺失，回填唯一非空值")
    return result, records


def _csv_value(value: Any) -> Any:
    # CSV opened in Excel must not execute untrusted DICOM strings as formulas.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def write_csv(path: Path, fields: list, records: list) -> None:
    with path.open("x", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows({key: _csv_value(value) for key, value in row.items()} for row in records)


def run(source_root: Any, selection_xlsx: Any, output_xlsx: Any,
        sheet_name: Optional[str] = None, control_dir: Any = None, workers: int = 4,
        progress_every: int = 500, execute: bool = False) -> dict:
    if workers < 1 or progress_every < 0:
        raise ValueError("workers必须大于0；progress-every不能小于0")
    source = Path(source_root).resolve(strict=True)
    selection = Path(selection_xlsx).resolve(strict=True)
    output_raw = Path(output_xlsx)
    if output_raw.exists() or output_raw.is_symlink():
        raise FileExistsError(f"输出文件已存在，拒绝覆盖: {output_raw}")
    output = output_raw.resolve()
    control = Path(control_dir).resolve() if control_dir else output.parent / (output.stem + "_设备控制")
    if not source.is_dir() or not selection.is_file():
        raise ValueError("影像根目录或Excel路径无效")
    if selection.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
        raise ValueError("输入和输出必须是.xlsx工作簿")
    if output == selection or within(output, source):
        raise ValueError("输出Excel不能覆盖输入或位于DICOM源目录内")
    if within(control, source) or within(source, control):
        raise ValueError("审计目录必须与DICOM源目录相互独立")
    if output == control or within(selection, output) or within(control, output):
        raise ValueError("输出文件与输入或审计目录路径冲突")
    started = datetime.now().astimezone().isoformat()
    input_hash = digest(selection)
    info, rules = read_rules(selection, sheet_name)
    print(f"设备回填 v{VERSION} | {'执行' if execute else '预检'} | 工作表: {info['sheet_name']}", flush=True)
    print(f"读取Excel数据行: {len(rules)}；DICOM头验证线程: {workers}", flush=True)
    control.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f") + "_" + uuid.uuid4().hex[:8]
    audit_path = control / f"设备信息回填明细_{stamp}.csv"
    file_audit_path = control / f"设备信息文件审计_{stamp}.csv"
    summary_path = control / f"设备信息回填汇总_{stamp}.json"
    count = 0
    lock = threading.Lock()

    def progress() -> None:
        nonlocal count
        with lock:
            count += 1
            if progress_every and count % progress_every == 0:
                print(f"已处理DICOM Header: {count}", flush=True)

    with ThreadPoolExecutor(max_workers=workers) as pool:
        outcomes = list(pool.map(lambda rule: inspect_rule(rule, source, progress), rules))
    rows = [outcome[0] for outcome in outcomes]
    files = [record for outcome in outcomes for record in outcome[1]]
    counts = dict(Counter(row["status"] for row in rows))
    values = {row["excel_row"]: (row["Manufacturer"], row["ManufacturerModelName"]) for row in rows}
    exit_code = int(any(row["status"] not in {"ok", "ignored"} for row in rows))
    summary = dict(version=VERSION, started_at=started, sheet_name=info["sheet_name"],
                   source_root=str(source), selection_xlsx=str(selection), selection_sha256=input_hash,
                   mode="execute" if execute else "preview", workbook_rows=len(rows),
                   anonymous_codes=len({r["anonymous_code"] for r in rules if r["include"].upper() in YES}),
                   included_rows=sum(r["include"].upper() in YES for r in rules),
                   series_rules=sum(r["include"].upper() in YES and not r["SOPUID"] for r in rules),
                   sop_rules=sum(r["include"].upper() in YES and bool(r["SOPUID"]) for r in rules),
                   counts=counts, files_checked=len(files),
                   both_fields_filled=sum(bool(r["Manufacturer"] and r["ManufacturerModelName"]) for r in rows),
                   manufacturer_filled=sum(bool(r["Manufacturer"]) for r in rows),
                   model_filled=sum(bool(r["ManufacturerModelName"]) for r in rows),
                   requested_output=str(output), output_path="", audit_path=str(audit_path),
                   file_audit_path=str(file_audit_path), summary_path=str(summary_path), exit_code=exit_code)
    write_csv(audit_path, ROW_FIELDS, rows)
    write_csv(file_audit_path, FILE_FIELDS, files)
    try:
        if digest(selection) != input_hash:
            raise ValueError("扫描期间输入Excel发生变化，请使用已保存的固定版本重跑")
        if execute:
            output.parent.mkdir(parents=True, exist_ok=True)
            write_equipment_workbook(selection, output, info["sheet_name"], values)
            summary["output_path"] = str(output)
            summary["output_sha256"] = digest(output)
            summary["dropped_blank_edge_cells"] = info.get("dropped_blank_cells", 0)
    except Exception as exc:
        summary.update(exit_code=2, error=f"{type(exc).__name__}: {exc}")
        raise
    finally:
        summary["finished_at"] = datetime.now().astimezone().isoformat()
        with summary_path.open("x", encoding="utf-8") as stream:
            json.dump(summary, stream, ensure_ascii=False, indent=2)
            stream.write("\n")
    print("行级结果: " + json.dumps(counts, ensure_ascii=False), flush=True)
    print(f"厂家和型号均有值: {summary['both_fields_filled']}/{len(rows)}", flush=True)
    print(f"明细审计: {audit_path}\n文件审计: {file_audit_path}\n汇总: {summary_path}", flush=True)
    print(f"新工作簿: {output}" if execute else "预检完成；未生成Excel，正式回填增加 --execute。", flush=True)
    return summary


def main(argv: Optional[list[str]] = None) -> int:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="从筛选后DICOM读取厂家/型号，写入modality之后两列；默认预检")
    parser.add_argument("source_root", help="03.筛选数据/01.筛选影像根目录")
    parser.add_argument("selection_xlsx", help="待补充的原Excel，仅支持.xlsx")
    parser.add_argument("output_xlsx", help="新的结果Excel，必须尚不存在")
    parser.add_argument("--sheet", dest="sheet_name", help="工作表名称，默认第一个工作表")
    parser.add_argument("--control-dir", help="独立审计目录")
    parser.add_argument("--workers", type=int, default=4, help="DICOM头验证线程数，默认4")
    parser.add_argument("--progress-every", type=int, default=500, help="每处理多少文件显示进度，0关闭")
    parser.add_argument("--execute", action="store_true", help="写入新的结果Excel；不会写入或复制DICOM")
    args = parser.parse_args(argv)
    try:
        return run(**vars(args))["exit_code"]
    except Exception as exc:
        print(f"设备回填失败: {type(exc).__name__}: {exc}", file=sys.stderr, flush=True)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
