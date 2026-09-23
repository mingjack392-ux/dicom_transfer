#!/usr/bin/env python3
"""Build a per-center, reviewable Study whitelist for enrolled patients."""

from __future__ import annotations

import argparse
import math
import re
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from enrolled_anonymization_common import (
    anonymous_code,
    atomic_save_workbook,
    choose_center,
    classify_period,
    is_yes,
    normalized_identifier,
    normalized_name,
    normalized_screening,
    normalized_text,
    read_table,
    text,
)


SELECTION_RULE_VERSION = "2026.09.14-enrolled-selection-v1.5"
INTERNAL_PATIENT_KEY = "__matched_patient_key__"
INTERNAL_MANUAL_BASIS = "__manual_confirmation_basis__"
INTERNAL_MANUAL_ROW = "__manual_confirmation_row__"
INTERNAL_MANUAL_PERSON = "__manual_confirmation_person__"
INTERNAL_MANUAL_TIME = "__manual_confirmation_time__"


SEQUENCE_REQUIRED = (
    "住院号",
    "患者",
    "StudyInstanceUID",
    "SeriesUID",
    "SOPUID",
    "AcqusitionDate",
    "时期",
)
ENROLLED_REQUIRED = ("住院号", "患者")
MANUAL_MATCH_REQUIRED = (
    "受试者筛选号",
    "患者",
    "StudyInstanceUID",
    "是否确认同一人",
    "确认依据",
)
REGISTRY_REQUIRED = (
    "#",
    "中心编号",
    "中心名称",
    "受试者筛选号",
    "受试者编号",
    "姓名",
    "住院号/门诊号/放射号",
    "手术日期",
)

MAIN_HEADERS = (
    "中心编号",
    "中心名称",
    "住院号",
    "患者",
    "受试者筛选号",
    "全局编号",
    "匿名编号",
    "StudyInstanceUID",
    "AcqusitionDate",
    "原时期",
    "标准访视期",
    "目标二级目录",
    "原明细行数",
    "是否进入匿名化",
    "匹配状态",
    "说明",
)

ANONYMOUS_MAP_HEADERS = (
    "匿名编号",
    "中心编号",
    "中心名称",
    "受试者筛选号",
    "原全局编号",
    "住院号",
    "患者",
    "编号状态",
)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="根据入组患者表和筛选序列表生成按中心审核的Study白名单。"
    )
    parser.add_argument("--center", required=True, help="中心简称、全称或中心编号")
    parser.add_argument("--sequence-workbook", required=True, help="筛选序列表.xls/.xlsx")
    parser.add_argument("--enrolled-workbook", required=True, help="符合入组患者表.xls/.xlsx")
    parser.add_argument("--registry-workbook", required=True, help="分中心影像表.xls/.xlsx")
    parser.add_argument(
        "--manual-match-workbook",
        help=(
            "缺少患者号患者的人工确认表.xls/.xlsx；仅处理明确填写“是”的"
            "筛选号+StudyInstanceUID绑定"
        ),
    )
    parser.add_argument("--manual-match-sheet", help="人工确认表工作表名称")
    parser.add_argument(
        "--anonymous-map",
        help="跨中心共享的患者匿名编号映射.xlsx；默认放在输出表同目录",
    )
    parser.add_argument("--output", required=True, help="输出的入组Study筛选表.xlsx")
    parser.add_argument("--sequence-sheet")
    parser.add_argument("--enrolled-sheet")
    parser.add_argument("--registry-sheet")
    parser.add_argument("--overwrite", action="store_true", help="确认覆盖已有输出表")
    return parser


def _mapping_key(center_code: Any, screening: Any) -> str:
    return f"{normalized_identifier(center_code)}|{normalized_screening(screening)}"


def _natural_sort_key(value: Any) -> tuple[Any, ...]:
    parts = re.split(r"(\d+)", normalized_screening(value))
    return tuple(int(part) if part.isdigit() else part for part in parts)


def _registry_assignment_sort_key(item: tuple[str, dict[str, Any]]) -> tuple[Any, ...]:
    patient_key, registry = item
    subject_number = text(registry.get("受试者编号"))
    if re.fullmatch(r"\d+(?:\.0+)?", subject_number):
        subject_key: tuple[Any, ...] = (0, int(float(subject_number)))
    else:
        subject_key = (1, _natural_sort_key(registry.get("受试者筛选号")))
    return subject_key + (_natural_sort_key(registry.get("受试者筛选号")), patient_key)


def validate_anonymous_map(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    """Validate and normalize the shared cross-center anonymous-code ledger."""

    normalized_rows: list[dict[str, Any]] = []
    key_to_code: dict[str, str] = {}
    code_to_key: dict[str, str] = {}
    for index, source_row in enumerate(rows, start=2):
        screening = text(source_row.get("受试者筛选号"))
        center_code = text(source_row.get("中心编号"))
        key = _mapping_key(center_code, screening)
        if not normalized_identifier(center_code) or not normalized_screening(screening):
            raise ValueError(f"匿名编号映射第{source_row.get('__row__', index)}行缺少中心编号或筛选号")
        try:
            code = anonymous_code(source_row.get("匿名编号"))
        except ValueError as exc:
            raise ValueError(
                f"匿名编号映射第{source_row.get('__row__', index)}行编号无效: {exc}"
            ) from exc
        previous_code = key_to_code.get(key)
        if previous_code and previous_code != code:
            raise ValueError(f"匿名编号映射冲突: {key}同时对应{previous_code}和{code}")
        previous_key = code_to_key.get(code)
        if previous_key and previous_key != key:
            raise ValueError(f"匿名编号映射冲突: 编号{code}同时对应{previous_key}和{key}")
        if previous_code:
            raise ValueError(f"匿名编号映射重复: {key}与编号{code}出现多行")
        key_to_code[key] = code
        code_to_key[code] = key
        normalized_rows.append(
            {
                "匿名编号": code,
                "中心编号": center_code,
                "中心名称": text(source_row.get("中心名称")),
                "受试者筛选号": screening,
                "原全局编号": text(source_row.get("原全局编号") or source_row.get("全局编号")),
                "住院号": text(source_row.get("住院号")),
                "患者": text(source_row.get("患者")),
                "编号状态": text(source_row.get("编号状态")) or "已分配",
            }
        )
    normalized_rows.sort(key=lambda row: int(row["匿名编号"]))
    return normalized_rows


def load_anonymous_map(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    table = read_table(
        path,
        ("匿名编号", "中心编号", "受试者筛选号"),
        sheet_name="患者匿名编号映射",
    )
    return validate_anonymous_map(table.rows)


def _assert_existing_mapping_identity(
    existing: dict[str, Any],
    registry: dict[str, Any],
    patient: dict[str, Any],
) -> None:
    checks = (
        ("中心名称", normalized_text(existing.get("中心名称")), normalized_text(registry.get("中心名称"))),
        ("原全局编号", normalized_identifier(existing.get("原全局编号")), normalized_identifier(registry.get("#"))),
        ("住院号", normalized_identifier(existing.get("住院号")), normalized_identifier(patient.get("住院号"))),
        ("患者", normalized_name(existing.get("患者")), normalized_name(patient.get("患者"))),
    )
    conflicts = [label for label, old, current in checks if old and current and old != current]
    if conflicts:
        raise ValueError(
            f"匿名编号{existing['匿名编号']}的既有映射与本次数据不一致: {','.join(conflicts)}；请人工核对共享编号账本"
        )


def assign_anonymous_codes(
    registry_matches: dict[str, dict[str, Any]],
    patients: dict[str, dict[str, Any]],
    center_code: str,
    center_name: str,
    anonymous_map_rows: list[dict[str, Any]],
) -> None:
    """Reuse existing codes and append new codes in deterministic subject order."""

    validated = validate_anonymous_map(anonymous_map_rows)
    anonymous_map_rows[:] = validated
    by_key = {
        _mapping_key(row["中心编号"], row["受试者筛选号"]): row
        for row in anonymous_map_rows
    }
    next_number = max((int(row["匿名编号"]) for row in anonymous_map_rows), default=0) + 1
    for patient_key, registry in sorted(
        registry_matches.items(), key=_registry_assignment_sort_key
    ):
        patient = patients[patient_key]
        screening = text(registry.get("受试者筛选号"))
        key = _mapping_key(center_code, screening)
        existing = by_key.get(key)
        if existing:
            _assert_existing_mapping_identity(existing, registry, patient)
            registry["匿名编号"] = existing["匿名编号"]
            continue
        code = str(next_number).zfill(3)
        next_number += 1
        mapping_row = {
            "匿名编号": code,
            "中心编号": center_code,
            "中心名称": center_name,
            "受试者筛选号": screening,
            "原全局编号": text(registry.get("#")),
            "住院号": patient["住院号"],
            "患者": patient["患者"],
            "编号状态": (
                "人工确认分配（缺少患者号）"
                if patient.get("人工确认")
                else "已分配"
            ),
        }
        anonymous_map_rows.append(mapping_row)
        by_key[key] = mapping_row
        registry["匿名编号"] = code


def _exception(
    kind: str,
    message: str,
    *,
    source: str,
    source_row: Any = "",
    hospital_id: Any = "",
    patient_name: Any = "",
    study_uid: Any = "",
) -> dict[str, Any]:
    return {
        "异常类型": kind,
        "来源表": source,
        "来源行": source_row,
        "住院号": text(hospital_id),
        "患者": text(patient_name),
        "StudyInstanceUID": text(study_uid),
        "说明": message,
    }


def _unique_by_id(rows: Iterable[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    output: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        key = normalized_identifier(row.get("住院号"))
        if key:
            output[key].append(row)
    return output


def _match_enrolled_to_sequence(
    enrolled_rows: Iterable[dict[str, Any]], sequence_rows: Iterable[dict[str, Any]]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, dict[str, Any]]]:
    sequence_by_id = _unique_by_id(sequence_rows)
    selected: list[dict[str, Any]] = []
    exceptions: list[dict[str, Any]] = []
    patients: dict[str, dict[str, Any]] = {}
    seen_enrolled_names: dict[str, str] = {}

    for enrolled in enrolled_rows:
        hospital_id = text(enrolled.get("住院号"))
        patient_name = text(enrolled.get("患者"))
        identifier_key = normalized_identifier(hospital_id)
        name_key = normalized_name(patient_name)
        if not identifier_key:
            exceptions.append(
                _exception(
                    "入组患者缺少住院号",
                    "不允许仅凭姓名自动匹配",
                    source="符合入组患者表",
                    source_row=enrolled.get("__row__"),
                    patient_name=patient_name,
                )
            )
            continue
        if identifier_key in seen_enrolled_names:
            previous_name = seen_enrolled_names[identifier_key]
            if previous_name and name_key and previous_name != name_key:
                exceptions.append(
                    _exception(
                        "入组患者重复冲突",
                        "符合入组患者表中同一住院号对应多个姓名",
                        source="符合入组患者表",
                        source_row=enrolled.get("__row__"),
                        hospital_id=hospital_id,
                        patient_name=patient_name,
                    )
                )
            continue
        seen_enrolled_names[identifier_key] = name_key
        candidates = sequence_by_id.get(identifier_key, [])
        if not candidates:
            exceptions.append(
                _exception(
                    "入组患者未找到筛选序列",
                    "筛选序列表中没有相同住院号",
                    source="符合入组患者表",
                    source_row=enrolled.get("__row__"),
                    hospital_id=hospital_id,
                    patient_name=patient_name,
                )
            )
            continue

        candidate_names = {normalized_name(row.get("患者")) for row in candidates}
        candidate_names.discard("")
        if name_key and candidate_names and candidate_names != {name_key}:
            exceptions.append(
                _exception(
                    "姓名不一致",
                    "同一住院号在入组表和筛选序列表中的姓名不一致，未自动合并",
                    source="符合入组患者表/筛选序列表",
                    source_row=enrolled.get("__row__"),
                    hospital_id=hospital_id,
                    patient_name=patient_name,
                )
            )
            continue
        if len(candidate_names) > 1:
            exceptions.append(
                _exception(
                    "筛选序列身份冲突",
                    "同一住院号对应多个患者姓名",
                    source="筛选序列表",
                    hospital_id=hospital_id,
                    patient_name=patient_name,
                )
            )
            continue

        match_basis = "住院号+姓名" if name_key and candidate_names else "唯一住院号（姓名缺失）"
        patients[identifier_key] = {
            "住院号": hospital_id,
            "患者": patient_name or text(candidates[0].get("患者")),
            "匹配依据": match_basis,
            "入组行": enrolled.get("__row__"),
        }
        selected.extend(candidates)
    return selected, exceptions, patients


def _match_registry(
    patients: dict[str, dict[str, Any]],
    registry_rows: Iterable[dict[str, Any]],
    center_code: str,
    center_name: str,
) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    registry_rows = list(registry_rows)
    center_rows = [
        row
        for row in registry_rows
        if normalized_identifier(row.get("中心编号")) == normalized_identifier(center_code)
        and text(row.get("中心名称")) == center_name
    ]
    registry_by_id: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in center_rows:
        key = normalized_identifier(row.get("住院号/门诊号/放射号"))
        if key:
            registry_by_id[key].append(row)

    matches: dict[str, dict[str, Any]] = {}
    exceptions: list[dict[str, Any]] = []
    for key, patient in patients.items():
        candidates = registry_by_id.get(key, [])
        if len(candidates) != 1:
            exceptions.append(
                _exception(
                    "分中心映射不唯一" if candidates else "分中心映射缺失",
                    f"当前中心相同住院号记录数={len(candidates)}",
                    source="分中心影像表A:Q",
                    hospital_id=patient["住院号"],
                    patient_name=patient["患者"],
                )
            )
            continue
        registry = candidates[0]
        registry_name = normalized_name(registry.get("姓名"))
        patient_name = normalized_name(patient.get("患者"))
        if registry_name and patient_name and registry_name != patient_name:
            exceptions.append(
                _exception(
                    "分中心姓名不一致",
                    "相同住院号在分中心影像表中的姓名不一致，未生成匿名编号",
                    source="分中心影像表A:Q",
                    source_row=registry.get("__row__"),
                    hospital_id=patient["住院号"],
                    patient_name=patient["患者"],
                )
            )
            continue
        screening = text(registry.get("受试者筛选号"))
        screening_key = normalized_screening(screening)
        if not screening_key:
            exceptions.append(
                _exception(
                    "筛选号缺失",
                    "分中心影像表没有受试者筛选号",
                    source="分中心影像表A:Q",
                    source_row=registry.get("__row__"),
                    hospital_id=patient["住院号"],
                    patient_name=patient["患者"],
                )
            )
            continue
        center_number = ""
        if re.fullmatch(r"\d+(?:\.0+)?", text(center_code)):
            center_number = str(int(float(text(center_code))))
        if center_number and not screening_key.startswith(f"{center_number}-"):
            exceptions.append(
                _exception(
                    "筛选号中心前缀不一致",
                    f"筛选号{screening}与中心编号{center_code}不一致",
                    source="分中心影像表A:Q",
                    source_row=registry.get("__row__"),
                    hospital_id=patient["住院号"],
                    patient_name=patient["患者"],
                )
            )
            continue
        matches[key] = dict(registry)
    return matches, exceptions


def _screening_matches_center(screening: Any, center_code: Any) -> bool:
    screening_key = normalized_screening(screening)
    center_number = ""
    if re.fullmatch(r"\d+(?:\.0+)?", text(center_code)):
        center_number = str(int(float(text(center_code))))
    return not center_number or screening_key.startswith(f"{center_number}-")


def _apply_manual_missing_id_matches(
    manual_rows: Iterable[dict[str, Any]],
    sequence_rows: list[dict[str, Any]],
    enrolled_rows: list[dict[str, Any]],
    registry_rows: list[dict[str, Any]],
    center_code: str,
    center_name: str,
    patients: dict[str, dict[str, Any]],
    registry_matches: dict[str, dict[str, Any]],
    selected: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], set[str]]:
    """Apply explicit Study-to-subject confirmations for enrolled rows lacking IDs.

    The public ``住院号`` value remains blank.  A private in-memory patient key keeps
    these rows separate without inventing a clinical identifier.
    """

    exceptions: list[dict[str, Any]] = []
    resolved_enrolled_rows: set[str] = set()
    active_rows: list[dict[str, Any]] = []
    recognized_no_values = {"否", "n", "no", "false", "0"}
    for row in manual_rows:
        confirmation = normalized_text(row.get("是否确认同一人"))
        if is_yes(row.get("是否确认同一人")):
            active_rows.append(row)
        elif confirmation and confirmation not in recognized_no_values:
            exceptions.append(
                _exception(
                    "人工确认值无效",
                    "“是否确认同一人”只能填写“是”或“否”；该行未采用",
                    source="缺失患者号人工确认表",
                    source_row=row.get("__row__"),
                    patient_name=row.get("患者"),
                    study_uid=row.get("StudyInstanceUID"),
                )
            )

    if not active_rows:
        return exceptions, resolved_enrolled_rows

    center_rows = [
        row
        for row in registry_rows
        if normalized_identifier(row.get("中心编号"))
        == normalized_identifier(center_code)
        and text(row.get("中心名称")) == center_name
    ]
    registry_by_screening: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in center_rows:
        screening_key = normalized_screening(row.get("受试者筛选号"))
        if screening_key:
            registry_by_screening[screening_key].append(row)

    missing_enrolled_by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in enrolled_rows:
        if normalized_identifier(row.get("住院号")):
            continue
        name_key = normalized_name(row.get("患者"))
        if name_key:
            missing_enrolled_by_name[name_key].append(row)

    sequence_by_study: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in sequence_rows:
        study_uid = text(row.get("StudyInstanceUID"))
        if study_uid:
            sequence_by_study[study_uid].append(row)

    regular_screenings = {
        normalized_screening(registry.get("受试者筛选号"))
        for registry in registry_matches.values()
    }
    regular_study_owners: dict[str, set[str]] = defaultdict(set)
    for row in selected:
        study_uid = text(row.get("StudyInstanceUID"))
        patient_key = text(row.get(INTERNAL_PATIENT_KEY))
        if study_uid and patient_key:
            regular_study_owners[study_uid].add(patient_key)

    candidates: list[dict[str, Any]] = []
    for row in active_rows:
        screening = text(row.get("受试者筛选号"))
        screening_key = normalized_screening(screening)
        patient_name = text(row.get("患者"))
        name_key = normalized_name(patient_name)
        study_uid = text(row.get("StudyInstanceUID"))
        confirmation_basis = text(row.get("确认依据"))
        missing_fields = [
            label
            for label, value in (
                ("受试者筛选号", screening_key),
                ("患者", name_key),
                ("StudyInstanceUID", study_uid),
                ("确认依据", confirmation_basis),
            )
            if not value
        ]
        if missing_fields:
            exceptions.append(
                _exception(
                    "人工确认信息不完整",
                    f"明确确认的行缺少字段: {','.join(missing_fields)}；该行未采用",
                    source="缺失患者号人工确认表",
                    source_row=row.get("__row__"),
                    patient_name=patient_name,
                    study_uid=study_uid,
                )
            )
            continue
        if not _screening_matches_center(screening, center_code):
            exceptions.append(
                _exception(
                    "人工确认筛选号中心不一致",
                    f"筛选号{screening}与当前中心编号{center_code}不一致；该行未采用",
                    source="缺失患者号人工确认表",
                    source_row=row.get("__row__"),
                    patient_name=patient_name,
                    study_uid=study_uid,
                )
            )
            continue
        registry_candidates = registry_by_screening.get(screening_key, [])
        if len(registry_candidates) != 1:
            exceptions.append(
                _exception(
                    "人工确认分中心映射不唯一"
                    if registry_candidates
                    else "人工确认分中心映射缺失",
                    f"当前中心相同筛选号记录数={len(registry_candidates)}；该行未采用",
                    source="缺失患者号人工确认表/分中心影像表A:Q",
                    source_row=row.get("__row__"),
                    patient_name=patient_name,
                    study_uid=study_uid,
                )
            )
            continue
        registry = registry_candidates[0]
        registry_name = normalized_name(registry.get("姓名"))
        if registry_name and registry_name != name_key:
            exceptions.append(
                _exception(
                    "人工确认姓名不一致",
                    "人工确认表患者与分中心影像表姓名/缩写不一致；该行未采用",
                    source="缺失患者号人工确认表/分中心影像表A:Q",
                    source_row=row.get("__row__"),
                    patient_name=patient_name,
                    study_uid=study_uid,
                )
            )
            continue
        enrolled_candidates = missing_enrolled_by_name.get(name_key, [])
        if len(enrolled_candidates) != 1:
            exceptions.append(
                _exception(
                    "人工确认入组患者不唯一"
                    if enrolled_candidates
                    else "人工确认未找到入组患者",
                    f"缺少患者号且姓名/缩写相同的入组记录数={len(enrolled_candidates)}；该行未采用",
                    source="缺失患者号人工确认表/符合入组患者表",
                    source_row=row.get("__row__"),
                    patient_name=patient_name,
                    study_uid=study_uid,
                )
            )
            continue
        study_rows = sequence_by_study.get(study_uid, [])
        if not study_rows:
            exceptions.append(
                _exception(
                    "人工确认Study不存在",
                    "筛选序列表中没有该StudyInstanceUID；该行未采用",
                    source="缺失患者号人工确认表/筛选序列表",
                    source_row=row.get("__row__"),
                    patient_name=patient_name,
                    study_uid=study_uid,
                )
            )
            continue
        sequence_names = {
            normalized_name(item.get("患者"))
            for item in study_rows
            if normalized_name(item.get("患者"))
        }
        if sequence_names and sequence_names != {name_key}:
            exceptions.append(
                _exception(
                    "人工确认Study姓名不一致",
                    "该Study在筛选序列表中的患者姓名/缩写与人工确认表不一致；该行未采用",
                    source="缺失患者号人工确认表/筛选序列表",
                    source_row=row.get("__row__"),
                    patient_name=patient_name,
                    study_uid=study_uid,
                )
            )
            continue
        sequence_ids = {
            normalized_identifier(item.get("住院号"))
            for item in study_rows
            if normalized_identifier(item.get("住院号"))
        }
        if len(sequence_ids) > 1:
            exceptions.append(
                _exception(
                    "人工确认Study身份冲突",
                    "同一StudyInstanceUID在筛选序列表中出现多个非空住院号；该行未采用",
                    source="缺失患者号人工确认表/筛选序列表",
                    source_row=row.get("__row__"),
                    patient_name=patient_name,
                    study_uid=study_uid,
                )
            )
            continue
        if screening_key in regular_screenings:
            exceptions.append(
                _exception(
                    "人工确认筛选号已匹配",
                    "该筛选号已经通过常规住院号规则匹配，不允许再次人工覆盖",
                    source="缺失患者号人工确认表",
                    source_row=row.get("__row__"),
                    patient_name=patient_name,
                    study_uid=study_uid,
                )
            )
            continue
        if regular_study_owners.get(study_uid):
            exceptions.append(
                _exception(
                    "人工确认Study已归属",
                    "该Study已通过常规住院号规则归属其他入组记录，不允许人工覆盖",
                    source="缺失患者号人工确认表",
                    source_row=row.get("__row__"),
                    patient_name=patient_name,
                    study_uid=study_uid,
                )
            )
            continue
        candidates.append(
            {
                "manual_row": row,
                "screening": screening,
                "screening_key": screening_key,
                "patient_name": patient_name,
                "name_key": name_key,
                "study_uid": study_uid,
                "confirmation_basis": confirmation_basis,
                "registry": registry,
                "enrolled": enrolled_candidates[0],
                "study_rows": study_rows,
            }
        )

    study_owners: dict[str, set[str]] = defaultdict(set)
    screening_names: dict[str, set[str]] = defaultdict(set)
    enrolled_screenings: dict[tuple[str, str], set[str]] = defaultdict(set)
    for candidate in candidates:
        study_owners[candidate["study_uid"]].add(candidate["screening_key"])
        screening_names[candidate["screening_key"]].add(candidate["name_key"])
        enrolled_key = (
            text(candidate["enrolled"].get("__row__")),
            candidate["name_key"],
        )
        enrolled_screenings[enrolled_key].add(candidate["screening_key"])

    conflicting_studies = {
        study_uid for study_uid, owners in study_owners.items() if len(owners) > 1
    }
    conflicting_screenings = {
        screening for screening, names in screening_names.items() if len(names) > 1
    }
    conflicting_enrolled = {
        enrolled_key
        for enrolled_key, screenings in enrolled_screenings.items()
        if len(screenings) > 1
    }
    accepted_pairs: set[tuple[str, str]] = set()
    for candidate in candidates:
        enrolled_key = (
            text(candidate["enrolled"].get("__row__")),
            candidate["name_key"],
        )
        conflict_message = ""
        if candidate["study_uid"] in conflicting_studies:
            conflict_message = "同一StudyInstanceUID被人工分配给多个筛选号"
        elif candidate["screening_key"] in conflicting_screenings:
            conflict_message = "同一筛选号在人工确认表中对应多个患者姓名/缩写"
        elif enrolled_key in conflicting_enrolled:
            conflict_message = "同一条缺号入组记录被人工分配给多个筛选号"
        if conflict_message:
            exceptions.append(
                _exception(
                    "人工确认归属冲突",
                    f"{conflict_message}；相关行均未采用",
                    source="缺失患者号人工确认表",
                    source_row=candidate["manual_row"].get("__row__"),
                    patient_name=candidate["patient_name"],
                    study_uid=candidate["study_uid"],
                )
            )
            continue

        pair = (candidate["screening_key"], candidate["study_uid"])
        if pair in accepted_pairs:
            continue
        accepted_pairs.add(pair)
        patient_key = (
            f"manual:{normalized_identifier(center_code)}:{candidate['screening_key']}"
        )
        existing_patient = patients.get(patient_key)
        if existing_patient and normalized_name(existing_patient.get("患者")) != candidate["name_key"]:
            exceptions.append(
                _exception(
                    "人工确认患者冲突",
                    "同一人工患者键对应不同姓名/缩写；该行未采用",
                    source="缺失患者号人工确认表",
                    source_row=candidate["manual_row"].get("__row__"),
                    patient_name=candidate["patient_name"],
                    study_uid=candidate["study_uid"],
                )
            )
            continue
        patients.setdefault(
            patient_key,
            {
                "住院号": "",
                "患者": candidate["patient_name"],
                "匹配依据": "人工确认（缺少患者号；筛选号+StudyInstanceUID）",
                "入组行": candidate["enrolled"].get("__row__"),
                "人工确认": True,
            },
        )
        registry_matches.setdefault(patient_key, dict(candidate["registry"]))
        for source_row in candidate["study_rows"]:
            selected_row = dict(source_row)
            selected_row[INTERNAL_PATIENT_KEY] = patient_key
            selected_row[INTERNAL_MANUAL_BASIS] = candidate["confirmation_basis"]
            selected_row[INTERNAL_MANUAL_ROW] = candidate["manual_row"].get("__row__")
            selected_row[INTERNAL_MANUAL_PERSON] = text(
                candidate["manual_row"].get("确认人")
            )
            selected_row[INTERNAL_MANUAL_TIME] = text(
                candidate["manual_row"].get("确认时间")
            )
            selected.append(selected_row)
        resolved_enrolled_rows.add(text(candidate["enrolled"].get("__row__")))

    return exceptions, resolved_enrolled_rows


def prepare_selection(
    sequence_rows: Iterable[dict[str, Any]],
    enrolled_rows: Iterable[dict[str, Any]],
    registry_rows: Iterable[dict[str, Any]],
    center_code: str,
    center_name: str,
    anonymous_map_rows: list[dict[str, Any]] | None = None,
    manual_match_rows: Iterable[dict[str, Any]] | None = None,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    sequence_rows = list(sequence_rows)
    enrolled_rows = list(enrolled_rows)
    registry_rows = list(registry_rows)
    if anonymous_map_rows is None:
        anonymous_map_rows = []
    selected, exceptions, patients = _match_enrolled_to_sequence(enrolled_rows, sequence_rows)
    selected = [
        {
            **row,
            INTERNAL_PATIENT_KEY: normalized_identifier(row.get("住院号")),
        }
        for row in selected
    ]
    registry_matches, registry_exceptions = _match_registry(
        patients, registry_rows, center_code, center_name
    )
    exceptions.extend(registry_exceptions)
    manual_exceptions, resolved_enrolled_rows = _apply_manual_missing_id_matches(
        manual_match_rows or [],
        sequence_rows,
        enrolled_rows,
        registry_rows,
        center_code,
        center_name,
        patients,
        registry_matches,
        selected,
    )
    if resolved_enrolled_rows:
        exceptions = [
            item
            for item in exceptions
            if not (
                item.get("异常类型") == "入组患者缺少住院号"
                and text(item.get("来源行")) in resolved_enrolled_rows
            )
        ]
    exceptions.extend(manual_exceptions)
    assign_anonymous_codes(
        registry_matches,
        patients,
        center_code,
        center_name,
        anonymous_map_rows,
    )

    valid_patient_keys = set(registry_matches)
    selected = [
        row
        for row in selected
        if text(row.get(INTERNAL_PATIENT_KEY)) in valid_patient_keys
    ]

    study_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in selected:
        study_uid = text(row.get("StudyInstanceUID"))
        if not study_uid:
            exceptions.append(
                _exception(
                    "StudyInstanceUID缺失",
                    "该明细不会进入匿名化",
                    source="筛选序列表",
                    source_row=row.get("__row__"),
                    hospital_id=row.get("住院号"),
                    patient_name=row.get("患者"),
                )
            )
            continue
        study_groups[study_uid].append(row)

    main_rows: list[dict[str, Any]] = []
    detail_rows: list[dict[str, Any]] = []
    mapping_rows: list[dict[str, Any]] = []
    study_owner: dict[str, str] = {}

    for patient_key, registry in sorted(
        registry_matches.items(), key=lambda item: int(item[1]["匿名编号"])
    ):
        patient = patients[patient_key]
        mapping_rows.append(
            {
                "中心编号": center_code,
                "中心名称": center_name,
                "住院号": patient["住院号"],
                "患者": patient["患者"],
                "受试者筛选号": text(registry.get("受试者筛选号")),
                "全局编号": text(registry.get("#")),
                "匿名编号": registry["匿名编号"],
                "匹配依据": patient["匹配依据"],
            }
        )

    study_infos: list[dict[str, Any]] = []
    for study_uid, rows in study_groups.items():
        patient_keys = {text(row.get(INTERNAL_PATIENT_KEY)) for row in rows}
        if len(patient_keys) != 1:
            exceptions.append(
                _exception(
                    "Study归属多个患者",
                    "同一StudyInstanceUID在筛选序列表中对应多个住院号",
                    source="筛选序列表",
                    study_uid=study_uid,
                )
            )
            continue
        patient_key = next(iter(patient_keys))
        previous_owner = study_owner.get(study_uid)
        if previous_owner and previous_owner != patient_key:
            exceptions.append(
                _exception(
                    "Study归属冲突",
                    "同一StudyInstanceUID不能分配给多个患者",
                    source="筛选序列表",
                    study_uid=study_uid,
                )
            )
            continue
        study_owner[study_uid] = patient_key
        patient = patients[patient_key]
        registry = registry_matches[patient_key]
        code = registry["匿名编号"]
        period = classify_period(row.get("时期") for row in rows)
        study_infos.append(
            {
                "study_uid": study_uid,
                "rows": rows,
                "patient_key": patient_key,
                "patient": patient,
                "registry": registry,
                "code": code,
                "period": period,
            }
        )

    twelve_month_candidates: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for info in study_infos:
        if info["period"].is_12m_candidate:
            twelve_month_candidates[info["patient_key"]].append(info)

    selected_twelve_month: dict[str, dict[str, Any]] = {}
    tied_twelve_month_studies: set[str] = set()
    for patient_key, candidates in twelve_month_candidates.items():
        minimum_distance = min(abs(info["period"].followup_month - 12) for info in candidates)
        nearest = [
            info
            for info in candidates
            if math.isclose(
                abs(info["period"].followup_month - 12),
                minimum_distance,
                rel_tol=0,
                abs_tol=1e-9,
            )
        ]
        if len(nearest) == 1:
            selected_twelve_month[patient_key] = nearest[0]
            continue
        patient = patients[patient_key]
        tied_twelve_month_studies.update(info["study_uid"] for info in nearest)
        candidate_display = "；".join(
            f"{info['study_uid']}={info['period'].followup_month:g}月" for info in nearest
        )
        exceptions.append(
            _exception(
                "12M候选并列",
                f"多个Study与12个月距离相同，请人工选择其中一个: {candidate_display}",
                source="筛选序列表",
                hospital_id=patient["住院号"],
                patient_name=patient["患者"],
                study_uid="、".join(info["study_uid"] for info in nearest),
            )
        )

    for info in study_infos:
        study_uid = info["study_uid"]
        rows = info["rows"]
        patient_key = info["patient_key"]
        patient = info["patient"]
        registry = info["registry"]
        code = info["code"]
        period = info["period"]
        original_periods = "、".join(
            dict.fromkeys(text(row.get("时期")) or "<空>" for row in rows)
        )
        acquisition_dates = "、".join(
            dict.fromkeys(text(row.get("AcqusitionDate")) for row in rows if text(row.get("AcqusitionDate")))
        )
        manually_confirmed = bool(patient.get("人工确认"))
        approved = period.valid
        match_status = "已匹配" if period.valid else "时期待确认"
        note = period.note
        target = f"{code}-{period.target_suffix}" if period.valid else ""
        if period.is_12m_candidate:
            target = f"{code}-12M"
            selected = selected_twelve_month.get(patient_key)
            month = period.followup_month
            if study_uid in tied_twelve_month_studies:
                approved = False
                match_status = "12M并列待确认"
                note = f"12M候选月份={month:g}；与另一Study同样接近12个月，请人工仅选择一个"
            elif selected and selected["study_uid"] == study_uid:
                approved = True
                match_status = "已选为最接近12M"
                if 9 <= month <= 15:
                    note = f"12M候选月份={month:g}；位于9～15个月窗口且为该患者最接近12个月的数据"
                else:
                    note = f"12M候选月份={month:g}；该患者无更接近数据，按9个月以后最近值纳入12M"
            else:
                approved = False
                match_status = "12M候选未选"
                selected_month = selected["period"].followup_month if selected else None
                selected_uid = selected["study_uid"] if selected else ""
                note = (
                    f"12M候选月份={month:g}；未选择，另一个Study更接近12个月: "
                    f"{selected_uid}={selected_month:g}月"
                )
        identity_note = patient["匹配依据"]
        if manually_confirmed:
            match_status = (
                "人工确认进入（缺少患者号）"
                if match_status == "已匹配"
                else f"人工确认（缺少患者号）；{match_status}"
            )
            confirmation_bases = "、".join(
                dict.fromkeys(
                    text(row.get(INTERNAL_MANUAL_BASIS))
                    for row in rows
                    if text(row.get(INTERNAL_MANUAL_BASIS))
                )
            )
            confirmers = "、".join(
                dict.fromkeys(
                    text(row.get(INTERNAL_MANUAL_PERSON))
                    for row in rows
                    if text(row.get(INTERNAL_MANUAL_PERSON))
                )
            )
            confirmation_times = "、".join(
                dict.fromkeys(
                    text(row.get(INTERNAL_MANUAL_TIME))
                    for row in rows
                    if text(row.get(INTERNAL_MANUAL_TIME))
                )
            )
            identity_note += f"；确认依据={confirmation_bases}"
            if confirmers:
                identity_note += f"；确认人={confirmers}"
            if confirmation_times:
                identity_note += f"；确认时间={confirmation_times}"
        main_rows.append(
            {
                "中心编号": center_code,
                "中心名称": center_name,
                "住院号": patient["住院号"],
                "患者": patient["患者"],
                "受试者筛选号": text(registry.get("受试者筛选号")),
                "全局编号": text(registry.get("#")),
                "匿名编号": code,
                "StudyInstanceUID": study_uid,
                "AcqusitionDate": acquisition_dates,
                "原时期": original_periods,
                "标准访视期": period.category,
                "目标二级目录": target,
                "原明细行数": len(rows),
                "是否进入匿名化": "是" if approved else "否",
                "匹配状态": match_status,
                "说明": f"{identity_note}；{note}",
            }
        )
        for row in rows:
            detail = {
                key: value for key, value in row.items() if not key.startswith("__")
            }
            detail.update(
                {
                    "来源行": row.get("__row__"),
                    "中心编号": center_code,
                    "中心名称": center_name,
                    "受试者筛选号": text(registry.get("受试者筛选号")),
                    "匿名编号": code,
                    "标准访视期": period.category,
                    "目标二级目录": target,
                    "是否进入匿名化": "是" if approved else "否",
                    "匹配状态": match_status,
                    "选择说明": note,
                }
            )
            if manually_confirmed:
                detail.update(
                    {
                        "人工确认表来源行": row.get(INTERNAL_MANUAL_ROW),
                        "人工确认依据": row.get(INTERNAL_MANUAL_BASIS),
                        "确认人": row.get(INTERNAL_MANUAL_PERSON),
                        "确认时间": row.get(INTERNAL_MANUAL_TIME),
                    }
                )
            detail_rows.append(detail)

    main_rows.sort(key=lambda row: (int(row["匿名编号"]), row["标准访视期"], row["StudyInstanceUID"]))
    return main_rows, detail_rows, mapping_rows, exceptions


def _write_sheet(worksheet: Any, headers: list[str], rows: Iterable[dict[str, Any]]) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])

    white = PatternFill("solid", fgColor="FFFFFF")
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in worksheet.iter_rows():
        for cell in row:
            cell.fill = white
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 30
    for column_cells in worksheet.columns:
        letter = column_cells[0].column_letter
        maximum = max((len(text(cell.value)) for cell in column_cells[:200]), default=8)
        worksheet.column_dimensions[letter].width = min(max(maximum + 2, 10), 42)


def create_workbook(
    main_rows: list[dict[str, Any]],
    detail_rows: list[dict[str, Any]],
    mapping_rows: list[dict[str, Any]],
    exceptions: list[dict[str, Any]],
) -> Any:
    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError(
            "生成.xlsx需要openpyxl，请执行: python -m pip install openpyxl"
        ) from exc

    workbook = Workbook()
    main = workbook.active
    main.title = "入组Study清单"
    _write_sheet(main, list(MAIN_HEADERS), main_rows)
    yes_column = MAIN_HEADERS.index("是否进入匿名化") + 1
    if main.max_row >= 2:
        validation = DataValidation(type="list", formula1='"是,否"', allow_blank=False)
        main.add_data_validation(validation)
        validation.add(f"{main.cell(1, yes_column).column_letter}2:{main.cell(main.max_row, yes_column).column_letter}{main.max_row}")

    detail = workbook.create_sheet("筛选序列明细")
    detail_headers: list[str] = []
    for row in detail_rows:
        for key in row:
            if key not in detail_headers:
                detail_headers.append(key)
    _write_sheet(detail, detail_headers or ["说明"], detail_rows)

    mapping = workbook.create_sheet("患者匿名编号映射")
    mapping_headers = [
        "中心编号",
        "中心名称",
        "住院号",
        "患者",
        "受试者筛选号",
        "全局编号",
        "匿名编号",
        "匹配依据",
    ]
    _write_sheet(mapping, mapping_headers, mapping_rows)

    exception_sheet = workbook.create_sheet("匹配异常")
    exception_headers = ["异常类型", "来源表", "来源行", "住院号", "患者", "StudyInstanceUID", "说明"]
    _write_sheet(exception_sheet, exception_headers, exceptions)

    rules = workbook.create_sheet("规则说明")
    rule_rows = [
        {"项目": "筛选规则版本", "规则": SELECTION_RULE_VERSION},
        {"项目": "处理单位", "规则": "同一StudyInstanceUID的全部明细先聚合，再判断时期。"},
        {"项目": "术前", "规则": "同一Study全部时期为“术前”。"},
        {"项目": "术中", "规则": "同一Study全部时期为“术中”。"},
        {"项目": "术前与术中", "规则": "同一Study同时出现“术前”和“术中”。"},
        {"项目": "6M", "规则": "数值或“术后N个月”满足3≤N<9。"},
        {"项目": "12M优先窗口", "规则": "12个月正负3个月，即9≤N≤15；先按患者寻找离12个月最近的Study。"},
        {"项目": "12M窗口外补选", "规则": "若没有更近数据，可从N>15的Study中选择离12个月最近的一项，例如仅有18.4个月时纳入12M。"},
        {"项目": "12M唯一性", "规则": "同一患者只默认选择一个12M Study；其他候选为“否”，距离并列时全部待人工确认。"},
        {"项目": "时期待确认", "规则": "空值、无法识别、小于3个月、跨类别混合均默认不进入匿名化。"},
        {"项目": "匿名编号", "规则": "使用跨中心共享编号账本；首位患者从001开始，新中心从账本最大编号加1继续，既有中心+筛选号始终复用原编号。"},
        {"项目": "原全局编号", "规则": "分中心影像表A:Q主表的#列仅保留用于内部追溯，不再决定匿名编号。"},
        {"项目": "缺少患者号人工确认", "规则": "不按姓名自动匹配；仅采用人工确认表中明确为“是”、且通过筛选号、StudyUID及身份冲突校验的绑定。"},
        {"项目": "匿名化准入确认", "规则": "仅“是否进入匿名化”为“是”的Study才会被第二步脚本处理。"},
        {"项目": "敏感性", "规则": "本工作簿含原筛选号、住院号和姓名，只能存放在内部控制目录，不得交付。"},
    ]
    _write_sheet(rules, ["项目", "规则"], rule_rows)
    return workbook


def create_anonymous_map_workbook(mapping_rows: list[dict[str, Any]]) -> Any:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - environment-specific
        raise RuntimeError(
            "生成.xlsx需要openpyxl，请执行: python -m pip install openpyxl"
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "患者匿名编号映射"
    _write_sheet(
        sheet,
        list(ANONYMOUS_MAP_HEADERS),
        sorted(mapping_rows, key=lambda row: int(row["匿名编号"])),
    )
    notes = workbook.create_sheet("使用说明")
    _write_sheet(
        notes,
        ["项目", "说明"],
        [
            {"项目": "筛选规则版本", "说明": SELECTION_RULE_VERSION},
            {"项目": "用途", "说明": "跨中心连续分配匿名编号；所有中心必须使用同一个文件。"},
            {"项目": "唯一键", "说明": "中心编号+受试者筛选号。"},
            {"项目": "分配规则", "说明": "空账本从001开始；新患者从当前最大编号加1；既有患者复用原编号。"},
            {"项目": "安全要求", "说明": "本表含筛选号、住院号和姓名，只能存放在内部控制目录，不得交付或删除重建。"},
            {"项目": "并发要求", "说明": "一次只运行一个中心，避免两个任务同时修改同一账本。"},
        ],
    )
    return workbook


def run(args: argparse.Namespace) -> dict[str, Any]:
    output = Path(args.output).resolve()
    if output.exists() and not args.overwrite:
        raise FileExistsError(f"输出文件已存在，未覆盖: {output}；确认后使用--overwrite")
    anonymous_map_arg = getattr(args, "anonymous_map", None)
    anonymous_map_path = (
        Path(anonymous_map_arg).resolve()
        if anonymous_map_arg
        else output.parent / "患者匿名编号映射.xlsx"
    )
    if anonymous_map_path == output:
        raise ValueError("--anonymous-map不能与--output使用同一个文件")
    manual_match_arg = getattr(args, "manual_match_workbook", None)
    manual_match_path = Path(manual_match_arg).resolve() if manual_match_arg else None
    if manual_match_path in {output, anonymous_map_path}:
        raise ValueError("--manual-match-workbook不能与--output或--anonymous-map使用同一个文件")
    anonymous_map_rows = load_anonymous_map(anonymous_map_path)
    existing_mapping_count = len(anonymous_map_rows)
    sequence = read_table(
        args.sequence_workbook,
        SEQUENCE_REQUIRED,
        sheet_name=args.sequence_sheet,
        header_aliases={
            "AcqusitionDate": ("AcquisitionDate", "StudyDate", "SeriesDate"),
            "SeriesUID": ("SeriesInstanceUID",),
            "SOPUID": ("SOPInstanceUID",),
        },
    )
    enrolled = read_table(
        args.enrolled_workbook, ENROLLED_REQUIRED, sheet_name=args.enrolled_sheet
    )
    manual_match_rows: Iterable[dict[str, Any]] = []
    if manual_match_path is not None:
        manual_match = read_table(
            manual_match_path,
            MANUAL_MATCH_REQUIRED,
            sheet_name=getattr(args, "manual_match_sheet", None),
            header_aliases={
                "受试者筛选号": ("筛选号",),
                "患者": ("姓名", "姓名缩写"),
                "StudyInstanceUID": ("StudyUID",),
                "是否确认同一人": ("是否确认",),
                "确认依据": ("人工确认依据",),
            },
        )
        manual_match_rows = manual_match.rows
    registry = read_table(
        args.registry_workbook,
        REGISTRY_REQUIRED,
        sheet_name=args.registry_sheet,
        max_columns=17,
    )
    center_code, center_name = choose_center(registry.rows, args.center)
    main_rows, detail_rows, mapping_rows, exceptions = prepare_selection(
        sequence.rows,
        enrolled.rows,
        registry.rows,
        center_code,
        center_name,
        anonymous_map_rows,
        manual_match_rows,
    )
    workbook = create_workbook(main_rows, detail_rows, mapping_rows, exceptions)
    map_workbook = create_anonymous_map_workbook(anonymous_map_rows)
    atomic_save_workbook(map_workbook, anonymous_map_path, True)
    atomic_save_workbook(workbook, output, args.overwrite)
    return {
        "center_code": center_code,
        "center_name": center_name,
        "patients": len(mapping_rows),
        "studies": len(main_rows),
        "approved_by_default": sum(row["是否进入匿名化"] == "是" for row in main_rows),
        "pending": sum(row["是否进入匿名化"] != "是" for row in main_rows),
        "exceptions": len(exceptions),
        "mapping_total": len(anonymous_map_rows),
        "mapping_new": len(anonymous_map_rows) - existing_mapping_count,
        "manual_patients": sum(
            "人工确认" in text(row.get("匹配依据")) for row in mapping_rows
        ),
        "manual_studies": sum(
            "人工确认" in text(row.get("匹配状态")) for row in main_rows
        ),
        "manual_match_workbook": str(manual_match_path) if manual_match_path else "",
        "anonymous_map": str(anonymous_map_path),
        "output": str(output.resolve()),
    }


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        summary = run(args)
    except (OSError, ValueError, RuntimeError) as exc:
        print(f"错误: {exc}", file=sys.stderr)
        return 1
    print(f"筛选规则版本: {SELECTION_RULE_VERSION}")
    print(f"中心: {summary['center_code']} {summary['center_name']}")
    print(f"入组患者: {summary['patients']}")
    print(f"Study: {summary['studies']}（默认进入 {summary['approved_by_default']}，待确认 {summary['pending']}）")
    print(f"匹配异常: {summary['exceptions']}")
    if summary["manual_match_workbook"]:
        print(
            "缺号人工确认: "
            f"患者{summary['manual_patients']}，Study{summary['manual_studies']}"
        )
        print(f"人工确认表: {summary['manual_match_workbook']}")
    print(f"共享匿名编号账本: 共{summary['mapping_total']}人，本次新增{summary['mapping_new']}人")
    print(f"编号账本: {summary['anonymous_map']}")
    print(f"输出: {summary['output']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
