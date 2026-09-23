#!/usr/bin/env python3
"""Linux-safe Excel writer for the center image-web report.

The Windows report uses the desktop artifact runtime.  This module intentionally
uses openpyxl so the same report can be generated on a normal Linux server with
public Python dependencies only.
"""

from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any, Iterable, Sequence


EXCEPTION_HEADERS = [
    "中心",
    "异常类型",
    "患者",
    "住院号",
    "StudyInstanceUID",
    "SeriesUID",
    "SOPUID",
    "来源文件或行",
    "说明",
]

SUMMARY_HEADERS = [
    "中心工作表",
    "中心目录名",
    "分中心表中心名称",
    "中心匹配依据",
    "扫描文件数",
    "DICOM数",
    "患者数",
    "Study数",
    "Series数",
    "输出行数",
    "已匹配行数",
    "异常数",
]


def require_openpyxl() -> None:
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:  # pragma: no cover - depends on target machine
        raise RuntimeError(
            "Linux版生成Excel需要openpyxl；请运行: "
            "python3 -m pip install -r requirements-linux.txt"
        ) from exc


def _definitions(payload: dict[str, Any]) -> list[list[Any]]:
    return [
        ["中心工作表", "一个中心一个工作表", "工作表名称取转存根目录的一级中心目录名"],
        ["一行的含义", "第四批混合粒度", "单帧DICOM按Series汇总；NumberOfFrames>1按SOP输出"],
        ["住院号", "住院号/门诊号/放射号", "先限定中心，再按规范姓名匹配；住院号缺失时保持空白，不使用筛选号代填"],
        ["编号回退", "姓名不一致或缺失时的关联规则", "患者目录15_H001或01_Q001规范为15-H001或01-Q001，并与分中心表受试者筛选号唯一匹配"],
        ["患者", "患者显示名称", "优先分中心表姓名；缺失时使用姓名缩写；仍缺失时使用DICOM姓名或筛选号"],
        [
            "AcqusitionDate",
            "影像采集日期",
            "优先AcquisitionDate；缺失时依次使用StudyDate、SeriesDate；表头保留第四批历史拼写",
        ],
        ["时期", "相对手术日期", "负数=术前；0=术中；正数=相差天数/30，保留两位小数"],
        ["SOPUID", "多帧对象标识", "多帧按SOP保留；单帧Series汇总行填NA"],
        ["NumberOfFrames", "多帧对象帧数", "多帧SOP读取DICOM标签；单帧Series填NA"],
        ["帧数", "单帧序列实例数", "按Series内去重后的SOPInstanceUID计数；多帧SOP填NA"],
        ["日期回退", "AcquisitionDate缺失", "依次使用StudyDate、SeriesDate，并在处理异常中记录日期来源"],
        ["日期冲突", "同一Series有多个候选日期", "主表使用当前优先级字段的最早日期，同时进入处理异常"],
        ["缺失日期", "无法计算时期", "AcquisitionDate、StudyDate和SeriesDate均缺失时时期留空"],
        ["数据安全", "只读分析", "不修改、不移动、不删除转存后DICOM及V3状态文件"],
        [
            "规则版本",
            payload.get("rule_version", ""),
            f"生成时间 {payload.get('generated_at', '')}",
        ],
    ]


def _style_grid(worksheet: Any, *, max_row: int, max_column: int) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    white = PatternFill(fill_type="solid", fgColor="FFFFFF")
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    data_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_border = Border(left=thin, right=thin, top=thin, bottom=medium)

    worksheet.sheet_view.showGridLines = False
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=max_row,
        min_col=1,
        max_col=max_column,
    ):
        for cell in row:
            cell.fill = white
            cell.border = data_border
            cell.font = Font(name="等线", size=11, color="000000")
            cell.alignment = Alignment(vertical="center")

    for cell in worksheet[1][:max_column]:
        cell.fill = white
        cell.border = header_border
        cell.font = Font(name="Times New Roman", size=11, bold=True, color="000000")
        cell.alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 24


def _add_table(worksheet: Any, reference: str, name: str) -> None:
    from openpyxl.worksheet.table import Table, TableStyleInfo

    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight1",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def _set_widths(worksheet: Any, widths: Sequence[float]) -> None:
    from openpyxl.utils import get_column_letter

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width


def _append_dict_rows(
    worksheet: Any,
    headers: Sequence[str],
    rows: Iterable[dict[str, Any]],
) -> int:
    count = 0
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])
        count += 1
    return count


def _add_center_sheet(
    workbook: Any,
    center: dict[str, Any],
    headers: Sequence[str],
    index: int,
) -> None:
    from openpyxl.styles import Alignment

    worksheet = workbook.create_sheet(center["sheet_name"])
    worksheet.append(list(headers))
    row_count = _append_dict_rows(worksheet, headers, center.get("rows", []))
    end_row = row_count + 1
    _style_grid(worksheet, max_row=end_row, max_column=len(headers))
    worksheet.freeze_panes = "A2"
    _set_widths(
        worksheet,
        [15, 12, 46, 46, 46, 14, 12, 48, 42, 14, 17, 12, 12, 17, 17],
    )

    for row in worksheet.iter_rows(min_row=2, max_row=end_row, min_col=1, max_col=5):
        for cell in row:
            cell.number_format = "@"
    for cell in worksheet["F"][1:end_row]:
        cell.number_format = "0"
    for cell in worksheet["G"][1:end_row]:
        cell.number_format = "0.00"
    for column in ("J", "K", "L"):
        for cell in worksheet[column][1:end_row]:
            cell.alignment = Alignment(horizontal="right", vertical="center")
    for column in ("N", "O"):
        for cell in worksheet[column][1:end_row]:
            cell.number_format = "0.0"
    if row_count:
        _add_table(worksheet, f"A1:O{end_row}", f"CenterImageTable{index + 1}")


def _add_exception_sheet(workbook: Any, exceptions: Sequence[dict[str, Any]]) -> None:
    from openpyxl.styles import Alignment

    worksheet = workbook.create_sheet("处理异常")
    worksheet.append(EXCEPTION_HEADERS)
    row_count = _append_dict_rows(worksheet, EXCEPTION_HEADERS, exceptions)
    end_row = row_count + 1
    _style_grid(worksheet, max_row=end_row, max_column=len(EXCEPTION_HEADERS))
    worksheet.freeze_panes = "A2"
    _set_widths(worksheet, [14, 28, 14, 18, 42, 42, 42, 54, 58])
    for row_number in range(2, end_row + 1):
        worksheet.row_dimensions[row_number].height = 32
        for column in range(1, 10):
            worksheet.cell(row_number, column).alignment = Alignment(
                vertical="top",
                wrap_text=(column == 9),
            )
        for column in range(4, 8):
            worksheet.cell(row_number, column).number_format = "@"
    if row_count:
        _add_table(worksheet, f"A1:I{end_row}", "ProcessingExceptionsTable")


def _add_summary_sheet(workbook: Any, centers: Sequence[dict[str, Any]]) -> None:
    worksheet = workbook.create_sheet("运行摘要")
    worksheet.append(SUMMARY_HEADERS)
    for center in centers:
        counters = center.get("counters", {})
        worksheet.append(
            [
                center.get("sheet_name", ""),
                center.get("source_center_name", ""),
                center.get("reference_center_name", ""),
                counters.get("center_match_basis", ""),
                counters.get("files_seen", 0),
                counters.get("dicom", 0),
                counters.get("patients", 0),
                counters.get("studies", 0),
                counters.get("series", 0),
                counters.get("output_rows", 0),
                counters.get("matched_rows", 0),
                counters.get("exceptions", 0),
            ]
        )
    end_row = len(centers) + 1
    _style_grid(worksheet, max_row=end_row, max_column=len(SUMMARY_HEADERS))
    worksheet.freeze_panes = "A2"
    _set_widths(worksheet, [18, 18, 28, 30, 14, 12, 12, 12, 12, 14, 16, 12])
    for row in worksheet.iter_rows(min_row=2, max_row=end_row, min_col=5, max_col=12):
        for cell in row:
            cell.number_format = "#,##0"
    if centers:
        _add_table(worksheet, f"A1:L{end_row}", "RunSummaryTable")


def _add_definition_sheet(workbook: Any, payload: dict[str, Any]) -> None:
    from openpyxl.styles import Alignment

    worksheet = workbook.create_sheet("字段说明")
    worksheet.append(["字段/规则", "含义", "来源或计算方法"])
    for row in _definitions(payload):
        worksheet.append(row)
    _style_grid(worksheet, max_row=worksheet.max_row, max_column=3)
    worksheet.freeze_panes = "A2"
    _set_widths(worksheet, [23, 35, 68])
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_center_workbook_linux(output_path: Path, payload: dict[str, Any]) -> None:
    """Write one Linux-compatible image-web workbook atomically."""

    require_openpyxl()
    from openpyxl import Workbook

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = payload.get("headers", [])
    centers = payload.get("centers", [])
    for index, center in enumerate(centers):
        _add_center_sheet(workbook, center, headers, index)
    _add_exception_sheet(workbook, payload.get("exceptions", []))
    _add_summary_sheet(workbook, centers)
    _add_definition_sheet(workbook, payload)

    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=output_path.stem + "_",
        suffix=".xlsx",
        dir=str(output_path.parent),
    )
    os.close(file_descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        workbook.close()
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()
