#!/usr/bin/env python3
"""DICOM 转存 V3 Linux 入口。

与 ``dicom_organize_v3.py`` 使用同一套扫描、身份归组和复制规则，仅将 Excel
实现替换为公开的 openpyxl，并提供 ``--report-only`` 补报表模式。
"""

from __future__ import annotations

import argparse
import itertools
import json
import os
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Optional, Sequence

import dicom_organize_v3 as v3
from build_dicom_study_workbook_linux import (
    require_openpyxl,
    write_study_workbook_linux,
)


REPORT_SPECS = [
    {
        "sheet_name": "患者身份映射",
        "table_name": "PatientIdentityMapTable",
        "columns": v3.IDENTITY_MAPPING_COLUMNS,
        "widths": [16, 18, 17, 28, 24, 14, 11, 28, 18, 28, 25, 42, 38],
    },
    {
        "sheet_name": "患者身份待确认",
        "table_name": "PatientIdentityReviewTable",
        "columns": v3.IDENTITY_REVIEW_COLUMNS,
        "widths": [20, 17, 18, 17, 18, 16, 16, 11, 11, 28, 28, 18, 18, 48, 54],
    },
    {
        "sheet_name": "来源目录身份审计",
        "table_name": "SourceIdentityAuditTable",
        "columns": v3.SOURCE_AUDIT_COLUMNS,
        "widths": [32, 18, 18, 18, 18, 12, 12, 28, 42],
    },
    {
        "sheet_name": "转存异常",
        "table_name": "TransferExceptionTable",
        "columns": v3.v2.EXCEPTION_COLUMNS,
        "widths": [20, 17, 18, 38, 48, 46, 46],
    },
]


def _reports(
    mapping_rows: Sequence[dict[str, Any]],
    review_rows: Sequence[dict[str, Any]],
    source_audit_rows: Sequence[dict[str, Any]],
    exception_rows: Sequence[dict[str, Any]],
) -> list[dict[str, Any]]:
    row_groups = [mapping_rows, review_rows, source_audit_rows, exception_rows]
    return [
        {**spec, "rows": list(rows)}
        for spec, rows in zip(REPORT_SPECS, row_groups)
    ]


def _write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=path.stem + "_", suffix=".tmp", dir=str(path.parent)
    )
    os.close(file_descriptor)
    temporary = Path(temporary_name)
    try:
        with temporary.open("w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, separators=(",", ":"))
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _report_snapshot_path(destination_root: Path) -> Path:
    return destination_root / ".dicom_v3_state" / "report_snapshot.json"


def linux_write_study_workbook(
    output_path: Path,
    study_rows: Sequence[dict[str, Any]],
    mapping_rows: Sequence[dict[str, Any]],
    identity_review_rows: Sequence[dict[str, Any]],
    source_audit_rows: Sequence[dict[str, Any]],
    exception_rows: Sequence[dict[str, Any]],
    source_root: Path,
    destination_root: Path,
    counters: dict[str, int],
    manifest_csv_path: Optional[Path] = None,
    preview_dir: Optional[Path] = None,
) -> None:
    del preview_dir  # Linux生产入口不生成图片预览。
    study_rows = v3.sort_study_rows(study_rows, mapping_rows)
    reports = _reports(
        mapping_rows,
        identity_review_rows,
        source_audit_rows,
        exception_rows,
    )
    generated_at = datetime.now().isoformat(timespec="seconds")
    snapshot = {
        "rule_version": v3.RULE_VERSION,
        "generated_at": generated_at,
        "source_root": str(source_root),
        "destination_root": str(destination_root),
        "counters": dict(counters),
        "studies": [
            {
                key: value
                for key, value in row.items()
                if key != "series_summaries"
            }
            for row in study_rows
        ],
        "reports": reports,
        "manifest_csv_path": str(manifest_csv_path or ""),
    }
    _write_json_atomic(_report_snapshot_path(destination_root), snapshot)
    write_study_workbook_linux(
        output_path,
        study_rows,
        reports,
        counters,
        rule_version=v3.RULE_VERSION,
        generated_at=generated_at,
        manifest_csv_path=manifest_csv_path,
    )


def _excel_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).strftime(
                "%Y-%m-%d"
            )
        except (OverflowError, ValueError):
            return ""
    return v3._study_date(value)


def _read_sheet_matrix(path: Path, sheet_name: str) -> list[list[object]]:
    require_openpyxl()
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return []
        return [list(row) for row in workbook[sheet_name].iter_rows(values_only=True)]
    finally:
        workbook.close()


def load_identity_mapping_linux(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() != ".xlsx":
        return _ORIGINAL_LOAD_IDENTITY_MAPPING(path)
    if not path.is_file():
        return []
    matrix = _read_sheet_matrix(path, "患者身份映射")
    if not matrix:
        return []
    headers = [str(value or "").strip() for value in matrix[0]]
    required = {"master_patient_key", "canonical_name", "patient_id"}
    if not required.issubset(headers):
        return []
    rows: list[dict[str, str]] = []
    for values in matrix[1:]:
        source = {
            header: str(values[index] or "").strip() if index < len(values) else ""
            for index, header in enumerate(headers)
        }
        if source.get("patient_id"):
            rows.append(
                {column: source.get(column, "") for column in v3.IDENTITY_MAPPING_COLUMNS}
            )
    return rows


def load_study_inventory_linux(
    state_path: Path,
    excel_path: Path,
) -> list[dict[str, Any]]:
    if state_path.is_file():
        with state_path.open("r", encoding="utf-8") as stream:
            payload = json.load(stream)
        rows = payload.get("rows", []) if isinstance(payload, dict) else payload
        return [
            dict(row)
            for row in rows
            if isinstance(row, dict) and row.get("study_uid")
        ]
    if not excel_path.is_file():
        return []
    matrix = _read_sheet_matrix(excel_path, "影像检查明细")
    header_index = next(
        (
            index
            for index, values in enumerate(matrix)
            if "StudyInstanceUID" in {str(value or "").strip() for value in values}
        ),
        -1,
    )
    if header_index < 0:
        return []
    headers = [str(value or "").strip() for value in matrix[header_index]]
    rows: list[dict[str, Any]] = []
    for values in matrix[header_index + 1 :]:
        row: dict[str, Any] = {}
        for index, header in enumerate(headers):
            key = v3.STUDY_WORKBOOK_COLUMN_MAP.get(header)
            if key:
                row[key] = values[index] if index < len(values) else ""
        if not row.get("study_uid"):
            continue
        row["study_uid"] = str(row["study_uid"]).strip()
        row["exam_date"] = _excel_date(row.get("exam_date"))
        for key in ("series_count", "file_count"):
            try:
                row[key] = int(float(row.get(key) or 0))
            except (TypeError, ValueError):
                row[key] = 0
        try:
            row["confidence"] = float(row.get("confidence") or 0)
        except (TypeError, ValueError):
            row["confidence"] = 0.0
        rows.append(row)
    return rows


def _review_rows_from_mapping(
    mapping_rows: Sequence[dict[str, str]],
) -> list[dict[str, str]]:
    grouped: dict[str, list[dict[str, str]]] = {}
    for row in mapping_rows:
        if row.get("match_status") != "grouped_needs_review":
            continue
        grouped.setdefault(row.get("master_patient_key", ""), []).append(row)
    review_rows: list[dict[str, str]] = []
    for rows in grouped.values():
        for left, right in itertools.combinations(
            sorted(rows, key=lambda row: row.get("patient_id", "")), 2
        ):
            if left.get("birth_date") == right.get("birth_date"):
                continue
            review_rows.append(
                {
                    "review_status": "grouped_needs_review",
                    "left_patient_id": left.get("patient_id", ""),
                    "left_name": left.get("canonical_name", ""),
                    "right_patient_id": right.get("patient_id", ""),
                    "right_name": right.get("canonical_name", ""),
                    "left_birth_date": left.get("birth_date", ""),
                    "right_birth_date": right.get("birth_date", ""),
                    "left_sex": left.get("sex", ""),
                    "right_sex": right.get("sex", ""),
                    "left_institutions": left.get("institutions", ""),
                    "right_institutions": right.get("institutions", ""),
                    "left_estimated_birth_year": left.get("estimated_birth_year", ""),
                    "right_estimated_birth_year": right.get("estimated_birth_year", ""),
                    "match_basis": left.get("match_basis", "") or right.get("match_basis", ""),
                    "suggested_action": "已归入同一患者目录；出生日期冲突，仍需人工确认自然人身份",
                }
            )
    return review_rows


def report_only(argv: Sequence[str]) -> int:
    parser = argparse.ArgumentParser(
        description="从V3状态文件补生成Linux版综合Excel，不重新扫描或复制DICOM"
    )
    parser.add_argument("src_dir", help="原始输入目录（仅写入报表元信息）")
    parser.add_argument("dst_dir", help="V3转存目标目录")
    parser.add_argument("--report-only", action="store_true")
    parser.add_argument("--excel", help="输出Excel路径")
    parser.add_argument("--identity-map", help="身份映射JSON/CSV/XLSX路径")
    parser.add_argument("--manifest", nargs="?", const="AUTO")
    args, _ = parser.parse_known_args(argv)

    require_openpyxl()
    source_root = Path(args.src_dir).resolve()
    destination_root = Path(args.dst_dir).resolve()
    excel_path = (
        Path(args.excel).resolve()
        if args.excel
        else destination_root / "影像检查明细_V3.xlsx"
    )
    snapshot_path = _report_snapshot_path(destination_root)
    if snapshot_path.is_file():
        with snapshot_path.open("r", encoding="utf-8") as stream:
            snapshot = json.load(stream)
        studies = list(snapshot.get("studies", []))
        reports = list(snapshot.get("reports", []))
        mapping_rows = next(
            (
                list(report.get("rows", []))
                for report in reports
                if report.get("sheet_name") == "患者身份映射"
            ),
            [],
        )
        counters = dict(snapshot.get("counters", {}))
        generated_at = datetime.now().isoformat(timespec="seconds")
        manifest_text = str(snapshot.get("manifest_csv_path", "") or "")
        manifest_path = Path(manifest_text) if manifest_text else None
    else:
        study_state = destination_root / ".dicom_v3_state" / "study_inventory.json"
        studies = load_study_inventory_linux(study_state, excel_path)
        identity_path = (
            Path(args.identity_map).resolve()
            if args.identity_map
            else destination_root / ".dicom_v3_state" / "identity_mapping.json"
        )
        mapping_rows = load_identity_mapping_linux(identity_path)
        review_rows = _review_rows_from_mapping(mapping_rows)
        reports = _reports(mapping_rows, review_rows, [], [])
        counters = {
            "dicom": sum(int(row.get("file_count") or 0) for row in studies),
            "studies_current_run": 0,
            "studies": len(studies),
            "identity_groups": len(
                {
                    row.get("master_patient_key")
                    for row in mapping_rows
                    if row.get("master_patient_key")
                }
            ),
            "identity_review_pairs": len(review_rows),
            "mixed_source_records": 0,
        }
        generated_at = datetime.now().isoformat(timespec="seconds")
        manifest_path = (
            Path(args.manifest).resolve()
            if args.manifest and args.manifest != "AUTO"
            else None
        )

    studies = v3.sort_study_rows(studies, mapping_rows)

    if not studies and not any(report.get("rows") for report in reports):
        print(
            "补报表失败：目标目录中没有可用的V3状态。请确认"
            ".dicom_v3_state/identity_mapping.json和study_inventory.json存在。",
            file=sys.stderr,
        )
        return 1
    write_study_workbook_linux(
        excel_path,
        studies,
        reports,
        counters,
        rule_version=v3.RULE_VERSION,
        generated_at=generated_at,
        manifest_csv_path=manifest_path,
    )
    print("=" * 60)
    print("Linux版V3报表补生成完成（未扫描、未复制DICOM）")
    print(f"累计Study: {len(studies)}")
    print(f"综合Excel: {excel_path}")
    return 0


_ORIGINAL_LOAD_IDENTITY_MAPPING = v3.load_identity_mapping


def main() -> int:
    if "--report-only" in sys.argv[1:]:
        return report_only(sys.argv[1:])

    # 启动前检查，避免长时间扫描复制结束后才因Excel依赖缺失而失败。
    require_openpyxl()
    v3.write_study_workbook = linux_write_study_workbook
    v3.load_identity_mapping = load_identity_mapping_linux
    v3.load_study_inventory = load_study_inventory_linux
    return v3.main()


if __name__ == "__main__":
    raise SystemExit(main())
