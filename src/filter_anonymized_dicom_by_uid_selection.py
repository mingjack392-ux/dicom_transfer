#!/usr/bin/env python3
"""按Excel中的Series/SOP UID从已匿名化交付目录二次筛选DICOM。

输入目录必须是已完成匿名化后的标准层级::

    匿名编号/目标二级目录/StudyInstanceUID/SeriesInstanceUID/SOPInstanceUID.dcm

筛选规则::

* ``是否进入匿名化=是`` 的行参与筛选；
* ``SOPUID`` 为 ``NA``、空白或 ``/`` 时复制整条 Series；
* ``SOPUID`` 为具体值时只复制该 SOP；
* 目标目录完整保留源目录相对层级；
* 源文件只读，DICOM按字节复制，不重新序列化；
* 默认仅预检并写审计CSV，显式传入 ``--execute`` 才复制。

本程序是匿名化之后的下游筛选入口，不修改现有匿名化程序，也不替代
``filter_transferred_dicom_by_selection.py`` 的匿名化前筛选契约。
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import re
import sys
import threading
import unicodedata
import uuid
import warnings
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional, Sequence

from openpyxl import load_workbook
from pydicom import dcmread

import dicom_transfer_by_screening as transfer_core


RULE_VERSION = "2026.09.14-anonymized-uid-filter-v1.0"
NO_SOP_MARKERS = {"", "NA", "N/A", "NONE", "NULL", "/"}
YES_VALUES = {"是", "Y", "YES", "TRUE", "1"}
UID_PATH_PATTERN = re.compile(r"[0-9.]+")
MAX_COMPONENT_LENGTH = 240
HEADER_TAGS = ["StudyInstanceUID", "SeriesInstanceUID", "SOPInstanceUID"]
COPY_CHUNK_SIZE = 4 * 1024 * 1024
_COPY_LOCKS = tuple(threading.Lock() for _ in range(257))


@dataclass(frozen=True)
class SelectionRule:
    row_number: int
    anonymous_code: str
    target_folder: str
    study_uid: str
    series_uid: str
    sop_uid: str
    modality: str
    include_raw: str

    @property
    def included(self) -> bool:
        return normalized_text(self.include_raw).upper() in YES_VALUES

    @property
    def selection_level(self) -> str:
        return "SOP" if self.sop_uid else "Series"

    @property
    def key(self) -> tuple[str, str, str, str, str]:
        return (
            self.anonymous_code,
            self.target_folder,
            self.study_uid,
            self.series_uid,
            self.sop_uid,
        )

    @property
    def group_key(self) -> tuple[str, str]:
        return (self.anonymous_code, self.target_folder)


@dataclass(frozen=True)
class FileCandidate:
    rule: SelectionRule
    source: Path
    target: Path
    expected_sop_uid: str

    @property
    def validation_key(self) -> tuple[str, str, str, str]:
        return (
            str(self.source),
            self.rule.study_uid,
            self.rule.series_uid,
            self.expected_sop_uid,
        )


@dataclass
class RulePlan:
    rule: SelectionRule
    candidates: list[FileCandidate] = field(default_factory=list)
    issue_status: str = ""
    issue_message: str = ""


@dataclass(frozen=True)
class HeaderResult:
    status: str
    actual_study_uid: str = ""
    actual_series_uid: str = ""
    actual_sop_uid: str = ""
    file_size: int = 0
    message: str = ""


@dataclass(frozen=True)
class TargetState:
    status: str
    source_sha256: str = ""
    target_sha256: str = ""
    message: str = ""


@dataclass(frozen=True)
class CopyOutcome:
    status: str
    file_size: int = 0
    source_sha256: str = ""
    target_sha256: str = ""
    message: str = ""


@dataclass
class GroupStats:
    anonymous_code: str
    target_folder: str
    series_rules: int = 0
    sop_rules: int = 0
    matched_rules: int = 0
    missing_rules: int = 0
    invalid_rules: int = 0
    blocked_rules: int = 0
    redundant_selections: int = 0
    planned_files: int = 0
    copied: int = 0
    duplicate_same: int = 0
    conflicts: int = 0
    errors: int = 0


@dataclass
class RunSummary:
    sheet_name: str
    workbook_rows: int = 0
    included_rules: int = 0
    ignored_rules: int = 0
    series_rules: int = 0
    sop_rules: int = 0
    duplicate_rules: int = 0
    matched_rules: int = 0
    missing_rules: int = 0
    invalid_rules: int = 0
    blocked_rules: int = 0
    redundant_selections: int = 0
    selected_files: int = 0
    header_validated: int = 0
    unreadable_dicom: int = 0
    uid_mismatches: int = 0
    copied: int = 0
    duplicate_same: int = 0
    conflicts: int = 0
    errors: int = 0
    detail_audit_path: str = ""
    group_summary_path: str = ""

    @property
    def needs_attention(self) -> bool:
        return bool(
            self.missing_rules
            or self.invalid_rules
            or self.blocked_rules
            or self.unreadable_dicom
            or self.uid_mismatches
            or self.conflicts
            or self.errors
        )


DETAIL_HEADERS = [
    "Excel行号",
    "匿名编号",
    "目标二级目录",
    "modality",
    "筛选层级",
    "筛选StudyInstanceUID",
    "筛选SeriesInstanceUID",
    "筛选SOPInstanceUID",
    "源文件",
    "目标文件",
    "实际StudyInstanceUID",
    "实际SeriesInstanceUID",
    "实际SOPInstanceUID",
    "文件大小",
    "源SHA256",
    "目标SHA256",
    "状态",
    "说明",
]


GROUP_HEADERS = [
    "匿名编号",
    "目标二级目录",
    "Series规则数",
    "SOP规则数",
    "命中规则数",
    "未命中规则数",
    "无效规则数",
    "阻断规则数",
    "重复选择数",
    "计划或处理文件数",
    "成功复制",
    "内容重复",
    "目标冲突",
    "错误",
    "处理状态",
]


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized_text(value: str) -> str:
    return unicodedata.normalize("NFKC", value or "").strip()


def normalize_sop(value: str) -> str:
    normalized = normalized_text(value)
    return "" if normalized.upper() in NO_SOP_MARKERS else normalized


def choose_sheet(workbook: Any, requested_name: Optional[str]) -> Any:
    if requested_name:
        if requested_name not in workbook.sheetnames:
            raise ValueError(
                f"筛选表不存在工作表: {requested_name}；可用工作表: "
                + ", ".join(workbook.sheetnames)
            )
        return workbook[requested_name]
    for worksheet in workbook.worksheets:
        if worksheet.max_row >= 2 and worksheet.max_column >= 1:
            return worksheet
    raise ValueError("筛选表中没有可读取的数据工作表")


def resolve_header(headers: dict[str, int], *aliases: str) -> int:
    for alias in aliases:
        if alias in headers:
            return headers[alias]
    raise ValueError("筛选表缺少字段: " + "/".join(aliases))


def optional_header(headers: dict[str, int], *aliases: str) -> Optional[int]:
    for alias in aliases:
        if alias in headers:
            return headers[alias]
    return None


def row_value(row: Sequence[Any], index: Optional[int]) -> str:
    if index is None or index >= len(row):
        return ""
    return normalized_text(cell_text(row[index]))


def read_selection_rules(
    workbook_path: Path, sheet_name: Optional[str] = None
) -> tuple[str, list[SelectionRule]]:
    if not workbook_path.is_file():
        raise ValueError(f"筛选表不存在: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = choose_sheet(workbook, sheet_name)
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise ValueError("筛选表为空") from exc
        headers: dict[str, int] = {}
        for index, value in enumerate(header_row):
            header = normalized_text(cell_text(value))
            if header and header not in headers:
                headers[header] = index

        anonymous_index = resolve_header(headers, "匿名编号")
        target_index = resolve_header(headers, "目标二级目录")
        study_index = resolve_header(headers, "StudyInstanceUID")
        series_index = resolve_header(headers, "SeriesUID", "SeriesInstanceUID")
        sop_index = resolve_header(headers, "SOPUID", "SOPInstanceUID")
        include_index = resolve_header(headers, "是否进入匿名化")
        modality_index = optional_header(headers, "modality", "Modality", "影像类型")

        rules: list[SelectionRule] = []
        for row_number, row in enumerate(rows, start=2):
            if not any(normalized_text(cell_text(value)) for value in row):
                continue
            rules.append(
                SelectionRule(
                    row_number=row_number,
                    anonymous_code=row_value(row, anonymous_index),
                    target_folder=row_value(row, target_index),
                    study_uid=row_value(row, study_index),
                    series_uid=row_value(row, series_index),
                    sop_uid=normalize_sop(row_value(row, sop_index)),
                    modality=row_value(row, modality_index),
                    include_raw=row_value(row, include_index),
                )
            )
        return worksheet.title, rules
    finally:
        workbook.close()


def component_problem(label: str, value: str) -> str:
    if not value:
        return f"{label}为空"
    if value in {".", ".."}:
        return f"{label}不能为点目录"
    if "/" in value or "\\" in value or "\x00" in value:
        return f"{label}包含路径分隔符或NUL"
    if len(value) > MAX_COMPONENT_LENGTH:
        return f"{label}超过安全路径长度{MAX_COMPONENT_LENGTH}"
    return ""


def uid_problem(label: str, value: str) -> str:
    if not value:
        return f"{label}为空"
    if not UID_PATH_PATTERN.fullmatch(value):
        return f"{label}不是ASCII数字点号格式"
    if value.startswith(".") or value.endswith("."):
        return f"{label}以点开头或结尾"
    if len(value) > MAX_COMPONENT_LENGTH:
        return f"{label}超过安全路径长度{MAX_COMPONENT_LENGTH}"
    return ""


def uid_warnings(label: str, value: str) -> list[str]:
    result: list[str] = []
    if len(value) > 64:
        result.append(f"{label}长度{len(value)}超过DICOM标准上限64，仅记录")
    if ".." in value:
        result.append(f"{label}含连续点，仅记录")
    return result


def selection_issue(rule: SelectionRule) -> str:
    problems = [
        component_problem("匿名编号", rule.anonymous_code),
        component_problem("目标二级目录", rule.target_folder),
        uid_problem("StudyInstanceUID", rule.study_uid),
        uid_problem("SeriesInstanceUID", rule.series_uid),
    ]
    if rule.sop_uid:
        problems.append(uid_problem("SOPInstanceUID", rule.sop_uid))
    return "；".join(problem for problem in problems if problem)


def selection_warnings(rule: SelectionRule) -> list[str]:
    result = [
        *uid_warnings("StudyInstanceUID", rule.study_uid),
        *uid_warnings("SeriesInstanceUID", rule.series_uid),
    ]
    if rule.sop_uid:
        result.extend(uid_warnings("SOPInstanceUID", rule.sop_uid))
    return result


def is_within(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def validate_control_path(
    source_root: Path, destination_root: Path, control_dir: Path
) -> None:
    source = source_root.resolve()
    destination = destination_root.resolve()
    control = control_dir.resolve()
    if control == source or is_within(control, source) or is_within(source, control):
        raise ValueError("控制目录必须与匿名化源目录相互独立")
    if control == destination or is_within(control, destination) or is_within(destination, control):
        raise ValueError("控制目录必须与筛选结果目录相互独立")


def safe_resolved_source(path: Path, source_root: Path) -> tuple[Optional[Path], str]:
    try:
        resolved = path.resolve()
    except OSError as exc:
        return None, f"源路径无法解析: {exc}"
    root = source_root.resolve()
    if not is_within(resolved, root):
        return None, "源路径解析后越出匿名化源目录"
    return resolved, ""


def detail_base(rule: SelectionRule) -> dict[str, Any]:
    return {
        "Excel行号": rule.row_number,
        "匿名编号": rule.anonymous_code,
        "目标二级目录": rule.target_folder,
        "modality": rule.modality,
        "筛选层级": rule.selection_level,
        "筛选StudyInstanceUID": rule.study_uid,
        "筛选SeriesInstanceUID": rule.series_uid,
        "筛选SOPInstanceUID": rule.sop_uid or "NA",
        "源文件": "",
        "目标文件": "",
        "实际StudyInstanceUID": "",
        "实际SeriesInstanceUID": "",
        "实际SOPInstanceUID": "",
        "文件大小": "",
        "源SHA256": "",
        "目标SHA256": "",
        "状态": "",
        "说明": "",
    }


def candidate_detail(
    candidate: FileCandidate,
    result: Optional[HeaderResult] = None,
    *,
    status: str,
    message: str = "",
    source_sha256: str = "",
    target_sha256: str = "",
) -> dict[str, Any]:
    row = detail_base(candidate.rule)
    row.update(
        {
            "源文件": str(candidate.source),
            "目标文件": str(candidate.target),
            "实际StudyInstanceUID": result.actual_study_uid if result else "",
            "实际SeriesInstanceUID": result.actual_series_uid if result else "",
            "实际SOPInstanceUID": result.actual_sop_uid if result else "",
            "文件大小": result.file_size if result else "",
            "源SHA256": source_sha256,
            "目标SHA256": target_sha256,
            "状态": status,
            "说明": message,
        }
    )
    return row


def build_rule_plan(
    rule: SelectionRule, source_root: Path, destination_root: Path
) -> RulePlan:
    relative_series = (
        Path(rule.anonymous_code)
        / rule.target_folder
        / rule.study_uid
        / rule.series_uid
    )
    source_series = source_root / relative_series
    resolved_series, path_problem = safe_resolved_source(source_series, source_root)
    if path_problem:
        return RulePlan(rule, issue_status="unsafe_source_path", issue_message=path_problem)
    if resolved_series is None or not resolved_series.is_dir():
        return RulePlan(
            rule,
            issue_status="series_directory_not_found",
            issue_message="源匿名化目录中不存在对应Study/Series目录",
        )

    if rule.sop_uid:
        source_file = source_series / f"{rule.sop_uid}.dcm"
        resolved_file, file_problem = safe_resolved_source(source_file, source_root)
        if file_problem:
            return RulePlan(
                rule, issue_status="unsafe_source_path", issue_message=file_problem
            )
        if resolved_file is None or not resolved_file.is_file():
            return RulePlan(
                rule,
                issue_status="sop_file_not_found",
                issue_message="源匿名化目录中不存在对应SOP文件，不回退为整Series",
            )
        target = destination_root / relative_series / source_file.name
        return RulePlan(
            rule,
            candidates=[
                FileCandidate(
                    rule=rule,
                    source=source_file,
                    target=target,
                    expected_sop_uid=rule.sop_uid,
                )
            ],
        )

    source_files = sorted(
        (
            path
            for path in source_series.iterdir()
            if path.is_file() and path.suffix.casefold() == ".dcm"
        ),
        key=lambda path: path.name.casefold(),
    )
    if not source_files:
        return RulePlan(
            rule,
            issue_status="empty_series_directory",
            issue_message="对应Series目录中没有直接存放的.dcm文件",
        )
    invalid_names = [
        path.name for path in source_files if uid_problem("SOP文件名", path.stem)
    ]
    if invalid_names:
        display = "、".join(invalid_names[:5])
        if len(invalid_names) > 5:
            display += f"等{len(invalid_names)}个"
        return RulePlan(
            rule,
            issue_status="invalid_sop_filename",
            issue_message=f"整Series含非标准UID文件名，未复制该Series: {display}",
        )
    return RulePlan(
        rule,
        candidates=[
            FileCandidate(
                rule=rule,
                source=path,
                target=destination_root / relative_series / path.name,
                expected_sop_uid=path.stem,
            )
            for path in source_files
        ],
    )


def validate_candidate(candidate: FileCandidate, source_root: Path) -> HeaderResult:
    resolved, path_problem = safe_resolved_source(candidate.source, source_root)
    if path_problem:
        return HeaderResult("unsafe_source_path", message=path_problem)
    if resolved is None or not resolved.is_file():
        return HeaderResult("source_file_missing", message="预检期间源文件消失")
    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            dataset = dcmread(
                str(candidate.source),
                stop_before_pixels=True,
                force=True,
                specific_tags=HEADER_TAGS,
            )
        actual_study = normalized_text(cell_text(getattr(dataset, "StudyInstanceUID", "")))
        actual_series = normalized_text(cell_text(getattr(dataset, "SeriesInstanceUID", "")))
        actual_sop = normalized_text(cell_text(getattr(dataset, "SOPInstanceUID", "")))
        file_size = candidate.source.stat().st_size
    except Exception as exc:
        return HeaderResult("unreadable_dicom", message=str(exc))

    mismatches: list[str] = []
    expected = (
        ("StudyInstanceUID", candidate.rule.study_uid, actual_study),
        ("SeriesInstanceUID", candidate.rule.series_uid, actual_series),
        ("SOPInstanceUID", candidate.expected_sop_uid, actual_sop),
    )
    for label, expected_value, actual_value in expected:
        if not actual_value:
            mismatches.append(f"DICOM头{label}为空")
        elif actual_value != expected_value:
            mismatches.append(
                f"{label}不一致: 路径/清单={expected_value}，DICOM头={actual_value}"
            )
    warning_messages = [str(item.message) for item in caught]
    warning_messages.extend(selection_warnings(candidate.rule))
    warning_messages.extend(uid_warnings("实际SOPInstanceUID", actual_sop))
    warning_text = "；".join(dict.fromkeys(message for message in warning_messages if message))
    if mismatches:
        message = "；".join(mismatches)
        if warning_text:
            message += "；" + warning_text
        return HeaderResult(
            "uid_mismatch",
            actual_study,
            actual_series,
            actual_sop,
            file_size,
            message,
        )
    return HeaderResult(
        "validated",
        actual_study,
        actual_series,
        actual_sop,
        file_size,
        warning_text,
    )


def hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while True:
            block = stream.read(COPY_CHUNK_SIZE)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def target_path_problem(target: Path, destination_root: Path) -> str:
    try:
        resolved_root = destination_root.resolve()
        resolved_parent = target.parent.resolve()
    except OSError as exc:
        return f"目标路径无法解析: {exc}"
    if not is_within(resolved_parent, resolved_root):
        return "目标父目录解析后越出筛选结果根目录"
    if target.exists():
        try:
            resolved_target = target.resolve()
        except OSError as exc:
            return f"已存在目标无法解析: {exc}"
        if not is_within(resolved_target, resolved_root):
            return "已存在目标解析后越出筛选结果根目录"
    return ""


def inspect_existing_target(
    candidate: FileCandidate, destination_root: Path
) -> TargetState:
    target = candidate.target
    path_problem = target_path_problem(target, destination_root)
    if path_problem:
        return TargetState("error", message=path_problem)
    if not target.exists():
        return TargetState("missing")
    if not target.is_file():
        return TargetState("error", message="目标路径已存在但不是普通文件")
    try:
        source_sha = hash_file(candidate.source)
        target_sha = hash_file(target)
    except Exception as exc:
        return TargetState("error", message=f"目标冲突预检失败: {exc}")
    if candidate.source.stat().st_size == target.stat().st_size and source_sha == target_sha:
        return TargetState("duplicate_same", source_sha, target_sha)
    return TargetState(
        "conflict",
        source_sha,
        target_sha,
        "目标路径已有不同内容；未覆盖，也未写入冲突副本",
    )


def _copy_lock(target: Path) -> threading.Lock:
    return _COPY_LOCKS[hash(str(target)) % len(_COPY_LOCKS)]


def _fsync_directory(directory: Path) -> None:
    if os.name == "nt":
        return
    descriptor = os.open(str(directory), os.O_RDONLY)
    try:
        os.fsync(descriptor)
    finally:
        os.close(descriptor)


def copy_new_atomic(
    candidate: FileCandidate, destination_root: Path, durable: bool = False
) -> CopyOutcome:
    source = candidate.source
    target = candidate.target
    temporary: Optional[Path] = None
    try:
        with _copy_lock(target):
            existing = inspect_existing_target(candidate, destination_root)
            if existing.status != "missing":
                return CopyOutcome(
                    existing.status,
                    source.stat().st_size if source.is_file() else 0,
                    existing.source_sha256,
                    existing.target_sha256,
                    existing.message,
                )
            target.parent.mkdir(parents=True, exist_ok=True)
            path_problem = target_path_problem(target, destination_root)
            if path_problem:
                return CopyOutcome("error", message=path_problem)

            initial_stat = source.stat()
            temporary = target.parent / f".{target.name}.{uuid.uuid4().hex}.part"
            digest = hashlib.sha256()
            copied_size = 0
            with source.open("rb") as input_stream, temporary.open("xb") as output_stream:
                while True:
                    block = input_stream.read(COPY_CHUNK_SIZE)
                    if not block:
                        break
                    output_stream.write(block)
                    digest.update(block)
                    copied_size += len(block)
                output_stream.flush()
                if durable:
                    os.fsync(output_stream.fileno())
            final_stat = source.stat()
            if (
                copied_size != initial_stat.st_size
                or final_stat.st_size != initial_stat.st_size
                or final_stat.st_mtime_ns != initial_stat.st_mtime_ns
            ):
                return CopyOutcome("error", message="复制期间源文件发生变化")
            source_sha = digest.hexdigest()
            try:
                os.link(str(temporary), str(target))
            except FileExistsError:
                existing = inspect_existing_target(candidate, destination_root)
                return CopyOutcome(
                    existing.status,
                    copied_size,
                    source_sha or existing.source_sha256,
                    existing.target_sha256,
                    existing.message,
                )
            except OSError as exc:
                return CopyOutcome(
                    "error",
                    copied_size,
                    source_sha,
                    message="目标文件系统不支持安全的原子无覆盖写入: " + str(exc),
                )
            temporary.unlink()
            temporary = None
            if durable:
                _fsync_directory(target.parent)
            return CopyOutcome(
                "copied", copied_size, source_sha, source_sha, "按字节复制完成"
            )
    except Exception as exc:
        return CopyOutcome("error", message=str(exc))
    finally:
        if temporary is not None and temporary.exists():
            try:
                temporary.unlink()
            except OSError:
                pass


def unique_csv_path(directory: Path, prefix: str, timestamp: str) -> Path:
    candidate = directory / f"{prefix}_{timestamp}.csv"
    if not candidate.exists():
        return candidate
    return directory / f"{prefix}_{timestamp}_{uuid.uuid4().hex[:8]}.csv"


def write_csv(path: Path, headers: Sequence[str], rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(headers), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def group_stats_for(
    groups: dict[tuple[str, str], GroupStats], rule: SelectionRule
) -> GroupStats:
    if rule.group_key not in groups:
        groups[rule.group_key] = GroupStats(rule.anonymous_code, rule.target_folder)
    return groups[rule.group_key]


def run_filter(
    source_root: Path,
    workbook_path: Path,
    destination_root: Path,
    *,
    sheet_name: Optional[str] = None,
    execute: bool = False,
    workers: int = 4,
    copy_workers: int = 2,
    control_dir: Optional[Path] = None,
    progress_every: int = 500,
    durable: bool = False,
) -> RunSummary:
    source_root = Path(os.path.abspath(source_root))
    workbook_path = Path(os.path.abspath(workbook_path))
    destination_root = Path(os.path.abspath(destination_root))
    control_dir = Path(
        os.path.abspath(
            control_dir
            or destination_root.parent / f"{destination_root.name}_筛选控制"
        )
    )
    transfer_core.validate_paths(source_root, destination_root)
    validate_control_path(source_root, destination_root, control_dir)

    actual_sheet, all_rules = read_selection_rules(workbook_path, sheet_name)
    summary = RunSummary(sheet_name=actual_sheet, workbook_rows=len(all_rules))
    detail_rows: list[dict[str, Any]] = []
    groups: dict[tuple[str, str], GroupStats] = {}
    selected_rules: list[SelectionRule] = []
    for rule in all_rules:
        if not rule.included:
            summary.ignored_rules += 1
            row = detail_base(rule)
            row.update(
                {
                    "状态": "ignored_not_selected",
                    "说明": f"是否进入匿名化={rule.include_raw or '<空>'}，未参与筛选",
                }
            )
            detail_rows.append(row)
            continue
        selected_rules.append(rule)
        summary.included_rules += 1
        group = group_stats_for(groups, rule)
        if rule.selection_level == "Series":
            summary.series_rules += 1
            group.series_rules += 1
        else:
            summary.sop_rules += 1
            group.sop_rules += 1

    seen_rules: set[tuple[str, str, str, str, str]] = set()
    plans: list[RulePlan] = []
    for rule in selected_rules:
        group = group_stats_for(groups, rule)
        if rule.key in seen_rules:
            summary.duplicate_rules += 1
            row = detail_base(rule)
            row.update(
                {
                    "状态": "duplicate_rule_ignored",
                    "说明": "与前面的Excel规则完全重复，已去重",
                }
            )
            detail_rows.append(row)
            continue
        seen_rules.add(rule.key)
        issue = selection_issue(rule)
        if issue:
            summary.invalid_rules += 1
            group.invalid_rules += 1
            row = detail_base(rule)
            row.update({"状态": "invalid_selection_rule", "说明": issue})
            detail_rows.append(row)
            continue
        plan = build_rule_plan(rule, source_root, destination_root)
        if plan.issue_status:
            summary.missing_rules += 1
            group.missing_rules += 1
            row = detail_base(rule)
            row.update({"状态": plan.issue_status, "说明": plan.issue_message})
            detail_rows.append(row)
            continue
        plans.append(plan)

    unique_candidates: dict[tuple[str, str, str, str], FileCandidate] = {}
    for plan in plans:
        for candidate in plan.candidates:
            unique_candidates.setdefault(candidate.validation_key, candidate)
    validation_items = list(unique_candidates.items())
    validations: dict[tuple[str, str, str, str], HeaderResult] = {}
    if validation_items:
        with ThreadPoolExecutor(max_workers=max(1, workers)) as executor:
            results = executor.map(
                lambda item: validate_candidate(item[1], source_root), validation_items
            )
            for index, ((key, _candidate), result) in enumerate(
                zip(validation_items, results), start=1
            ):
                validations[key] = result
                if progress_every > 0 and index % progress_every == 0:
                    print(f"[UID预检] 已验证 {index}/{len(validation_items)} 个DICOM")

    validation_counts = Counter(result.status for result in validations.values())
    summary.header_validated = validation_counts["validated"]
    summary.unreadable_dicom = validation_counts["unreadable_dicom"]
    summary.uid_mismatches = validation_counts["uid_mismatch"]
    summary.errors += sum(
        count
        for status, count in validation_counts.items()
        if status not in {"validated", "unreadable_dicom", "uid_mismatch"}
    )

    seen_targets: set[str] = set()
    task_groups: list[tuple[SelectionRule, list[FileCandidate]]] = []
    for plan in plans:
        group = group_stats_for(groups, plan.rule)
        plan_results = [validations[item.validation_key] for item in plan.candidates]
        invalid_results = [result for result in plan_results if result.status != "validated"]
        if invalid_results:
            summary.blocked_rules += 1
            group.blocked_rules += 1
            for candidate, result in zip(plan.candidates, plan_results):
                status = (
                    result.status
                    if result.status != "validated"
                    else "blocked_by_rule_validation"
                )
                message = result.message
                if result.status == "validated":
                    message = "同一Excel规则内存在不可读或UID不一致文件，整条规则未复制"
                detail_rows.append(
                    candidate_detail(candidate, result, status=status, message=message)
                )
            continue

        summary.matched_rules += 1
        group.matched_rules += 1
        group_tasks: list[FileCandidate] = []
        for candidate, result in zip(plan.candidates, plan_results):
            target_key = os.path.normcase(os.path.abspath(candidate.target))
            if target_key in seen_targets:
                summary.redundant_selections += 1
                group.redundant_selections += 1
                detail_rows.append(
                    candidate_detail(
                        candidate,
                        result,
                        status="redundant_selection",
                        message="该文件已由前面的规则选中，未重复计划",
                    )
                )
                continue
            seen_targets.add(target_key)
            group_tasks.append(candidate)
            group.planned_files += 1
        if group_tasks:
            task_groups.append((plan.rule, group_tasks))

    summary.selected_files = sum(len(tasks) for _rule, tasks in task_groups)

    if not execute:
        for _rule, tasks in task_groups:
            for candidate in tasks:
                result = validations[candidate.validation_key]
                detail_rows.append(
                    candidate_detail(
                        candidate,
                        result,
                        status="planned",
                        message=result.message,
                    )
                )
    else:
        copy_queue: list[FileCandidate] = []
        for rule, tasks in task_groups:
            group = group_stats_for(groups, rule)
            states = [
                inspect_existing_target(candidate, destination_root)
                for candidate in tasks
            ]
            blocking = [state for state in states if state.status in {"conflict", "error"}]
            if blocking:
                summary.blocked_rules += 1
                group.blocked_rules += 1
                for candidate, state in zip(tasks, states):
                    result = validations[candidate.validation_key]
                    if state.status == "conflict":
                        summary.conflicts += 1
                        group.conflicts += 1
                        detail_rows.append(
                            candidate_detail(
                                candidate,
                                result,
                                status="conflict",
                                message=state.message,
                                source_sha256=state.source_sha256,
                                target_sha256=state.target_sha256,
                            )
                        )
                    elif state.status == "error":
                        summary.errors += 1
                        group.errors += 1
                        detail_rows.append(
                            candidate_detail(
                                candidate,
                                result,
                                status="error",
                                message=state.message,
                                source_sha256=state.source_sha256,
                                target_sha256=state.target_sha256,
                            )
                        )
                    else:
                        detail_rows.append(
                            candidate_detail(
                                candidate,
                                result,
                                status="blocked_by_rule_conflict",
                                message="同一Excel规则存在目标冲突，整条规则未新增复制",
                                source_sha256=state.source_sha256,
                                target_sha256=state.target_sha256,
                            )
                        )
                continue
            for candidate, state in zip(tasks, states):
                result = validations[candidate.validation_key]
                if state.status == "duplicate_same":
                    summary.duplicate_same += 1
                    group.duplicate_same += 1
                    detail_rows.append(
                        candidate_detail(
                            candidate,
                            result,
                            status="duplicate_same",
                            message="目标已有相同内容，未重复复制",
                            source_sha256=state.source_sha256,
                            target_sha256=state.target_sha256,
                        )
                    )
                else:
                    copy_queue.append(candidate)

        if copy_queue:
            with ThreadPoolExecutor(max_workers=max(1, copy_workers)) as executor:
                outcomes = executor.map(
                    lambda item: copy_new_atomic(item, destination_root, durable),
                    copy_queue,
                )
                for index, (candidate, outcome) in enumerate(
                    zip(copy_queue, outcomes), start=1
                ):
                    group = group_stats_for(groups, candidate.rule)
                    result = validations[candidate.validation_key]
                    if outcome.status == "copied":
                        summary.copied += 1
                        group.copied += 1
                    elif outcome.status == "duplicate_same":
                        summary.duplicate_same += 1
                        group.duplicate_same += 1
                    elif outcome.status == "conflict":
                        summary.conflicts += 1
                        group.conflicts += 1
                    else:
                        summary.errors += 1
                        group.errors += 1
                    detail_rows.append(
                        candidate_detail(
                            candidate,
                            result,
                            status=outcome.status,
                            message=outcome.message,
                            source_sha256=outcome.source_sha256,
                            target_sha256=outcome.target_sha256,
                        )
                    )
                    if progress_every > 0 and index % progress_every == 0:
                        print(f"[复制进度] 已处理 {index}/{len(copy_queue)} 个DICOM")

    group_rows: list[dict[str, Any]] = []
    for key in sorted(groups, key=lambda value: (value[0], value[1])):
        item = groups[key]
        attention = bool(
            item.missing_rules
            or item.invalid_rules
            or item.blocked_rules
            or item.conflicts
            or item.errors
        )
        group_rows.append(
            {
                "匿名编号": item.anonymous_code,
                "目标二级目录": item.target_folder,
                "Series规则数": item.series_rules,
                "SOP规则数": item.sop_rules,
                "命中规则数": item.matched_rules,
                "未命中规则数": item.missing_rules,
                "无效规则数": item.invalid_rules,
                "阻断规则数": item.blocked_rules,
                "重复选择数": item.redundant_selections,
                "计划或处理文件数": item.planned_files,
                "成功复制": item.copied,
                "内容重复": item.duplicate_same,
                "目标冲突": item.conflicts,
                "错误": item.errors,
                "处理状态": (
                    "needs_review"
                    if attention
                    else ("completed" if execute else "preview_passed")
                ),
            }
        )

    timestamp = datetime.now().strftime("%Y%m%dT%H%M%S_%f")
    detail_path = unique_csv_path(control_dir, "匿名化后UID筛选明细", timestamp)
    summary_path = unique_csv_path(control_dir, "匿名化后UID筛选汇总", timestamp)
    detail_rows.sort(
        key=lambda row: (
            int(row.get("Excel行号") or 0),
            str(row.get("源文件") or ""),
            str(row.get("状态") or ""),
        )
    )
    write_csv(detail_path, DETAIL_HEADERS, detail_rows)
    write_csv(summary_path, GROUP_HEADERS, group_rows)
    summary.detail_audit_path = str(detail_path)
    summary.group_summary_path = str(summary_path)
    return summary


def print_summary(summary: RunSummary, execute: bool) -> None:
    print(f"规则版本: {RULE_VERSION}")
    print(f"模式: {'执行' if execute else '预检'}")
    print(f"工作表: {summary.sheet_name}")
    print(f"Excel数据行: {summary.workbook_rows}")
    print(
        f"参与筛选规则: {summary.included_rules} "
        f"(Series={summary.series_rules}, SOP={summary.sop_rules})"
    )
    print(f"未选择规则: {summary.ignored_rules}")
    print(f"完全重复规则: {summary.duplicate_rules}")
    print(f"命中规则: {summary.matched_rules}")
    print(f"未命中规则: {summary.missing_rules}")
    print(f"无效规则: {summary.invalid_rules}")
    print(f"阻断规则: {summary.blocked_rules}")
    print(f"重复选择文件: {summary.redundant_selections}")
    print(f"计划或处理DICOM: {summary.selected_files}")
    print(f"DICOM头验证通过: {summary.header_validated}")
    print(f"不可读DICOM: {summary.unreadable_dicom}")
    print(f"UID不一致: {summary.uid_mismatches}")
    if execute:
        print(f"成功复制: {summary.copied}")
        print(f"内容重复: {summary.duplicate_same}")
        print(f"目标冲突: {summary.conflicts}")
    print(f"错误: {summary.errors}")
    print(f"明细审计: {summary.detail_audit_path}")
    print(f"目录汇总: {summary.group_summary_path}")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="按Excel中的Series/SOP UID筛选已匿名化DICOM，并保持目录层级不变"
    )
    parser.add_argument(
        "source_root",
        help="已匿名化交付影像根目录，层级为匿名编号/目标二级目录/Study/Series/SOP.dcm",
    )
    parser.add_argument("selection_xlsx", help="包含匿名编号、目标二级目录和UID的Excel")
    parser.add_argument("destination_root", help="独立筛选结果根目录")
    parser.add_argument("--sheet", help="工作表名称；默认读取第一个非空工作表")
    parser.add_argument(
        "--control-dir",
        help="独立审计目录；默认使用目标目录同级的<目标目录名>_筛选控制",
    )
    parser.add_argument("--workers", type=int, default=4, help="DICOM头验证线程数，默认4")
    parser.add_argument("--copy-workers", type=int, default=2, help="文件复制线程数，默认2")
    parser.add_argument(
        "--progress-every", type=int, default=500, help="每处理多少文件显示进度，0表示关闭"
    )
    parser.add_argument("--durable", action="store_true", help="复制时执行fsync，更安全但更慢")
    parser.add_argument("--execute", action="store_true", help="正式复制；默认仅预检")
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    args = build_parser().parse_args(argv)
    try:
        summary = run_filter(
            Path(args.source_root),
            Path(args.selection_xlsx),
            Path(args.destination_root),
            sheet_name=args.sheet,
            execute=args.execute,
            workers=max(1, args.workers),
            copy_workers=max(1, args.copy_workers),
            control_dir=Path(args.control_dir) if args.control_dir else None,
            progress_every=max(0, args.progress_every),
            durable=args.durable,
        )
    except Exception as exc:
        print(f"筛选失败: {exc}", file=sys.stderr)
        return 2
    print_summary(summary, args.execute)
    return 1 if summary.needs_attention else 0


if __name__ == "__main__":
    raise SystemExit(main())
